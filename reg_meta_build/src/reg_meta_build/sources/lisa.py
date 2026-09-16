"""Reader for SCB's pinned 2024/2025 LISA variable-list workbook."""

from __future__ import annotations

import re
from dataclasses import dataclass
from typing import TYPE_CHECKING, Any, Literal

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

from reg_meta_build.input_snapshot import SnapshotError
from reg_meta_build.normalization import normalize_text, normalize_token
from reg_meta_build.source_records import (
    NativeCoordinates,
    RecordLocator,
    ScopeInterval,
    SourceCoordinate,
    SourceField,
    SourceFields,
    SourceRecord,
    SourceRevision,
    SourceSubject,
    TemporalScope,
    value_field,
)

if TYPE_CHECKING:
    from collections.abc import Mapping
    from pathlib import Path


class LisaWorkbookError(SnapshotError):
    """The selected workbook differs from the one supported delivery layout."""


@dataclass(frozen=True)
class _ContextRow:
    kind: Literal["continuation-note", "worksheet-footnote"]
    cells: tuple[object | None, ...]
    preceding_declaration: tuple[int, str] | None = None


@dataclass(frozen=True)
class _TableSpec:
    title: str
    headers: tuple[str, ...]
    table_key: str
    period_column: int | None
    register_column: int
    text_rows: Mapping[int, str]
    sections: Mapping[int, tuple[str, str]]
    context_rows: Mapping[int, _ContextRow]
    population_rows: tuple[int, ...] = ()


@dataclass(frozen=True)
class _LisaSourceRead:
    records: tuple[SourceRecord, ...]
    worksheet_context: tuple[str, ...]


_INDIVID_TEXT_ROWS = {
    5: "Demografiska variabler",
    6: "Årsoberoende variabler, baseras på uppgifter avseende senast året av LISA",
    7: "Födelseland",
    31: "Avlidna",
    34: "In- och utvandring",
    70: "För hushåll:",
    105: "För familj:",
    120: "Utbildningsvariabler",
    138: "Sysselsättningsvariabler",
    139: (
        "Befolkningens arbetsmarknadsstatus (BAS) är källa från 2022 om inget "
        "annat år anges"
    ),
    201: "Yrke 7)",
    218: "För största förvärvskälla:",
    284: "För näst största förvärvskälla:",
    351: "För tredje största förvärvskälla:",
    405: "Inkomstvariabler",
    406: "Inkomst av förvärvskälla",
    446: "Arbetstidsrelaterad social inkomst",
    447: "Studier:",
    468: "Värnplikt/Grundläggande militär utbildning:",
    477: "Föräldraledighet:",
    503: "Sjukdom/Arbetsskada/Rehabilitering:",
    573: "Arbetslöshet:",
    587: "Arbetsmarknadspolitisk åtgärd:",
    613: "Sjukersättning (förtidspension)/Aktivitetsersättning (sjukbidrag):",
    648: "Delpension:",
    652: "Övrig arbetstidrelaterad social inkomst:",
    658: "Övrig inkomst",
    659: "Kapitalinkomst",
    663: "Annan inkomst",
    664: "Pension:",
    681: "Tjänstepension:",
    694: "Övrig pension:",
    699: "Summa inkomst av pensioner:",
    703: "Yrkes- och arbetsskadelivränta:",
    709: "Annan livränta:",
    712: "Efterlevandeförmån:",
    719: "Familjerelaterade inkomster:",
    732: "Övriga inkomster/ersättningar:",
    749: "Disponibel inkomst",
    766: "Huvudsaklig inkomstkälla, Befolkningens abetsmarknadsstatus",
    788: "Registerbaserad aktivitetsstatistik, RAKS",
}

_INDIVID_CONTEXT_ROWS = {
    323: _ContextRow(
        kind="continuation-note",
        cells=(
            None,
            (
                "Före 2010 finns inte kod 5 (företagare i eget AB). För åren "
                "1993- finns istället variabeln KU1Faman som anges som 1 om "
                "personen är en företagare i eget AB."
            ),
            None,
            None,
            None,
            "RAMS-Jobb",
        ),
        preceding_declaration=(322, "KU2YrkStalln"),
    ),
    815: _ContextRow(
        kind="worksheet-footnote",
        cells=(None, "_ftnref2", None, None, None, None),
    ),
    816: _ContextRow(
        kind="worksheet-footnote",
        cells=(
            None,
            (
                "1 Markering för variabler som innehåller känsliga personuppgifter "
                "(Ja/Nej/I vissa fall). Variabeln betraktas som en känslig "
                "personuppgift och kräver godkänd etikprövning hos "
                "Etikprövningsmyndigheten om utlämnandet görs för forskningsändamål."
            ),
            None,
            None,
            None,
            None,
        ),
    ),
    817: _ContextRow(
        kind="worksheet-footnote",
        cells=(
            None,
            (
                "2 Variabeln finns i två typer, kontakta SCB:s handläggare för "
                "vidare information"
            ),
            None,
            None,
            None,
            None,
        ),
    ),
    818: _ContextRow(
        kind="worksheet-footnote",
        cells=(
            None,
            (
                "5 Fr.o.m. 2014 fördelas familjeinkomster mellan sökande och "
                "medsökande, tidigare ingick dessa ersättningar endast för den "
                "sökande"
            ),
            None,
            None,
            None,
            None,
        ),
    ),
    819: _ContextRow(
        kind="worksheet-footnote",
        cells=(
            None,
            (
                "7 Från och med årgång 2020 används inte längre yrkeskoderna 01, "
                "02, 03 (militärer) och 3360 (poliser), individer med dessa "
                "yrkeskoder ingår i stället i gruppen med okänt yrke."
            ),
            None,
            None,
            None,
            None,
        ),
    ),
}

_TABLES: dict[str, _TableSpec] = {
    "Individ": _TableSpec(
        title="Individvariabler",
        headers=(
            "Kolumnnamn",
            "Beskrivning",
            "År",
            "Register",
            "Känslig1",
            "Grundregister",
        ),
        table_key="individual",
        period_column=3,
        register_column=4,
        text_rows=_INDIVID_TEXT_ROWS,
        sections={
            5: ("demographics", "Demografiska variabler"),
            6: ("demographics/year-independent", _INDIVID_TEXT_ROWS[6]),
            7: ("demographics/year-independent/country-of-birth", "Födelseland"),
            31: ("demographics/year-independent/deceased", "Avlidna"),
            34: (
                "demographics/year-independent/migration",
                "In- och utvandring",
            ),
            70: ("demographics/household", "För hushåll:"),
            105: ("demographics/family", "För familj:"),
            120: ("education", "Utbildningsvariabler"),
            138: ("employment", "Sysselsättningsvariabler"),
            201: ("employment/occupation", "Yrke 7)"),
            218: ("employment/primary-source", "För största förvärvskälla:"),
            284: ("employment/secondary-source", "För näst största förvärvskälla:"),
            351: ("employment/tertiary-source", "För tredje största förvärvskälla:"),
            405: ("income", "Inkomstvariabler"),
            406: ("income/employment", "Inkomst av förvärvskälla"),
            446: ("income/work-time-social", "Arbetstidsrelaterad social inkomst"),
            447: ("income/work-time-social/study", "Studier:"),
            468: (
                "income/work-time-social/military-service",
                "Värnplikt/Grundläggande militär utbildning:",
            ),
            477: ("income/work-time-social/parental-leave", "Föräldraledighet:"),
            503: (
                "income/work-time-social/health",
                "Sjukdom/Arbetsskada/Rehabilitering:",
            ),
            573: ("income/work-time-social/unemployment", "Arbetslöshet:"),
            587: (
                "income/work-time-social/labour-market-programme",
                "Arbetsmarknadspolitisk åtgärd:",
            ),
            613: (
                "income/work-time-social/sickness-compensation",
                "Sjukersättning (förtidspension)/Aktivitetsersättning (sjukbidrag):",
            ),
            648: ("income/work-time-social/partial-pension", "Delpension:"),
            652: (
                "income/work-time-social/other",
                "Övrig arbetstidrelaterad social inkomst:",
            ),
            658: ("income/other", "Övrig inkomst"),
            659: ("income/other/capital", "Kapitalinkomst"),
            663: ("income/other/miscellaneous", "Annan inkomst"),
            664: ("income/other/pension", "Pension:"),
            681: ("income/other/occupational-pension", "Tjänstepension:"),
            694: ("income/other/pension-other", "Övrig pension:"),
            699: ("income/pension-total", "Summa inkomst av pensioner:"),
            703: (
                "income/occupational-injury-annuity",
                "Yrkes- och arbetsskadelivränta:",
            ),
            709: ("income/other-annuity", "Annan livränta:"),
            712: ("income/survivor-benefit", "Efterlevandeförmån:"),
            719: ("income/family-related", "Familjerelaterade inkomster:"),
            732: ("income/other-benefits", "Övriga inkomster/ersättningar:"),
            749: ("income/disposable", "Disponibel inkomst"),
            766: (
                "income/main-source",
                "Huvudsaklig inkomstkälla, Befolkningens abetsmarknadsstatus",
            ),
            788: (
                "income/activity-statistics",
                "Registerbaserad aktivitetsstatistik, RAKS",
            ),
        },
        context_rows=_INDIVID_CONTEXT_ROWS,
    ),
    "Individ årsoberoende": _TableSpec(
        title="Årgångsoberoende individvariabler, 3 tabeller",
        headers=("Kolumnamn", "Beskrivning", "Register", "Grundregister"),
        table_key="individual-year-independent",
        period_column=None,
        register_column=3,
        text_rows={
            5: "Födelseland",
            30: "Avlidna",
            34: "In- och Utvandring",
        },
        sections={
            5: ("country-of-birth", "Födelseland"),
            30: ("deceased", "Avlidna"),
            34: ("migration", "In- och Utvandring"),
        },
        context_rows={},
    ),
    "Företag": _TableSpec(
        title="Företags- och organisationsvariabler",
        headers=("Kolumnnamn", "Beskrivning", "År", "Register", "Grundregister"),
        table_key="company",
        period_column=3,
        register_column=4,
        text_rows={
            5: "Populationen i LISA:s företagstabell är företag med minst 1 sysselsatt enligt RAMS/BAS.",
            6: "1990-2001: 16 år och äldre, 2002-2011: 16-84 år, 2012-: 16-74 år",
            8: "Karaktäristikor för organisationer",
            48: "Ekonomiska nyckeltal och ekonomisk grunddata finns för företag som ingår i Företagens ekonomi (FEK).",
            49: "FEK täcker näringslivet (exklusive de finansiella och offentliga sektorerna samt hushållens icke-vinstdrivande",
            50: " organisationer).",
            52: "Ekonomiska nyckeltal för företag",
            90: "Ekonomisk grunddata för företag",
            91: "Från 2024 inkluderas godkända resultaträkningar även om balansräkning är underkänd och tvärtom.",
            93: "Resultaträkning",
            111: "Balansräkning",
        },
        sections={
            8: ("organization-characteristics", "Karaktäristikor för organisationer"),
            52: ("economic-key-figures", "Ekonomiska nyckeltal för företag"),
            90: ("economic-fundamentals", "Ekonomisk grunddata för företag"),
            93: ("economic-fundamentals/income-statement", "Resultaträkning"),
            111: ("economic-fundamentals/balance-sheet", "Balansräkning"),
        },
        context_rows={},
        population_rows=(5, 6),
    ),
    "Arbetsställe": _TableSpec(
        title="Arbetsställevariabler",
        headers=(
            "VariabelNamn",
            "Variabel i klartext",
            "Årgång",
            "Register",
            "Grundregister",
        ),
        table_key="workplace",
        period_column=3,
        register_column=4,
        text_rows={
            5: "Populationen i LISA:s företagstabell är företag med minst 1 sysselsatt enligt RAMS/BAS.",
            6: "1990-2001: 16 år och äldre, 2002-2011: 16-84 år, 2012-: 16-74 år",
        },
        sections={},
        context_rows={},
        population_rows=(5, 6),
    ),
}

_PERIOD_PART_RE = re.compile(r"(?P<start>\d{4})(?:-(?P<end>\d{4}))?\Z")


def _cell_text(value: Any) -> str | None:
    if value is None:
        return None
    text = str(value)
    return None if not text.strip() else text


def _period_scope(value: Any, coordinate: str) -> TemporalScope:
    if isinstance(value, bool):
        raise LisaWorkbookError(
            f"unsupported LISA availability period at {coordinate}: {value!r}"
        )
    if isinstance(value, int):
        raw_parts = (str(value),)
    elif isinstance(value, str) and value.strip():
        raw_parts = tuple(value.strip().split())
    else:
        raise LisaWorkbookError(
            f"missing LISA availability period at {coordinate}; only the explicit "
            "year-independent sections may omit it"
        )
    intervals: list[ScopeInterval] = []
    for part in raw_parts:
        match = _PERIOD_PART_RE.fullmatch(part)
        if match is None:
            raise LisaWorkbookError(
                f"unsupported LISA availability period at {coordinate}: {value!r}"
            )
        start = int(match.group("start"))
        end = int(match.group("end") or start)
        if not 1850 <= start <= end <= 2100:
            raise LisaWorkbookError(
                f"invalid LISA availability period at {coordinate}: {value!r}"
            )
        intervals.append(ScopeInterval(start=str(start), end=str(end)))
    try:
        return TemporalScope(kind="intervals", intervals=tuple(intervals))
    except ValueError as exc:
        raise LisaWorkbookError(
            f"overlapping or unordered LISA availability period at {coordinate}: {value!r}"
        ) from exc


def _optional_text_field(value: Any) -> SourceField | None:
    text = _cell_text(value)
    if text is None:
        return None
    return value_field(normalize_text(text), raw=text)


def _sensitivity_field(value: Any, coordinate: str) -> SourceField | None:
    text = _cell_text(value)
    if text is None:
        return None
    normalized = normalize_text(text)
    values: dict[str, bool | str] = {
        "Ja": True,
        "Nej": False,
        "I vissa fall": "conditional",
    }
    if normalized not in values:
        raise LisaWorkbookError(
            f"unsupported LISA sensitivity value at {coordinate}: {value!r}"
        )
    return value_field(values[normalized], raw=text)


def _attach_context(
    records: list[SourceRecord],
    record_index: int,
    *,
    sheet_name: str,
    row_number: int,
    context_row: _ContextRow,
) -> None:
    record = records[record_index]
    cells: list[str] = []
    context: list[str] = []
    for column, value in enumerate(context_row.cells, start=1):
        if value is None:
            continue
        coordinate = f"{sheet_name}!{get_column_letter(column)}{row_number}"
        cells.append(coordinate)
        context.append(f"{context_row.kind}: {value}")
    primary_locator = record.locators[0]
    locator = primary_locator.model_copy(
        update={"physical_cells": (*primary_locator.physical_cells, *cells)}
    )
    updated_context = (*record.context, *context)
    record_id = SourceRecord._record_id(
        source=record.source,
        semantic_record_key=locator.semantic_record_key,
        subject=record.subject,
        edition_scope=record.edition_scope,
        reference_period_scope=record.reference_period_scope,
        fields=record.fields,
        code_set_references=record.code_set_references,
        original_period_text=record.original_period_text,
        context=updated_context,
        delivered_cells=record.delivered_cells,
    )
    records[record_index] = record.model_copy(
        update={
            "record_id": record_id,
            "locators": (locator, *record.locators[1:]),
            "context": updated_context,
        }
    )


def _context_entries(
    sheet_name: str, row_number: int, context_row: _ContextRow
) -> tuple[str, ...]:
    return tuple(
        f"{context_row.kind} {sheet_name}!{get_column_letter(column)}{row_number}: {value}"
        for column, value in enumerate(context_row.cells, start=1)
        if value is not None
    )


def read_lisa_source(path: Path, revision: SourceRevision) -> _LisaSourceRead:
    """Validate the one supported workbook delivery and emit every declaration."""
    try:
        workbook = load_workbook(path, read_only=False, data_only=True)
    except Exception as exc:
        raise LisaWorkbookError(
            f"cannot open selected LISA workbook {path}: {exc}"
        ) from exc
    try:
        actual_sheets = tuple(workbook.sheetnames)
        expected_sheets = tuple(_TABLES)
        if actual_sheets != expected_sheets:
            raise LisaWorkbookError(
                "unsupported LISA workbook sheets: expected "
                f"{expected_sheets!r}, got {actual_sheets!r}"
            )

        records: list[SourceRecord] = []
        worksheet_context: list[str] = []
        semantic_keys: set[tuple[str, ...]] = set()
        for sheet_name, spec in _TABLES.items():
            sheet = workbook[sheet_name]
            if sheet.max_column != len(spec.headers):
                raise LisaWorkbookError(
                    f"unsupported LISA column structure in {sheet_name}: expected "
                    f"{len(spec.headers)} columns, got {sheet.max_column}"
                )
            if _cell_text(sheet["A1"].value) != spec.title:
                raise LisaWorkbookError(
                    f"unsupported LISA table title at {sheet_name}!A1: expected "
                    f"{spec.title!r}, got {sheet['A1'].value!r}"
                )
            actual_header = tuple(
                _cell_text(sheet.cell(3, column).value)
                for column in range(1, len(spec.headers) + 1)
            )
            if actual_header != spec.headers:
                end = get_column_letter(len(spec.headers))
                raise LisaWorkbookError(
                    f"unsupported LISA header at {sheet_name}!A3:{end}3: expected "
                    f"{spec.headers!r}, got {actual_header!r}"
                )
            for row_number, expected in spec.text_rows.items():
                expected_row = (expected,) + (None,) * (len(spec.headers) - 1)
                actual_row = tuple(
                    sheet.cell(row_number, column).value
                    for column in range(1, len(spec.headers) + 1)
                )
                if actual_row != expected_row:
                    end = get_column_letter(len(spec.headers))
                    raise LisaWorkbookError(
                        f"unsupported LISA section structure at {sheet_name}!"
                        f"A{row_number}:{end}{row_number}: expected "
                        f"{expected_row!r}, got {actual_row!r}"
                    )
                worksheet_context.append(
                    f"worksheet-context {sheet_name}!A{row_number}: {actual_row[0]}"
                )
            for row_number, context_row in spec.context_rows.items():
                actual = tuple(
                    sheet.cell(row_number, column).value
                    for column in range(1, len(spec.headers) + 1)
                )
                if actual != context_row.cells:
                    end = get_column_letter(len(spec.headers))
                    raise LisaWorkbookError(
                        f"unsupported LISA context structure at {sheet_name}!"
                        f"A{row_number}:{end}{row_number}: expected "
                        f"{context_row.cells!r}, got {actual!r}"
                    )

            population_text = "\n".join(
                str(sheet.cell(row, 1).value) for row in spec.population_rows
            )
            population = (
                SourceCoordinate(status="value", name=population_text)
                if population_text
                else SourceCoordinate(status="unknown")
            )
            section_key = spec.table_key
            section_label = spec.title
            year_independent = spec.period_column is None
            seen_dated = False
            last_record_index: int | None = None
            for row_number in range(4, sheet.max_row + 1):
                if row_number in spec.sections:
                    section_key, section_label = spec.sections[row_number]
                    if row_number == 6 and sheet_name == "Individ":
                        year_independent = True
                    continue
                if row_number in spec.text_rows:
                    continue

                values = tuple(
                    sheet.cell(row_number, column).value
                    for column in range(1, len(spec.headers) + 1)
                )
                if context_row := spec.context_rows.get(row_number):
                    if context_row.kind == "continuation-note":
                        expected_preceding = context_row.preceding_declaration
                        if last_record_index is None or expected_preceding is None:
                            raise LisaWorkbookError(
                                f"unsupported LISA continuation at {sheet_name}!"
                                f"A{row_number}: expected preceding declaration "
                                f"{expected_preceding!r}"
                            )
                        preceding = records[last_record_index]
                        preceding_column = (
                            preceding.fields.column_name.value
                            if preceding.fields.column_name is not None
                            and preceding.fields.column_name.status == "value"
                            else None
                        )
                        if (
                            preceding.locators[0].physical_table != sheet_name
                            or preceding.locators[0].physical_record
                            != f"row:{expected_preceding[0]}"
                            or preceding_column != expected_preceding[1]
                        ):
                            raise LisaWorkbookError(
                                f"unsupported LISA continuation at {sheet_name}!"
                                f"A{row_number}: expected preceding declaration "
                                f"{expected_preceding!r}"
                            )
                        _attach_context(
                            records,
                            last_record_index,
                            sheet_name=sheet_name,
                            row_number=row_number,
                            context_row=context_row,
                        )
                    else:
                        worksheet_context.extend(
                            _context_entries(sheet_name, row_number, context_row)
                        )
                    continue
                column_text = _cell_text(values[0])
                description = _cell_text(values[1])
                register_text = _cell_text(values[spec.register_column - 1])
                if all(_cell_text(value) is None for value in values):
                    continue
                if column_text is None or description is None or register_text is None:
                    end = get_column_letter(len(spec.headers))
                    raise LisaWorkbookError(
                        f"unsupported LISA row structure at {sheet_name}!"
                        f"A{row_number}:{end}{row_number}: a declaration requires "
                        "identity, description, and register"
                    )
                if register_text.strip() != "LISA":
                    raise LisaWorkbookError(
                        f"unsupported LISA register at {sheet_name}!"
                        f"{get_column_letter(spec.register_column)}{row_number}: "
                        f"expected 'LISA', got {register_text!r}"
                    )

                original_period = (
                    None
                    if spec.period_column is None
                    else values[spec.period_column - 1]
                )
                if spec.period_column is None:
                    edition_scope = TemporalScope(kind="year_independent")
                elif original_period is None:
                    if not year_independent or seen_dated:
                        raise LisaWorkbookError(
                            f"missing LISA availability period at {sheet_name}!"
                            f"{get_column_letter(spec.period_column)}{row_number}"
                        )
                    edition_scope = TemporalScope(kind="year_independent")
                else:
                    if year_independent and not seen_dated:
                        section_key = "demographics/annual"
                        section_label = "Demografiska variabler — årsberoende"
                    year_independent = False
                    seen_dated = True
                    edition_scope = _period_scope(
                        original_period,
                        f"{sheet_name}!{get_column_letter(spec.period_column)}{row_number}",
                    )

                interpreted_column = normalize_token(column_text)
                semantic_key = (spec.table_key, section_key, interpreted_column)
                if semantic_key in semantic_keys:
                    raise LisaWorkbookError(
                        f"duplicate semantic LISA declaration at {sheet_name}!A{row_number}: "
                        f"{semantic_key!r}"
                    )
                semantic_keys.add(semantic_key)
                end = get_column_letter(len(spec.headers))
                locator = RecordLocator(
                    semantic_record_key=semantic_key,
                    physical_file=revision.artifact_path,
                    physical_table=sheet_name,
                    physical_record=f"row:{row_number}",
                    physical_cells=tuple(
                        f"{sheet_name}!{get_column_letter(column)}{row_number}"
                        for column in range(1, len(spec.headers) + 1)
                    ),
                )
                sensitivity = (
                    _sensitivity_field(values[4], f"{sheet_name}!E{row_number}")
                    if sheet_name == "Individ"
                    else None
                )
                base_register = _optional_text_field(values[-1])
                records.append(
                    SourceRecord.create(
                        revision=revision,
                        locators=(locator,),
                        subject=SourceSubject(
                            provider="scb",
                            register=SourceCoordinate(
                                status="value", name=normalize_text(register_text)
                            ),
                            variant=SourceCoordinate(
                                status="value", name=spec.table_key
                            ),
                            population=population,
                            member=SourceCoordinate(
                                status="value", name=interpreted_column
                            ),
                            native=NativeCoordinates(),
                        ),
                        edition_scope=edition_scope,
                        reference_period_scope=TemporalScope(kind="not_applicable"),
                        fields=SourceFields(
                            availability=value_field(True),
                            column_name=value_field(
                                interpreted_column, raw=column_text
                            ),
                            description=value_field(
                                normalize_text(description.strip(), multiline=True),
                                raw=description,
                            ),
                            source_attribution=value_field(
                                normalize_text(register_text), raw=register_text
                            ),
                            sensitivity=sensitivity,
                            base_register=base_register,
                        ),
                        original_period_text=(
                            None if original_period is None else str(original_period)
                        ),
                        context=(spec.title, section_label),
                    )
                )
                last_record_index = len(records) - 1
        return _LisaSourceRead(
            records=tuple(records), worksheet_context=tuple(worksheet_context)
        )
    finally:
        workbook.close()


__all__ = ["LisaWorkbookError", "read_lisa_source"]
