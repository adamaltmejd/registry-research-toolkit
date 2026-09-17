"""SCB event and export-schema declarations before catalog resolution."""

from __future__ import annotations

import io
import re
from typing import TYPE_CHECKING

from openpyxl import load_workbook

from reg_meta_build.db import EXPECTED_HEADERS, _open_scb_csv_prepared
from reg_meta_build.normalization import normalize_text, normalize_token
from reg_meta_build.source_records import (
    DeliveredCell,
    RecordLocator,
    SourceEvidenceRow,
    SourceEvidenceTable,
    SourceField,
    SourceRevision,
    canonical_sha256,
    value_field,
)
from reg_meta_build.source_reference_records import (
    CleanedSourceReferences,
    SourceColumnTypeDeclaration,
    SourceEntityKind,
    SourceEventAction,
    SourceEventDeclaration,
    SourceJoinKeyDeclaration,
)
from reg_meta_build.sources.code_lists import read_selected_bytes
from reg_meta_build.sources.scb_records import _delivered_cells
from reg_meta_build.sources.sos import _row_evidence
from reg_meta_build.sources.sos_records import _delivered_cells as _workbook_cells

if TYPE_CHECKING:
    from collections.abc import Mapping, Sequence
    from pathlib import Path

    from reg_meta_build.input_snapshot import ScbSnapshotReader


_SQL_CREATE_RE = re.compile(
    r"CREATE\s+TABLE\s+\[dbo\]\.\[(\w+)\]\s*\((.*?)\)\s*ON\s+\[PRIMARY\]",
    re.DOTALL | re.IGNORECASE,
)
_SQL_COL_RE = re.compile(
    r"\[(\w+)\]\s+\[(\w+)\](?:\((\d+)\))?\s*(NULL|NOT\s+NULL)?",
)


class ScbReferenceSourceError(ValueError):
    """A selected SCB reference input has an unsupported source contract."""


_ACTIONS: dict[str, SourceEventAction] = {
    "Avslutad": "retired",
    "Tidsseriebrott": "series_break",
    "Ersatt av": "replaced_by",
    "Ersätter": "replaces",
}
_ENTITIES: dict[str, SourceEntityKind] = {
    "Register": "register",
    "RegisterVariant": "variant",
    "RegisterVersion": "edition",
    "Variabel": "variable",
    # The legacy source schema documents AktuellVariabel as the native CVID.
    "AktuellVariabel": "member",
    "Population": "population",
}
_SQL_OUTSIDE = re.compile(
    r"(?:\s+|GO\b|/\*.*?\*/|--[^\r\n]*)*", re.IGNORECASE | re.DOTALL
)
_SQL_COMMENTS = re.compile(r"/\*.*?\*/|--[^\r\n]*", re.DOTALL)
_JOIN_HEADER = ("Tabell", "ID-kolumn", "Beskrivning")


def _text(cell: DeliveredCell, *, multiline: bool = False) -> SourceField:
    if not cell.present or cell.storage_type == "f":
        return SourceField(status="unknown", raw_value=cell.raw_value)
    # Workbook display formatting trims text; original strings retain paragraph
    # indentation. SCB CSV interpretation still supplies its encoding repair.
    text = (
        cell.raw_value
        if cell.raw_type == "str" and cell.storage_type in {"s", "inlineStr"}
        else cell.interpreted_value
    )
    text = normalize_text(text or "", multiline=multiline)
    return (
        value_field(text, raw=cell.raw_value)
        if text
        else SourceField(status="unknown", raw_value=cell.raw_value)
    )


def clean_timeseries_row(
    header: Sequence[str],
    row_number: int,
    cells: Mapping[str, tuple[bool, str | None, str]],
    revision: SourceRevision,
) -> SourceEventDeclaration:
    """Decode event vocabulary without parsing, redirecting or resolving IDs."""
    if list(header) != EXPECTED_HEADERS["Timeseries.csv"] or set(cells) != set(header):
        raise ScbReferenceSourceError("unsupported Timeseries.csv fields")
    delivered = _delivered_cells({name: cells[name] for name in header})
    fields = {cell.name: cell for cell in delivered}
    return SourceEventDeclaration(
        revision=revision,
        locator=RecordLocator(
            semantic_record_key=(
                "event",
                canonical_sha256([cell.model_dump(mode="json") for cell in delivered]),
            ),
            physical_file=revision.artifact_path,
            physical_table="Timeseries.csv",
            physical_record=f"row:{row_number}",
            physical_cells=tuple(f"row:{row_number}:{name}" for name in header),
        ),
        delivered_cells=delivered,
        name=_text(fields["Namn"]),
        action=_ACTIONS.get(
            normalize_token(fields["Handelse"].interpreted_value), "unknown"
        ),
        action_label=_text(fields["Handelse"]),
        description=_text(fields["Beskrivning"], multiline=True),
        entity_kind=_ENTITIES.get(
            normalize_token(fields["Entitet"].interpreted_value), "opaque"
        ),
        entity_label=_text(fields["Entitet"]),
        first_token=fields["ID1"],
        second_token=fields["ID2"],
        document_token=fields["FilID"],
    )


def read_scb_events(
    snapshot: ScbSnapshotReader, revision: SourceRevision
) -> CleanedSourceReferences:
    """Read the accepted event stream with the selected raw-artifact revision."""
    from reg_meta_build.input_snapshot import _decode_cell

    filename = "Timeseries.csv"
    artifact = next(
        (item for item in snapshot.manifest.files if item.name == filename), None
    )
    if (
        artifact is None
        or not artifact.present
        or revision.artifact_sha256 != artifact.raw_sha256
        or revision.artifact_size != artifact.raw_size
    ):
        raise ScbReferenceSourceError(
            "Timeseries.csv differs from its declared revision"
        )
    with _open_scb_csv_prepared(snapshot.root / filename, snapshot) as (header, rows):
        declarations = tuple(
            clean_timeseries_row(header, row, cells, revision) for row, cells in rows
        )
    header_cells = tuple(
        DeliveredCell(
            name=f"column:{index}",
            present=raw is not None,
            raw_value=raw,
            interpreted_value=raw or "",
        )
        for index, raw in enumerate(map(_decode_cell, artifact.header), start=1)
    )
    header_locator = RecordLocator(
        semantic_record_key=("source_header",),
        physical_file=revision.artifact_path,
        physical_table=filename,
        physical_record="row:1",
        physical_cells=tuple(f"row:1:{cell.name}" for cell in header_cells),
    )
    evidence = (
        SourceEvidenceRow(locator=header_locator, role="header", cells=header_cells),
    ) + tuple(
        SourceEvidenceRow(
            locator=item.locator, role="declaration", cells=item.delivered_cells
        )
        for item in declarations
    )
    return CleanedSourceReferences(
        revision=revision,
        declarations=declarations,
        tables=(
            SourceEvidenceTable(
                source=revision.dataset,
                source_revision_id=revision.revision_id,
                name=filename,
                rows=evidence,
            ),
        ),
    )


def _sql_cell(name: str, raw: str | None) -> DeliveredCell:
    return DeliveredCell(
        name=name,
        present=raw is not None,
        raw_value=raw,
        interpreted_value=raw or "",
        raw_type="str" if raw is not None else None,
        storage_type="sql",
    )


def read_scb_column_types(
    path: Path, revision: SourceRevision
) -> CleanedSourceReferences:
    """Decode the delivered CREATE TABLE grammar without executing its SQL."""
    payload = read_selected_bytes(path, revision, error_type=ScbReferenceSourceError)
    try:
        raw = payload.decode("cp1252")
    except UnicodeDecodeError as exc:
        raise ScbReferenceSourceError(
            f"invalid SQL encoding in {revision.artifact_path}: {exc}"
        ) from exc
    # Keep character offsets while excluding SQL-looking text inside comments.
    masked = _SQL_COMMENTS.sub(lambda match: " " * len(match.group()), raw)
    declarations = []
    cursor = 0
    for table in _SQL_CREATE_RE.finditer(masked):
        if not _SQL_OUTSIDE.fullmatch(raw[cursor : table.start()]):
            raise ScbReferenceSourceError(
                f"unsupported SQL before character {table.start()}"
            )
        body = table.group(2)
        body_cursor = 0
        columns = list(_SQL_COL_RE.finditer(body))
        if not columns:
            raise ScbReferenceSourceError(
                f"SQL table {table.group(1)!r} declares no columns"
            )
        for index, column in enumerate(columns):
            gap = body[body_cursor : column.start()].strip()
            if gap != ("" if index == 0 else ","):
                raise ScbReferenceSourceError(
                    f"unsupported SQL column syntax in {table.group(1)!r}"
                )
            offset = table.start(2)
            start, end = offset + column.start(), offset + column.end()
            delivered = tuple(
                _sql_cell(name, value)
                for name, value in zip(
                    ("table_name", "column_name", "data_type", "width", "nullable"),
                    (
                        raw[table.start(1) : table.end(1)],
                        *(
                            raw[
                                offset + column.start(group) : offset
                                + column.end(group)
                            ]
                            if column.start(group) >= 0
                            else None
                            for group in range(1, 5)
                        ),
                    ),
                    strict=True,
                )
            )
            fields = {cell.name: cell for cell in delivered}
            width, null = column.group(3), column.group(4)
            declarations.append(
                SourceColumnTypeDeclaration(
                    revision=revision,
                    locator=RecordLocator(
                        semantic_record_key=(
                            "column_type",
                            table.group(1),
                            column.group(1),
                        ),
                        physical_file=revision.artifact_path,
                        physical_table=table.group(1),
                        physical_record=f"chars:{start}:{end}",
                        physical_cells=(
                            f"chars:{table.start(1)}:{table.end(1)}",
                            *(
                                f"chars:{offset + column.start(group)}:{offset + column.end(group)}"
                                for group in range(1, 5)
                                if column.start(group) >= 0
                            ),
                        ),
                    ),
                    delivered_cells=delivered,
                    table_name=_text(fields["table_name"]),
                    column_name=_text(fields["column_name"]),
                    data_type=_text(fields["data_type"]),
                    declared_width=(
                        value_field(int(width), raw=width)
                        if width is not None
                        else SourceField(status="unknown")
                    ),
                    nullable=(
                        value_field(
                            not null.upper().startswith("NOT"),
                            raw=fields["nullable"].raw_value,
                        )
                        if null is not None
                        else SourceField(status="unknown")
                    ),
                )
            )
            body_cursor = column.end()
        if body[body_cursor:].strip():
            raise ScbReferenceSourceError(
                f"unsupported SQL after columns in {table.group(1)!r}"
            )
        cursor = table.end()
    if not declarations or not _SQL_OUTSIDE.fullmatch(raw[cursor:]):
        raise ScbReferenceSourceError("unsupported SQL outside the declared tables")
    document = _sql_cell("sql_text", raw)
    locator = RecordLocator(
        semantic_record_key=("source_document",),
        physical_file=revision.artifact_path,
        physical_table=path.name,
        physical_record=f"chars:0:{len(raw)}",
        physical_cells=(f"chars:0:{len(raw)}",),
    )
    return CleanedSourceReferences(
        revision=revision,
        declarations=tuple(declarations),
        tables=(
            SourceEvidenceTable(
                source=revision.dataset,
                source_revision_id=revision.revision_id,
                name=path.name,
                rows=(
                    SourceEvidenceRow(
                        locator=locator, role="declaration", cells=(document,)
                    ),
                ),
            ),
        ),
    )


def read_scb_join_keys(path: Path, revision: SourceRevision) -> CleanedSourceReferences:
    """Retain the delivered table/column descriptions without extracting links."""
    payload = read_selected_bytes(path, revision, error_type=ScbReferenceSourceError)
    workbook = load_workbook(io.BytesIO(payload), read_only=False, data_only=False)
    try:
        if len(workbook.sheetnames) != 1:
            raise ScbReferenceSourceError("expected one ID-kolumner worksheet")
        sheet = workbook.active
        if sheet is None or sheet.max_column != 3:
            raise ScbReferenceSourceError("unsupported ID-kolumner column layout")
        header = tuple(sheet[1])
        if tuple(cell.value for cell in header) != _JOIN_HEADER:
            raise ScbReferenceSourceError("unsupported ID-kolumner header")
        rows = []
        declarations = []
        for row_number, cells in enumerate(sheet.iter_rows(), start=1):
            evidence = _row_evidence(sheet, header, cells, {})
            delivered = _workbook_cells(evidence)
            locator = RecordLocator(
                semantic_record_key=(
                    "join_key",
                    canonical_sha256(
                        [cell.model_dump(mode="json") for cell in delivered]
                    ),
                ),
                physical_file=revision.artifact_path,
                physical_table=sheet.title,
                physical_record=f"row:{row_number}",
                physical_cells=tuple(
                    f"{sheet.title}!{cell.coordinate}" for cell in cells
                ),
            )
            rows.append(
                SourceEvidenceRow(
                    locator=locator,
                    role="header" if row_number == 1 else "declaration",
                    cells=delivered,
                )
            )
            if row_number > 1:
                declarations.append(
                    SourceJoinKeyDeclaration(
                        revision=revision,
                        locator=locator,
                        delivered_cells=delivered,
                        table_name=_text(delivered[0]),
                        column_name=_text(delivered[1]),
                        description=_text(delivered[2], multiline=True),
                    )
                )
        return CleanedSourceReferences(
            revision=revision,
            declarations=tuple(declarations),
            tables=(
                SourceEvidenceTable(
                    source=revision.dataset,
                    source_revision_id=revision.revision_id,
                    name=sheet.title,
                    rows=tuple(rows),
                ),
            ),
        )
    finally:
        workbook.close()


__all__ = [
    "ScbReferenceSourceError",
    "clean_timeseries_row",
    "read_scb_column_types",
    "read_scb_events",
    "read_scb_join_keys",
]
