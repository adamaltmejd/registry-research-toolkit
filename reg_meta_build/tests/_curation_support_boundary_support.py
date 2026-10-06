"""Synthetic SCB/SOS sources and curation TOML built through the real pipeline.

Shared by the curation split, support and topology boundary tests. Every source is
a readable fixture (SCB Registerinformation CSV rows from `_csv_fixtures.var_row`,
SOS workbooks from `_sos_fixtures.write_sos_input`) prepared and accepted through
the real input pipeline; assertions read the built SQLite catalog and the report
event ledger.
"""

from __future__ import annotations

import json
import sqlite3
from dataclasses import dataclass, replace
from typing import TYPE_CHECKING, Literal

from _csv_fixtures import (
    REGISTERINFORMATION_HEADER,
    var_row,
    write_input_bundle,
    write_scb_input,
)
from _pipeline_catalog_support import report_events
from _prepared_fixtures import accept_prepared
from _sos_fixtures import DEFAULT_REGISTERS, write_sos_input
from reg_meta_build.pipeline import build_catalog
from reg_meta_build.prepared_catalog import (
    open_prepared_catalog_sources,
    prepare_catalog_sources,
)
from reg_meta_build.source_naming import authored_naming_id

if TYPE_CHECKING:
    from pathlib import Path

SCB_SAMPLE = '[register]\nprovider = "scb"\nslug = "sample"\nnative_id = "1"\n'
SCB_VARIANT = '[[variant]]\nnative_id = "1.10"\nslug = "people"\n'


def sos_id(
    register: str,
    kind: Literal["register", "register_variant", "variable"],
    member: str | None = None,
) -> str:
    """The authored native ID of an SOS register, variant or variable."""
    return authored_naming_id(
        kind, provider="sos", register_key=register, member_key=member
    )


def sos_register(abbrev: str, title: str, variables, subsets=()):
    """A synthetic SOS workbook register derived from the default SYN fixture.

    ``variables`` are field overrides of the SYN ``DIAGNOS`` row (``name``,
    ``deldatamangd``, ``label``, ``description``, ``data_type``); each row covers
    2001-2020. ``subsets`` name the Deldatamängder sheet rows; an empty tuple
    drops that sheet, as some deliveries do.
    """
    base = DEFAULT_REGISTERS[0]
    subset_row = base.deldatamangder[0]
    variable_row = replace(
        base.variables[0], description=None, data_from=2001, data_to=2020
    )
    return replace(
        base,
        abbrev=abbrev,
        title_sv=title,
        kodlistor=(),
        deldatamangder=tuple(
            replace(subset_row, name=name, label=name, data_from=2001, data_to=2020)
            for name in subsets
        ),
        variables=tuple(replace(variable_row, **fields) for fields in variables),
    )


def sos_head(register: str, title: str, subsets=()) -> str:
    """The register TOML head naming an SOS register and its delivered subsets."""
    return (
        f'[register]\nprovider = "sos"\nslug = "{register}"\nname = "{title}"\n'
        f'native_id = "{sos_id(register, "register")}"\n'
        + "".join(
            f'[[variant]]\nnative_id = "{sos_id(register, "register_variant", name)}"\n'
            f'slug = "{name.lower().replace("_", "-")}"\n'
            for name in subsets
        )
    )


@dataclass(frozen=True)
class Built:
    result: dict
    events: list[dict]
    db: Path

    def issues(self, code: str | None = None) -> list[dict]:
        return [
            e
            for e in self.events
            if e["kind"] == "issue" and (code is None or e["code"] == code)
        ]

    def uses(self, records, field: str = "column_name") -> list[tuple]:
        """``(field value, key, use, variable)`` per source occurrence disposition.

        ``records`` are the prepared records whose ledger dispositions are read;
        ``key`` is the last semantic record key part (SCB ``member:<cvid>``).
        """
        by_id = {
            e["record_id"]: e["dispositions"]
            for e in self.events
            if e["kind"] == "source_occurrence"
        }
        return sorted(
            (
                (
                    value.value
                    if (value := getattr(record.fields, field)) is not None
                    else None,
                    record.locators[0].semantic_record_key[-1],
                    disposition["use"],
                    disposition["variable"],
                )
                for record in records
                for disposition in by_id.get(record.record_id, ())
            ),
            key=repr,
        )

    def rows(self, sql: str, *params) -> list[tuple]:
        with sqlite3.connect(self.db) as conn:
            return conn.execute(sql, params).fetchall()

    def columns(self, register: str) -> dict[str, set[str]]:
        """Each built variable slug of ``register`` and its delivery columns."""
        out: dict[str, set[str]] = {}
        for slug, column in self.rows(
            "SELECT v.slug, s.delivery_column_name FROM variable v "
            "JOIN register r USING (register_id) "
            "LEFT JOIN variable_state s USING (variable_id) "
            "WHERE r.slug = ? ORDER BY 1, 2",
            register,
        ):
            out.setdefault(slug, set())
            if column is not None:
                out[slug].add(column)
        return out


@dataclass(frozen=True)
class Sources:
    prepared: Path
    commit: str
    digest: str
    curation: Path

    def build(
        self, tmp_path: Path, label: str = "build", registers: tuple[str, ...] = ()
    ) -> Built:
        report = tmp_path / f"{label}-report"
        output = tmp_path / f"{label}.db"
        result = build_catalog(
            self.prepared,
            self.commit,
            self.digest,
            output,
            report,
            curation_dir=self.curation,
            registers=registers,
            diagnostic=True,
        )
        return Built(result, report_events(report), output)

    def records(self, source: str = "scb-registerinformation") -> tuple:
        """The accepted prepared source records of ``source``, as a curator reads them."""
        prepared = open_prepared_catalog_sources(
            self.prepared, input_commit=self.commit, expected_sha256=self.digest
        )
        return tuple(r for r in prepared.records.records if r.source == source)

    def append(self, relative: str, text: str) -> None:
        path = self.curation / "registers" / relative
        path.write_text(path.read_text(encoding="utf-8") + text, encoding="utf-8")

    def write(self, relative: str, text: str) -> None:
        path = self.curation / "registers" / relative
        path.parent.mkdir(parents=True, exist_ok=True)
        path.write_text(text, encoding="utf-8")


def toml_inline(value) -> str:
    """A JSON-shaped value (``model_dump(mode="json")``) as an inline TOML value."""
    if isinstance(value, dict):
        items = ", ".join(
            f"{key} = {toml_inline(item)}"
            for key, item in value.items()
            if item is not None
        )
        return "{ " + items + " }"
    if isinstance(value, (list, tuple)):
        return "[" + ", ".join(toml_inline(item) for item in value) + "]"
    if isinstance(value, bool):
        return "true" if value else "false"
    if isinstance(value, str):
        return json.dumps(value)
    return str(value)


def summary_rows(rows: list[str]) -> list[str]:
    """One non-sensitive, non-identifier Unika summary row per delivered column.

    The summary carries each variable's flags; without it the build withholds the
    variable as `unresolved_flag`.
    """
    header = REGISTERINFORMATION_HEADER.split("|")
    out = []
    for row in rows:
        cells = dict(zip(header, row.split("|"), strict=True))
        if cells["Kolumnnamn"] in {"", '""'}:
            continue  # A member without a physical column has no column summary.
        year = cells["Registerversion_ForstaGodkannandeDatum"][:4]
        line = "|".join(
            (
                cells["Registernamn"],
                cells["Registerrubrik"],
                cells["Registervariantnamn"],
                cells["Registervariantrubrik"],
                cells["Variabelnamn"],
                cells["Kolumnnamn"],
                year,
                year,
                "0",
                "0",
                "0",
            )
        )
        if line not in out:
            out.append(line)
    # The reader needs a non-empty summary file; an unmatched row is only a warning.
    return out or [
        "TESTREG|Testregistret|Individer|Individer|GenericVar|VALUE|2020|2020|0|0|0"
    ]


def _deliver_blank_join_markers(path: Path) -> None:
    """Add a delivered, blank ``Kopplingsvariabel`` column to the variable sheet.

    SOS reads a delivered blank linkage marker as "not an identifier"; without the
    column the flag is unknown and the build withholds every variable.
    """
    from openpyxl import load_workbook

    workbook = load_workbook(path)
    sheet = workbook["Metadata - Variabelnivå"]
    sheet.cell(row=1, column=sheet.max_column + 1, value="Kopplingsvariabel")
    workbook.save(path)
    workbook.close()


def prepare_sources(
    tmp_path: Path,
    *,
    scb_rows: list[str] | None = None,
    vardemangder_rows: list[str] | None = None,
    sos_registers: tuple = (),
    curation: dict[str, str],
) -> Sources:
    """Prepare SCB rows and SOS workbooks, then write ``curation`` register TOMLs.

    ``curation`` maps a path under ``registers/`` (``scb/sample.toml``) to its text.
    """
    source = tmp_path / "source"
    if not scb_rows:
        # The SCB input set is mandatory; an SOS-only case adds one named SCB variable.
        scb_rows = [var_row(cvid=1001, var_id=101, colname="VALUE")]
        curation = {
            "scb/sample.toml": SCB_SAMPLE
            + SCB_VARIANT
            + '[[variable]]\nnative_id = "1.101"\nslug = "value"\n',
            **curation,
        }
    write_scb_input(
        source,
        registerinformation_rows=scb_rows or [],
        vardemangder_rows=vardemangder_rows or [],
        # Every value item is valid over the whole fixture window.
        valid_dates_rows=sorted(
            {
                f"{row.split('|')[-1]}|2000-01-01|2030-12-31"
                for row in vardemangder_rows or []
            }
        ),
        unika_rows=summary_rows(scb_rows or []),
        include=(
            ("registerinformation", "unika", "vardemangder", "valid_dates")
            if vardemangder_rows
            else ("registerinformation", "unika")
        ),
    )
    if sos_registers:
        for path in write_sos_input(source, registers=sos_registers).glob("*.xlsx"):
            _deliver_blank_join_markers(path)
    bundle = write_input_bundle(tmp_path / "inputs", source)
    prepared = tmp_path / "prepared" / "catalog"
    manifest = prepare_catalog_sources(bundle, prepared)
    commit = accept_prepared(prepared)
    root = tmp_path / "curation"
    (root / "classifications").mkdir(parents=True)
    sources = Sources(prepared, commit, manifest.sha256, root)
    for relative, text in curation.items():
        sources.write(relative, text)
    return sources
