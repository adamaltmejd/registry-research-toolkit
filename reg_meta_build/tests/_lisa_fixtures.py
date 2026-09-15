"""Synthetic workbook matching the four observed LISA delivery layouts."""

from __future__ import annotations

from typing import TYPE_CHECKING

from openpyxl import Workbook
from reg_meta_build.sources.lisa import _TABLES

if TYPE_CHECKING:
    from pathlib import Path

    from openpyxl.worksheet.worksheet import Worksheet


def write_lisa_workbook(path: Path) -> Path:
    workbook = Workbook()
    workbook.remove(workbook.active)
    for sheet_name, spec in _TABLES.items():
        sheet = workbook.create_sheet(sheet_name)
        sheet.cell(1, 1, spec.title)
        for column, header in enumerate(spec.headers, start=1):
            sheet.cell(3, column, header)
        for row, text in spec.text_rows.items():
            sheet.cell(row, 1, text)
        for row, context in spec.context_rows.items():
            _write_row(sheet, row, context.cells)

    individual = workbook["Individ"]
    _write_row(
        individual,
        8,
        ("PersonNr", "Personnummer", None, "LISA", "Nej", "RTB"),
    )
    _write_row(
        individual,
        44,
        (
            "RekrBidr",
            "Rekryteringsbidrag",
            "2003-2008\n2013",
            "LISA",
            "Ja",
            None,
        ),
    )
    _write_row(
        individual,
        322,
        (
            "KU2YrkStalln",
            "Yrkesställning enligt kontrolluppgift",
            "1990-2018",
            "LISA",
            "I vissa fall",
            "RAMS-Jobb",
        ),
    )
    _write_row(
        individual,
        600,
        (
            "AmPolTyp",
            "Typ av arbetsmarknadspolitisk åtgärd",
            "1990-2024",
            "LISA",
            "Nej",
            "Arbetsförmedlingen",
        ),
    )

    independent = workbook["Individ årsoberoende"]
    for row, description in (
        (6, "Personnummer i födelselandstabellen"),
        (31, "Personnummer i tabellen över avlidna"),
        (35, "Personnummer i migrationstabellen"),
    ):
        _write_row(independent, row, ("PersonNr", description, "LISA", "RTB"))

    _write_row(
        workbook["Företag"],
        9,
        ("ForetagKod", "Företagskod", 2024, "LISA", "FDB"),
    )
    _write_row(
        workbook["Arbetsställe"],
        7,
        (
            "Ast_LoneSum",
            "Lönesumma per arbetsställe",
            "1990-2021 2024",
            "LISA",
            "RAMS/BAS",
        ),
    )

    path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(path)
    workbook.close()
    return path


def _write_row(sheet: Worksheet, row: int, values: tuple[object, ...]) -> None:
    for column, value in enumerate(values, start=1):
        sheet.cell(row, column, value)


__all__ = ["write_lisa_workbook"]
