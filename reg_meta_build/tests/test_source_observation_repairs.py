"""Regressions for complete and focused SCB observation traversal."""

from __future__ import annotations

import gzip
import hashlib
import json
from contextlib import contextmanager
from dataclasses import replace
from typing import TYPE_CHECKING

from _csv_fixtures import (
    REGISTERINFORMATION_HEADER,
    _var_row,
    write_input_bundle,
    write_scb_input,
)
from _lisa_fixtures import write_lisa_workbook
from openpyxl import load_workbook
from reg_meta_build.db import _open_scb_csv_raw
from reg_meta_build.input_snapshot import LisaWorkbookSelection, open_input_bundle
from reg_meta_build.source_inspection import (
    CensusCompletion,
    inspect_bundle_source_records,
    write_scb_observation_census,
)
from reg_meta_build.source_records import SourceRevision
from reg_meta_build.sources import lisa as lisa_module
from reg_meta_build.sources.lisa import read_lisa_source

from reg_meta_build import db as db_module

if TYPE_CHECKING:
    from collections.abc import Iterator
    from pathlib import Path
    from typing import Any

    import pytest


def _read_census(path: Path) -> list[dict[str, Any]]:
    with gzip.open(path, "rt", encoding="utf-8") as source:
        return [json.loads(line) for line in source]


def _census(tmp_path: Path, name: str, rows: list[str]) -> list[dict[str, Any]]:
    input_dir = tmp_path / name / "source"
    write_scb_input(input_dir, registerinformation_rows=rows)
    selection = write_input_bundle(tmp_path / name / "accepted", input_dir)
    destination = tmp_path / f"{name}.jsonl.gz"
    write_scb_observation_census(
        open_input_bundle(selection), destination, code_commit="c" * 40
    )
    return _read_census(destination)


def _observation(lines: list[dict[str, Any]], member_id: int) -> dict[str, Any]:
    return next(
        line
        for line in lines
        if line["type"] == "observation"
        and line["record"]["subject"]["native"]["member_id"] == member_id
    )


def _workbook_revision(path: Path) -> SourceRevision:
    content = path.read_bytes()
    return SourceRevision.create(
        dataset="scb-lisa-variable-list-workbook",
        publisher="SCB",
        purpose="annotated-row relocation test",
        upstream_revision="2024-2025",
        artifact_path=path.name,
        artifact_size=len(content),
        artifact_sha256=hashlib.sha256(content).hexdigest(),
    )


def test_raw_scb_traversal_reuses_source_row_without_presence_sidecar(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    source_row: list[str | None] = ["kept", ""]

    @contextmanager
    def open_source(
        _path: Path, _snapshot: object
    ) -> Iterator[tuple[list[str], Iterator[list[str | None]]]]:
        yield ["first", "second"], iter((source_row,))

    monkeypatch.setattr(db_module, "_open_scb_source_raw", open_source)

    with _open_scb_csv_raw(tmp_path / "Other.csv") as (_header, rows):
        row_number, fields = next(rows)

    assert row_number == 2
    assert fields is source_row
    assert type(fields) is list
    assert fields == ["kept", ""]


def test_focused_lisa_inspection_skips_unrelated_malformed_native_row(
    tmp_path: Path,
) -> None:
    unrelated = _var_row(colname="Other", cvid=1, var_id=1)
    unrelated_fields = unrelated.split("|")
    header = REGISTERINFORMATION_HEADER.split("|")
    unrelated_fields[header.index("RegVerID")] = "not-an-id"
    lisa = _var_row(
        colname="AmPolTyp",
        cvid=2,
        var_id=31619,
        year="2020",
        regver_id=200,
        register=("LISA", 34, 153),
    )
    input_dir = tmp_path / "source"
    write_scb_input(
        input_dir,
        registerinformation_rows=["|".join(unrelated_fields), lisa],
    )
    workbook = write_lisa_workbook(input_dir / "docs" / "lisa.xlsx")
    selection = write_input_bundle(
        tmp_path / "accepted",
        input_dir,
        lisa_workbook=LisaWorkbookSelection(
            path=workbook,
            upstream_revision="2024-2025",
            sha256=hashlib.sha256(workbook.read_bytes()).hexdigest(),
        ),
    )

    report = inspect_bundle_source_records(
        open_input_bundle(selection), code_commit="c" * 40
    )

    scb_records = [
        record
        for record in report.source_records
        if record.source == "scb-registerinformation"
    ]
    assert len(scb_records) == 1
    assert scb_records[0].subject.native.register_id == 34
    assert scb_records[0].locators[0].physical_record == "row:3"
    assert report.summary.scb_total_occurrences == 1


def test_scb_identity_survives_reorder_revision_and_unrelated_edit(
    tmp_path: Path,
) -> None:
    target = _var_row(
        colname="Signal",
        cvid=2181,
        var_id=1880,
        vardef="Selected definition",
    )
    unrelated = _var_row(colname="Other", cvid=9001, var_id=9000)
    unrelated_edit = _var_row(
        colname="Other",
        cvid=9001,
        var_id=9000,
        vardef="An unrelated edit",
    )

    original = _census(tmp_path, "identity-original", [target, unrelated])
    reordered = _census(tmp_path, "identity-reordered", [unrelated_edit, target])
    changed_payload = _census(
        tmp_path,
        "identity-payload",
        [
            unrelated_edit,
            _var_row(
                colname="Signal",
                cvid=2181,
                var_id=1880,
                vardef="Changed definition",
            ),
        ],
    )
    changed_context = _census(
        tmp_path,
        "identity-context",
        [
            unrelated_edit,
            _var_row(
                colname="Signal",
                cvid=2181,
                var_id=1880,
                vardef="Selected definition",
                population_name="Another population",
            ),
        ],
    )

    original_record = _observation(original, 2181)["record"]
    reordered_record = _observation(reordered, 2181)["record"]
    assert original_record["record_id"] == reordered_record["record_id"]
    assert (
        original_record["source_revision_id"] != reordered_record["source_revision_id"]
    )
    assert original_record["locators"][0]["physical_record"] == "row:2"
    assert reordered_record["locators"][0]["physical_record"] == "row:3"

    original_completion = original[-1]
    reordered_completion = reordered[-1]
    for record, completion in (
        (original_record, original_completion),
        (reordered_record, reordered_completion),
    ):
        assert record["source_revision_id"] == completion["pins"]["source_revision_id"]
        assert (
            completion["pins"]["source_revision_id"]
            == completion["source_revision"]["revision_id"]
        )
    for pin in (
        "bundle_manifest_sha256",
        "scb_snapshot_manifest_sha256",
        "source_revision_id",
    ):
        assert original_completion["pins"][pin] != reordered_completion["pins"][pin]
    assert (
        original_completion["source_revision"]["artifact_sha256"]
        != reordered_completion["source_revision"]["artifact_sha256"]
    )

    assert (
        _observation(changed_payload, 2181)["record"]["record_id"]
        != original_record["record_id"]
    )
    assert (
        _observation(changed_context, 2181)["record"]["record_id"]
        != original_record["record_id"]
    )


def test_lisa_annotation_identity_survives_physical_row_relocation(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    original_path = write_lisa_workbook(tmp_path / "original.xlsx")
    relocated_path = write_lisa_workbook(tmp_path / "relocated.xlsx")
    original_read = read_lisa_source(original_path, _workbook_revision(original_path))

    workbook = load_workbook(relocated_path)
    sheet = workbook["Individ"]
    for column in range(1, 7):
        sheet.cell(324, column).value = sheet.cell(323, column).value
        sheet.cell(323, column).value = None
    workbook.save(relocated_path)
    workbook.close()

    original_spec = lisa_module._TABLES["Individ"]
    context_rows = dict(original_spec.context_rows)
    context_rows[324] = context_rows.pop(323)
    monkeypatch.setitem(
        lisa_module._TABLES,
        "Individ",
        replace(original_spec, context_rows=context_rows),
    )
    relocated_read = read_lisa_source(
        relocated_path, _workbook_revision(relocated_path)
    )

    def ku2(records: tuple[Any, ...]) -> Any:
        return next(
            record
            for record in records
            if record.fields.column_name is not None
            and record.fields.column_name.value == "KU2YrkStalln"
        )

    original = ku2(original_read.records)
    relocated = ku2(relocated_read.records)
    assert original.source_revision_id != relocated.source_revision_id
    assert original.record_id == relocated.record_id
    assert original.context == relocated.context
    assert original.locators[0].semantic_record_key == (
        relocated.locators[0].semantic_record_key
    )
    assert original.locators[0].physical_cells[-2:] == (
        "Individ!B323",
        "Individ!F323",
    )
    assert relocated.locators[0].physical_cells[-2:] == (
        "Individ!B324",
        "Individ!F324",
    )


def test_context_group_keeps_members_when_one_context_has_own_alternatives(
    tmp_path: Path,
) -> None:
    def row(data_length: str, population_name: str) -> str:
        return _var_row(
            colname="Contextual",
            cvid=7001,
            var_id=7000,
            year="2004",
            regver_id=700,
            data_type="char",
            data_length=data_length,
            population_name=population_name,
        )

    rows = [
        row("10", "Population A"),
        row("11", "Population A"),
        row("10", "Population B"),
    ]

    lines = _census(tmp_path, "contexts", rows)
    same_context = [
        line for line in lines if line["type"] == "same_complete_context_alternatives"
    ]
    separated = [
        line for line in lines if line["type"] == "context_separated_alternatives"
    ]

    assert len(same_context) == 1
    assert len(separated) == 1
    assert {
        member["locator"]["physical_record"]
        for alternative in separated[0]["alternatives"]
        for member in alternative["members"]
    } == {"row:2", "row:3", "row:4"}
    assert (
        len(
            {
                member["context_fingerprint"]
                for alternative in separated[0]["alternatives"]
                for member in alternative["members"]
            }
        )
        == 2
    )
    completion = CensusCompletion.model_validate_json(
        json.dumps(lines[-1], ensure_ascii=False)
    )
    assert completion.counts.context_separated_groups == 1
    assert completion.affected_memberships.context_separated_cvids == (7001,)


def test_temporal_group_crosses_edition_specific_cvids_without_expanding_pool(
    tmp_path: Path,
) -> None:
    rows = [
        _var_row(
            colname="Temporal",
            cvid=8101,
            var_id=8100,
            year="2001",
            versionname="2001-2003",
            regver_id=801,
            data_type="char",
            data_length="10",
        ),
        _var_row(
            colname="Temporal",
            cvid=8102,
            var_id=8100,
            year="2004",
            regver_id=802,
            data_type="char",
            data_length="11",
        ),
    ]

    lines = _census(tmp_path, "temporal", rows)
    groups = [
        line for line in lines if line["type"] == "unproved_temporal_alternatives"
    ]

    assert len(groups) == 1
    assert not any(part.startswith("member:") for part in groups[0]["comparison_key"])
    members = [
        member
        for alternative in groups[0]["alternatives"]
        for member in alternative["members"]
    ]
    assert {member["native"]["member_id"] for member in members} == {8101, 8102}
    assert {member["native"]["edition_id"] for member in members} == {801, 802}
    assert len(members) == 2
    pooled = next(
        line
        for line in lines
        if line["type"] == "observation"
        and line["record"]["subject"]["native"]["member_id"] == 8101
    )
    assert pooled["record"]["edition_scope"]["kind"] == "pooled"
    assert pooled["interpretation_issue"]["kind"] == "pooled_period"
    completion = CensusCompletion.model_validate_json(
        json.dumps(lines[-1], ensure_ascii=False)
    )
    assert completion.counts.unproved_temporal_groups == 1
    assert completion.counts.unproved_temporal_cvids == 2
    assert completion.affected_memberships.unproved_temporal_cvids == (8101, 8102)
