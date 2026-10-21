"""Materialize and prepare a `cases/build/` source spec.

A spec's SCB rows (`_csv_fixtures`) and SoS workbooks (`_sos_fixtures`) are written,
prepared and accepted through the real input pipeline.

Prepared inputs are cached by the content hash of their source spec, so cases that
share a source set prepare it once per session. An entry is written to a private
staging directory and renamed into place: concurrent xdist workers never write the
same path, and a failed preparation never leaves a partial entry. Entries
persist across sessions and worktrees in a generation of the shared fixture cache
(the reader cache's location), keyed by everything that shapes a prepared input.
"""

from __future__ import annotations

import functools
import hashlib
import json
import re
import shutil
import subprocess
import tempfile
import zipfile
from dataclasses import dataclass
from datetime import datetime
from pathlib import Path
from typing import TYPE_CHECKING

import pytest
from _csv_fixtures import summary_rows, write_scb_input
from _pipeline_catalog_support import (
    prepare_accepted,
    write_authored,
    write_delivered,
)
from _registerinformation_rows import replace_registerinformation_cell, var_row
from _sos_fixtures import (
    SosRegister,
    SosSubset,
    SosVariable,
    write_sos_input,
)
from openpyxl import load_workbook
from reader_artifacts import build_inputs_digest, generation_dir
from reg_meta_build.prepared_catalog import (
    ReferenceEvidence,
    open_prepared_catalog_sources,
)
from reg_meta_build.source_evidence import canonical_sha256, evidence_sha256
from reg_meta_build.source_value_bindings import (
    ValueBindingResult,
    bind_code_lists,
    open_value_bindings,
)

if TYPE_CHECKING:
    from reg_meta_build.source_records import SourceRecord


SOURCES = Path(__file__).with_name("cases") / "build" / "_sources"


def source_spec(reference: str | dict) -> dict:
    """A source spec: inline, or the name of a shared `_sources/<name>.json`."""
    if isinstance(reference, dict):
        return reference
    return json.loads((SOURCES / f"{reference}.json").read_text(encoding="utf-8"))


def scb_row(row: dict) -> str:
    """One Registerinformation row: `var_row` arguments, then raw `cells` overrides."""
    text = var_row(
        **{
            key: tuple(value) if key == "register" else value
            for key, value in row.items()
            if key != "cells"
        }
    )
    for name, value in row.get("cells", {}).items():
        text = replace_registerinformation_cell(text, name, value)
    return text


def write_sources(spec: dict, source: Path) -> None:
    """Materialize a source spec as provider deliveries under ``source``."""
    scb = spec["scb"]
    rows = [scb_row(row) for row in scb["registerinformation"]]
    values = scb.get("vardemangder", [])
    # `"unika": null` delivers no Unika summary file at all.
    unika = scb.get("unika", summary_rows(rows))
    write_scb_input(
        source,
        registerinformation_rows=rows,
        vardemangder_rows=values,
        unika_rows=unika or [],
        # Default: every value item is valid over the whole fixture window.
        valid_dates_rows=scb.get(
            "valid_dates",
            sorted({f"{row.split('|')[-1]}|2000-01-01|2030-12-31" for row in values}),
        ),
        include=("registerinformation",)
        + (("unika",) if unika is not None else ())
        + (("vardemangder", "valid_dates") if values else ()),
    )
    write_delivered(spec, source)
    for register in spec.get("sos", ()):
        workbook_spec = SosRegister(
            abbrev=register["abbrev"],
            title_sv=register["title"],
            description_sv=register.get("description"),
            variables=tuple(
                SosVariable(**{k: v for k, v in row.items() if k != "linkage"})
                for row in register["variables"]
            ),
            deldatamangder=tuple(
                SosSubset(**row) for row in register.get("subsets", ())
            ),
        )
        sos_dir = write_sos_input(source, registers=(workbook_spec,))
        path = (
            sos_dir / f"Metadata {register['title']} ({register['abbrev']})_webb.xlsx"
        )
        # A delivered blank `Kopplingsvariabel` column is SOS's explicit "not an
        # identifier" claim; without it every variable's flag is unknown and withheld.
        workbook = load_workbook(path)
        sheet = workbook["Metadata - Variabelnivå"]
        linkage = sheet.max_column + 1
        sheet.cell(row=1, column=linkage, value="Kopplingsvariabel")
        # A variable's `linkage` text fills its cell: a declared linkage variable.
        for row, variable in enumerate(register["variables"], start=2):
            if variable.get("linkage"):
                sheet.cell(row=row, column=linkage, value=variable["linkage"])
        # A delivered Kodlista names its variable in a `Variabelnamn` preamble row;
        # without it the build cannot bind the list (`unresolved_list_reference`).
        code_lists = {
            f"Kodlista_{name}": [
                ["Variabelnamn", name],
                ["Tidsperiod", "Kod", "Beskrivning"],
                *rows,
            ]
            for name, rows in register.get("code_lists", {}).items()
        }
        for name, sheet_rows in {**code_lists, **register.get("sheets", {})}.items():
            extra = workbook.create_sheet(name)
            for row in sheet_rows:
                extra.append(row)
        if register.get("blank_dataset"):
            # The `Datamängd` cell names the register; blank, the workbook has none.
            workbook["Generell information"]["C4"] = None
        workbook.properties.created = workbook.properties.modified = _XLSX_EPOCH
        workbook.save(path)
        workbook.close()
        _fix_zip_times(path)


# openpyxl stamps the save time into docProps/core.xml and every zip entry, so two
# prepares of one spec would deliver different workbook bytes (and revisions), and a
# drift case could go stale on the timestamp alone.
_XLSX_EPOCH = datetime(2000, 1, 1)  # openpyxl writes naive times as UTC


def _fix_zip_times(path: Path) -> None:
    with zipfile.ZipFile(path) as archive:
        entries = [(info, archive.read(info)) for info in archive.infolist()]
    with zipfile.ZipFile(path, "w") as archive:
        for info, data in entries:
            if info.filename == "docProps/core.xml":
                # `save` restamps `modified` with the current time.
                data = re.sub(
                    rb"(<dcterms:modified[^>]*>)[^<]*",
                    rb"\g<1>2000-01-01T00:00:00Z",
                    data,
                )
            fixed = zipfile.ZipInfo(info.filename, date_time=(1980, 1, 1, 0, 0, 0))
            fixed.compress_type = info.compress_type
            archive.writestr(fixed, data)


@dataclass(frozen=True)
class PreparedSet:
    """One accepted prepared input: what `build_catalog` and authoring read."""

    prepared: Path
    commit: str
    digest: str

    def records(self) -> tuple[SourceRecord, ...]:
        return _records(self.prepared, self.commit, self.digest)

    def bindings(self, record: SourceRecord) -> ValueBindingResult:
        """The source lists the prepared value store binds to one record."""
        with open_value_bindings(self.opened().value_sources) as sessions:
            return bind_code_lists(record, sessions)

    def claims(self, records: tuple[SourceRecord, ...]) -> list:
        """The code-list claims the prepared value store binds to ``records``."""
        return [claim for record in records for claim in self.bindings(record).claims]

    def coding_sha256(self, records: tuple[SourceRecord, ...]) -> list[str]:
        """The bound physical code-list evidence of ``records``, as curators pin it."""
        return [evidence_sha256(claim) for claim in self.claims(records)]

    def opened(self):
        return _opened(self.prepared, self.commit, self.digest)

    def relationship(self, table: str):
        """The one literal relationship declaration delivered in ``table``."""
        (declaration,) = (
            evidence.declaration
            for evidence in self.opened().iter_evidence()
            if isinstance(evidence, ReferenceEvidence)
            and evidence.declaration.locator.physical_table == table
        )
        return declaration

    def table_sha256(self, table: str) -> str:
        """The content hash of the one prepared evidence table named ``table``."""
        (found,) = (t for t in self.opened().records.iter_tables() if t.name == table)
        return evidence_sha256(found)


# Authoring reads open a cache entry's accepted sources once per process: an entry
# is immutable once renamed into place, and the build under test still opens and
# checks it itself.
@functools.cache
def _opened(prepared: Path, commit: str, digest: str):
    return open_prepared_catalog_sources(
        prepared, expected_sha256=digest, input_commit=commit
    )


@functools.cache
def _records(prepared: Path, commit: str, digest: str) -> tuple[SourceRecord, ...]:
    return tuple(_opened(prepared, commit, digest).records.records)


class PreparedCache:
    """Accepted prepared inputs keyed by the content hash of their source spec."""

    def __init__(self, root: Path) -> None:
        self.root = root
        root.mkdir(parents=True, exist_ok=True)

    def get(self, spec: dict) -> PreparedSet:
        key = hashlib.sha256(
            json.dumps(spec, sort_keys=True, ensure_ascii=False).encode()
        ).hexdigest()
        entry = self.root / key
        if not (entry / "selection.json").is_file():
            staging = Path(tempfile.mkdtemp(prefix=".staging-", dir=self.root))
            try:
                write_sources(spec, staging / "source")
                write_authored(spec, staging / "authored")
                _, commit, digest = prepare_accepted(
                    staging, staging / "source", staging / "authored"
                )
                # Only the accepted prepared repository is read after preparation.
                shutil.rmtree(staging / "source")
                shutil.rmtree(staging / "authored")
                shutil.rmtree(staging / "inputs")
                (staging / "selection.json").write_text(
                    json.dumps({"commit": commit, "digest": digest}), encoding="utf-8"
                )
                try:
                    staging.rename(entry)
                except OSError:
                    if not (entry / "selection.json").is_file():
                        raise
            finally:
                shutil.rmtree(staging, ignore_errors=True)
        selection = json.loads((entry / "selection.json").read_text(encoding="utf-8"))
        return PreparedSet(
            entry / "prepared" / "catalog", selection["commit"], selection["digest"]
        )


@pytest.fixture(scope="session")
def prepared_cache() -> PreparedCache:
    """The prepared inputs shared by the build cases and `catalog`.

    They live in a generation of the shared fixture cache, under the reader
    cache's one location policy (`reader_artifacts.fixture_cache_dir`):
    `$REG_FIXTURE_CACHE`, else the system temp directory. Entries are reused across
    sessions, worktrees and xdist workers; a fresh CI runner starts cold.
    """
    return PreparedCache(fixture_generation() / "build-case-prepared")


@functools.cache
def fixture_generation() -> Path:
    """This checkout's generation of the shared fixture cache.

    Keyed by `_prepared_inputs_digest`, so any cache under it (the prepared inputs,
    the `cases/cli/` artifact) is invalidated by a builder or fixture-module edit.
    """
    return generation_dir(_prepared_inputs_digest())


# The test modules that write and accept a source spec's provider deliveries; an
# edit to any of them changes what a prepared entry holds.
_PREPARING_MODULES = (
    "_build_case_sources.py",
    "_csv_fixtures.py",
    "_registerinformation_rows.py",
    "_input_repo_fixtures.py",
    "_sos_fixtures.py",
    "_prepared_fixtures.py",
    "_pipeline_catalog_support.py",
)


def _prepared_inputs_digest() -> str:
    """Everything a prepared entry depends on besides its source spec.

    The reader cache's build inputs (the `reg_meta_build` sources, the native
    extension sources, the installed
    distributions, Python and SQLite), the modules that prepare a spec, and the Git
    that commits and checks the accepted repository.
    """
    here = Path(__file__).parent
    return canonical_sha256(
        {
            "build_inputs": build_inputs_digest(),
            "modules": {
                name: hashlib.sha256((here / name).read_bytes()).hexdigest()
                for name in _PREPARING_MODULES
            },
            "git": subprocess.run(
                ["git", "--version"], capture_output=True, text=True, check=True
            ).stdout.strip(),
        }
    )
