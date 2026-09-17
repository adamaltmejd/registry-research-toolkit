"""Lossless source-record boundary for Socialstyrelsen workbooks."""

from __future__ import annotations

import hashlib
from dataclasses import replace
from datetime import datetime
from typing import TYPE_CHECKING

import pytest
from reg_meta_build.source_records import (
    DeliveredCell,
    NativeCoordinates,
    SourceCoordinate,
    SourceRevision,
    value_field,
)
from reg_meta_build.source_reference_records import (
    SourceCodeCrosswalkDeclaration,
    SourceDerivationDeclaration,
)
from reg_meta_build.sources.sos import SosParseError, SosParseIssue, parse_register_file
from reg_meta_build.sources.sos_records import (
    clean_sos_source,
    clean_sos_variable,
    iter_sos_variable_records,
)

if TYPE_CHECKING:
    from pathlib import Path


_CLASSIFICATION_URL = "https://example.test/classifications/ssyk"


@pytest.mark.parametrize(
    ("start", "end", "expected"),
    [
        # Excel calendar cells have no timezone.
        (datetime(2014, 2, 3), datetime(2015, 4, 5), ("2014-02-03", "2015-04-05")),  # noqa: DTZ001
        ("201402", "201503", ("2014-02-01", "2015-03-31")),
        ("20140203", "20150405", ("2014-02-03", "2015-04-05")),
        ("2014-02-03", "2015-04-05", ("2014-02-03", "2015-04-05")),
        ("2020-02-03", "2020", ("2020-02-03", "2020-12-31")),
        ("202002", "2020", ("2020-02-01", "2020-12-31")),
        ("2020", "20200203", ("2020-01-01", "2020-02-03")),
        ("2020", "2020", ("2020", "2020")),
        ("20140203", None, None),
        ("20140230", "20150405", None),
        ("20150405", "20140203", None),
        ("=2001", "2020", None),
    ],
)
def test_variable_coverage_preserves_explicit_date_bounds(
    tmp_path: Path, start: object, end: object, expected: tuple[str, str] | None
) -> None:
    import openpyxl

    path = tmp_path / "Metadata Patientregistret (PAR)_webb.xlsx"
    _write_source_workbook(path)
    workbook = openpyxl.load_workbook(path)
    sheet = workbook["Metadata - Variabelnivå"]
    sheet["H2"], sheet["I2"] = start, end
    workbook.save(path)
    register = parse_register_file(path)
    record = clean_sos_variable(register, register.variables[0], _revision(path))
    if expected is None:
        assert record.edition_scope.kind == "unknown"
    else:
        assert record.edition_scope.kind == "intervals"
        assert tuple((i.start, i.end) for i in record.edition_scope.intervals) == (
            expected,
        )
    assert record.fields.classification_declared == value_field(_CLASSIFICATION_URL)
    assert next(
        cell for cell in record.delivered_cells if cell.name == "Data från"
    ).raw_value == str(start)
    assert next(
        cell for cell in record.delivered_cells if cell.name == "Data till"
    ).raw_value == (str(end) if end is not None else "")


@pytest.mark.parametrize(
    ("header", "field"),
    [
        ("Variabelnamn", "name"),
        (" variabelNAMN ", "name"),
        ("Datavynamn", "deldatamangd"),
        ("Data från", "data_from"),
        ("Datatyp", "data_type"),
    ],
)
def test_variable_headers_reject_duplicate_semantic_fields(
    tmp_path: Path, header: str, field: str
) -> None:
    import openpyxl

    path = tmp_path / "source.xlsx"
    _write_source_workbook(path)
    workbook = openpyxl.load_workbook(path)
    sheet = workbook["Metadata - Variabelnivå"]
    sheet["L1"], sheet["L2"] = header, "SECOND_VALUE"
    workbook.save(path)
    original = path.read_bytes()

    with pytest.raises(SosParseError, match=f"ambiguous headers for '{field}'"):
        parse_register_file(path)
    assert path.read_bytes() == original


def test_absent_formula_cache_does_not_change_existing_source_payloads() -> None:
    original = {
        "name": "Kod",
        "present": True,
        "raw_value": "001",
        "interpreted_value": "001",
        "raw_type": "str",
        "storage_type": "s",
        "number_format": "General",
        "hyperlink_target": None,
        "hyperlink_location": None,
    }
    assert DeliveredCell.model_validate(original).model_dump(mode="json") == original
    for raw_value, raw_type in [("", "str"), ("", "none"), (" ", "str")]:
        with_cache = {
            **original,
            "cached_raw_value": raw_value,
            "cached_raw_type": raw_type,
        }
        assert (
            DeliveredCell.model_validate(with_cache).model_dump(mode="json")
            == with_cache
        )


def _write_source_workbook(path: Path) -> None:
    import openpyxl
    from openpyxl.worksheet.hyperlink import Hyperlink

    workbook = openpyxl.Workbook()
    general = workbook.active
    general.title = "Generell information"
    general.append(["", "Om datamängden version", None])
    general.append(["", "Datamängd", "Patientregistret källa"])
    general.append(["", "Version", "2026:1"])

    variables = workbook.create_sheet("Metadata - Variabelnivå")
    variables.append(
        [
            "Deldatamängdsnamn",
            "Variabelnamn",
            "Variabeletikett",
            "Variabelbeskrivning",
            "Värdemängd",
            "Länk kodverk",
            "Datatyp",
            "Data från",
            "Data till",
            "Specificera källa",
            "Eget källfält",
        ]
    )
    base = [
        "PAR_OV",
        "HDIA",
        " Huvuddiagnos ",
        "Första raden\r\nandra raden  ",
        "Se kodlista",
        _CLASSIFICATION_URL,
        "Heltal",
        2001,
        2020,
        "Patientregistret",
        "bevaras",
    ]
    variables.append(base)
    variables.append(base)
    variables.append([*base[:2], "Annan etikett", *base[3:]])
    variables.append([*base[:7], "2001", 2020, *base[9:]])
    for row_number in range(2, 6):
        variables[f"F{row_number}"].hyperlink = _CLASSIFICATION_URL
    variables.append(
        [
            "PAR_OV",
            "PARTIELL",
            "Partiell",
            None,
            None,
            None,
            "Sträng (text)",
            2010,
            None,
            None,
            None,
        ]
    )
    variables["F6"] = "Visad kodlista"
    variables["F6"].hyperlink = Hyperlink(
        ref="F6",
        location="'Kodlista_HDIA'!A1",
        display="Visad kodlista",
    )
    variables.append(
        [
            "PAR_OV",
            "MALFORMED",
            "Malformed",
            None,
            None,
            None,
            "Heltal",
            "+2001",
            "2020",
            None,
            None,
        ]
    )
    variables.append(
        [
            None,
            "UTAN_DEL",
            "Utan deldatamängd",
            None,
            None,
            None,
            "Datum",
            None,
            None,
            None,
            None,
        ]
    )

    codes = workbook.create_sheet("Kodlista_HDIA")
    codes.append(["Tidsperiod", "Kod", "Beskrivning"])
    codes.append(["2001-2020", 1, "Kod ett"])
    codes["B2"].number_format = "000"
    workbook.save(path)


def _revision(path: Path) -> SourceRevision:
    payload = path.read_bytes()
    return SourceRevision.create(
        dataset="sos-metadata",
        publisher="Socialstyrelsen",
        purpose="Socialstyrelsen source-record fixture",
        upstream_revision="2026:1",
        artifact_path=path.name,
        artifact_size=len(payload),
        artifact_sha256=hashlib.sha256(payload).hexdigest(),
    )


def test_parser_retains_xlsx_row_cells_and_code_display_provenance(
    tmp_path: Path,
) -> None:
    path = tmp_path / "Metadata Patientregistret (PAR)_webb.xlsx"
    _write_source_workbook(path)

    register = parse_register_file(path)

    evidence = register.variables[0].source_evidence
    assert evidence is not None
    assert evidence.sheet_name == "Metadata - Variabelnivå"
    assert evidence.row_number == 2
    coverage_start = next(
        cell for cell in evidence.cells if cell.field_name == "data_from"
    )
    assert (
        coverage_start.coordinate,
        coverage_start.raw_value,
        coverage_start.data_type,
        coverage_start.number_format,
        coverage_start.display_value,
    ) == ("H2", 2001, "n", "General", "2001")
    assert next(
        cell for cell in evidence.cells if cell.header == "Eget källfält"
    ).raw_value == ("bevaras")
    classification = next(
        cell for cell in evidence.cells if cell.field_name == "external_classification"
    )
    assert (
        classification.coordinate,
        classification.raw_value,
        classification.data_type,
        classification.number_format,
        classification.hyperlink_target,
    ) == ("F2", _CLASSIFICATION_URL, "s", "General", _CLASSIFICATION_URL)

    linked_evidence = register.variables[4].source_evidence
    assert linked_evidence is not None
    internal_link = next(
        cell
        for cell in linked_evidence.cells
        if cell.field_name == "external_classification"
    )
    assert (
        internal_link.raw_value,
        internal_link.hyperlink_target,
        internal_link.hyperlink_location,
    ) == ("Visad kodlista", None, "'Kodlista_HDIA'!A1")

    code = register.kodlistor[0].rows[0]
    assert code.kod == "001"
    assert code.source_evidence is not None
    code_cell = next(
        cell for cell in code.source_evidence.cells if cell.field_name == "kod"
    )
    assert (
        code_cell.coordinate,
        code_cell.raw_value,
        code_cell.data_type,
        code_cell.number_format,
        code_cell.display_value,
    ) == ("B2", 1, "n", "000", "001")


def test_source_records_keep_occurrences_conflicts_and_native_sos_coordinates(
    tmp_path: Path,
) -> None:
    path = tmp_path / "Metadata Patientregistret (PAR)_webb.xlsx"
    _write_source_workbook(path)
    register = parse_register_file(path)

    records = tuple(iter_sos_variable_records(register, _revision(path)))

    assert len(records) == 7
    first, duplicate, conflict, storage_variant, partial, malformed, no_subset = records
    assert first.record_id == duplicate.record_id
    assert first.locators[0].physical_record == "row:2"
    assert duplicate.locators[0].physical_record == "row:3"
    assert conflict.record_id != first.record_id
    assert storage_variant.fields.model_dump(
        exclude={"coverage_from"}
    ) == first.fields.model_dump(exclude={"coverage_from"})
    assert storage_variant.fields.coverage_from is not None
    assert first.fields.coverage_from is not None
    assert (
        storage_variant.fields.coverage_from.value
        == first.fields.coverage_from.value
        == "2001"
    )
    assert storage_variant.fields.coverage_from.raw_value == "2001"
    assert first.fields.coverage_from.raw_value == 2001
    assert storage_variant.edition_scope == first.edition_scope
    assert storage_variant.record_id != first.record_id
    first_start = next(
        cell for cell in first.delivered_cells if cell.name == "Data från"
    )
    variant_start = next(
        cell for cell in storage_variant.delivered_cells if cell.name == "Data från"
    )
    assert first_start.raw_type == "int"
    assert variant_start.raw_type == "str"
    assert first.locators[0].semantic_record_key == (
        "register:Patientregistret källa",
        "deldatamangd:PAR_OV",
        "variable:HDIA",
    )
    assert first.locators[0].physical_file == path.name
    assert first.locators[0].physical_table == "Metadata - Variabelnivå"
    assert "Metadata - Variabelnivå!B2" in first.locators[0].physical_cells

    assert first.subject.register_name.name == "Patientregistret källa"
    assert first.subject.variant.name == "PAR_OV"
    assert first.subject.variable == SourceCoordinate(
        status="value", native_id="HDIA", name="Huvuddiagnos"
    )
    assert conflict.subject.variable.native_id == first.subject.variable.native_id
    assert conflict.subject.variable.name == "Annan etikett"
    assert first.subject.member.name == "HDIA"
    assert first.subject.native == NativeCoordinates()
    assert first.fields.column_name == value_field("HDIA")
    assert first.fields.name == value_field("Huvuddiagnos", raw=" Huvuddiagnos ")
    assert first.fields.description == value_field(
        "Första raden\nandra raden", raw="Första raden\nandra raden  "
    )
    assert first.fields.data_type == value_field("integer", raw="Heltal")
    assert first.fields.representation == value_field("Se kodlista")
    assert first.fields.source_attribution == value_field("Patientregistret")
    assert first.fields.reference_period is None
    classification_cell = next(
        cell for cell in first.delivered_cells if cell.name == "Länk kodverk"
    )
    assert (
        classification_cell.raw_type,
        classification_cell.storage_type,
        classification_cell.number_format,
        classification_cell.hyperlink_target,
    ) == ("str", "s", "General", _CLASSIFICATION_URL)
    assert first.edition_scope.kind == "intervals"
    assert first.edition_scope.intervals[0].start == "2001"
    assert first.edition_scope.intervals[0].end == "2020"
    assert first.edition_period_scope.kind == "not_applicable"

    assert partial.edition_scope.kind == "unknown"
    assert partial.edition_scope.label == "Data från=2010; Data till=<blank>"
    partial_link = next(
        cell for cell in partial.delivered_cells if cell.name == "Länk kodverk"
    )
    assert (
        partial_link.raw_value,
        partial_link.hyperlink_target,
        partial_link.hyperlink_location,
    ) == ("Visad kodlista", None, "'Kodlista_HDIA'!A1")
    assert malformed.edition_scope.kind == "unknown"
    assert malformed.edition_scope.label == "Data från=+2001; Data till=2020"
    assert no_subset.subject.variant.status == "unknown"
    assert no_subset.subject.variable.native_id == "UTAN_DEL"
    assert no_subset.locators[0].semantic_record_key[1] == "deldatamangd:<unknown>"
    assert no_subset.edition_scope.kind == "unknown"
    assert not any(
        "_default" in part for part in no_subset.locators[0].semantic_record_key
    )

    blank_subset = next(
        cell for cell in no_subset.delivered_cells if cell.name == "Deldatamängdsnamn"
    )
    assert (
        blank_subset.present,
        blank_subset.raw_value,
        blank_subset.interpreted_value,
    ) == (
        True,
        "",
        "",
    )


def _write_complete_workbook(path: Path) -> None:
    import openpyxl

    _write_source_workbook(path)
    workbook = openpyxl.load_workbook(path)
    dcat = workbook.create_sheet("Metadata-Datamängd (DCAT-AP)")
    dcat.append(["Attribut", "Definition", "Svenska", "Engelska"])
    dcat.append(["Titel", None, "Första titeln", "First title"])
    dcat.append(["Titel", None, "Andra titeln", "Second title"])
    dcat.append(["Beskrivning", None, "  indragen rad\n    tabell  kolumn", None])
    dcat.append(["Okänt attribut", None, "Obehandlad uppgift", "Unmapped fact"])
    subsets = workbook.create_sheet("Deldatamängder")
    subsets.append(
        ["Deldatamängdsnamn", "Deldatamängdsetikett", "Data från", "Data till"]
    )
    subsets.append(["PAR_OV", "Öppenvård", 1900, 2025])
    codes = workbook["Kodlista_HDIA"]
    codes.delete_rows(1, codes.max_row)
    codes.append(["Variabelnamn", "HDIA"])
    codes.append(["Variabelnamn", "ANNAN_VAR"])
    codes.append(["Tidsperiod", "Kod", "Beskrivning (kliniknamn)"])
    codes.append(["2010-2012", None, None])
    for description in ["Klinik ett", "Klinik ett", "Annan klinik"]:
        codes.append([None, 1, description])
        codes.cell(codes.max_row, 2).number_format = "000"
    raw = workbook.create_sheet("Kodlista_RAW")
    raw.append(["Sjukhus", "Adress"])
    raw.append(["Klinik", "Gatan 1"])
    workbook.save(path)
    workbook.close()


def test_common_parent_metadata_preserves_languages_conflicts_and_raw_context(
    tmp_path: Path,
) -> None:
    path = tmp_path / "Metadata Test.xlsx"
    _write_complete_workbook(path)
    cleaned = clean_sos_source(parse_register_file(path), _revision(path))

    parents = [
        record
        for record in cleaned.records
        if record.subject.member.status == "not_applicable"
    ]
    assert all(record.subject.variable.status == "not_applicable" for record in parents)
    assert all(record.fields.name is None for record in parents)
    titles = [
        record
        for record in parents
        if record.language and record.parent_facts[0].fields.name
    ]
    observed_titles = set()
    for record in titles:
        name = record.parent_facts[0].fields.name
        assert name is not None
        observed_titles.add((record.language, name.value))
    assert observed_titles == {
        ("sv", "Första titeln"),
        ("sv", "Andra titeln"),
        ("en", "First title"),
        ("en", "Second title"),
    }
    swedish = [record for record in titles if record.language == "sv"]
    assert (
        swedish[0].locators[0].semantic_record_key
        == swedish[1].locators[0].semantic_record_key
    )
    assert swedish[0].record_id != swedish[1].record_id
    description = next(
        record.parent_facts[0].fields.description
        for record in parents
        if record.language == "sv" and record.parent_facts[0].fields.description
    )
    assert description.value == "  indragen rad\n    tabell  kolumn"
    subset = next(
        record for record in parents if record.subject.variant.name == "PAR_OV"
    )
    assert subset.edition_scope.intervals[0].start == "1900"
    assert subset.parent_facts[0].kind == "variant"
    assert subset.parent_field_locators(0, "coverage_from")[0].physical_cells == (
        "Deldatamängder!C2",
    )
    assert swedish[0].parent_field_locators(0, "name")[0].physical_cells == (
        "Metadata-Datamängd (DCAT-AP)!C2",
    )
    partial = next(
        record for record in cleaned.records if record.subject.member.name == "PARTIELL"
    )
    assert partial.edition_scope.kind == "unknown"
    assert partial.fields.coverage_from is not None
    assert partial.fields.coverage_to is not None
    assert partial.fields.coverage_from.value == "2010"
    assert partial.fields.coverage_to.status == "unknown"
    dcat = next(
        table
        for table in cleaned.tables
        if table.name == "Metadata-Datamängd (DCAT-AP)"
    )
    assert any(
        cell.raw_value == "Unmapped fact" for row in dcat.rows for cell in row.cells
    )


def test_common_values_keep_duplicate_associations_competing_labels_and_period_origins(
    tmp_path: Path,
) -> None:
    path = tmp_path / "Metadata Test.xlsx"
    _write_complete_workbook(path)
    cleaned = clean_sos_source(parse_register_file(path), _revision(path))

    assert len(cleaned.values) == 2
    assert {value.normalized_content for value in cleaned.values.values()} == {
        ("001", "Klinik ett"),
        ("001", "Annan klinik"),
    }
    first, duplicate, conflict = cleaned.associations
    assert first.value_key == duplicate.value_key != conflict.value_key
    assert [association.row_number for association in cleaned.associations] == [5, 6, 7]
    assert first.supplied_period is None
    assert first.section_period == "2010-2012"
    assert first.section_locator is not None
    assert first.section_locator.physical_record == "row:4"
    descriptor = cleaned.descriptors["sheet:Kodlista_HDIA"]
    assert [(hint.role, hint.value) for hint in descriptor.member_hints] == [
        ("sheet_suffix", "HDIA"),
        ("list_header", "HDIA"),
        ("list_header", "ANNAN_VAR"),
    ]
    assert len(cleaned.values[first.value_key].locators) == 2
    raw = next(table for table in cleaned.tables if table.name == "Kodlista_RAW")
    assert all(row.role == "unparsed" for row in raw.rows)
    assert not any(
        association.descriptor_key == "sheet:Kodlista_RAW"
        for association in cleaned.associations
    )
    assert all(not record.code_set_references for record in cleaned.records)


def test_source_cleaning_blocks_on_parser_failure_with_original_evidence_available(
    tmp_path: Path,
) -> None:
    path = tmp_path / "Metadata Test.xlsx"
    _write_complete_workbook(path)
    parsed = parse_register_file(path)
    failed = replace(
        parsed,
        parse_issues=(
            SosParseIssue("Kodlista_HDIA", "parse_failure", "bad code layout"),
        ),
    )

    with pytest.raises(ValueError, match="Kodlista_HDIA: bad code layout"):
        clean_sos_source(failed, _revision(path))
    assert failed.source_sheets == parsed.source_sheets


def test_formatted_code_section_heading_is_evidence_not_a_value(tmp_path: Path) -> None:
    import openpyxl
    from openpyxl.styles import Font

    path = tmp_path / "Metadata Test.xlsx"
    _write_source_workbook(path)
    workbook = openpyxl.load_workbook(path)
    sheet = workbook["Kodlista_HDIA"]
    sheet.append([None, "Description of the following code digits", None])
    sheet["B3"].font = Font(bold=True)
    sheet.append([None, "01-19", "Literal code range"])
    for cell in sheet[4]:
        cell.font = Font(bold=True)
    sheet.append([None, "20", None])
    workbook.save(path)

    parsed = parse_register_file(path)
    cleaned = clean_sos_source(parsed, _revision(path))

    assert [row.kod for row in parsed.kodlistor[0].rows] == ["001", "01-19", "20"]
    assert {value.code for value in cleaned.values.values()} == {"001", "01-19", "20"}
    table = next(table for table in cleaned.tables if table.name == sheet.title)
    assert table.rows[2].role == "section"
    assert (
        table.rows[2].cells[1].raw_value == "Description of the following code digits"
    )


def _clean_code_rows(tmp_path: Path, rows: list[list[object]]):
    import openpyxl

    path = tmp_path / "Metadata Test.xlsx"
    _write_source_workbook(path)
    workbook = openpyxl.load_workbook(path)
    del workbook["Kodlista_HDIA"]
    sheet = workbook.create_sheet("Kodlista_Arbitrary")
    for row in rows:
        sheet.append(row)
    workbook.save(path)
    return clean_sos_source(parse_register_file(path), _revision(path))


def test_known_hidden_support_sheet_is_preserved_but_unknown_sheet_blocks(
    tmp_path: Path,
) -> None:
    import openpyxl

    path = tmp_path / "Metadata Test.xlsx"
    _write_source_workbook(path)
    workbook = openpyxl.load_workbook(path)
    support = workbook.create_sheet("Ej relevant_listor")
    support.sheet_state = "veryHidden"
    support.append(["Binär", "Datatyp"])
    support.append(["Ja", "Heltal"])
    workbook.save(path)

    cleaned = clean_sos_source(parse_register_file(path), _revision(path))
    table = next(table for table in cleaned.tables if table.name == support.title)
    assert [row.role for row in table.rows] == ["unparsed", "unparsed"]
    assert [[cell.raw_value for cell in row.cells] for row in table.rows] == [
        ["Binär", "Datatyp"],
        ["Ja", "Heltal"],
    ]
    assert table.rows[1].locator.physical_cells == (
        "Ej relevant_listor!A2",
        "Ej relevant_listor!B2",
    )
    assert len(cleaned.associations) == 1

    support.title = "New source structure"
    workbook.save(path)
    parsed = parse_register_file(path)
    with pytest.raises(
        ValueError, match="New source structure: unrecognized worksheet"
    ):
        clean_sos_source(parsed, _revision(path))
    assert parsed.source_sheets[-1].rows[1].source_evidence.cells[0].raw_value == "Ja"


def test_crosswalk_retains_explicit_recode_direction_and_literal_period(
    tmp_path: Path,
) -> None:
    cleaned = _clean_code_rows(
        tmp_path,
        [
            ["Variabelnamn", "Tidsperiod", "Indata Kod", "Kodat till", "Beskrivning"],
            ["FAMSIT", "1990/91-1994", "02", 3, " Annan familjesituation "],
            ["FAMSIT", "1990/91-1994", "02", 3, " Annan familjesituation "],
        ],
    )
    first, duplicate = cleaned.declarations
    assert isinstance(first, SourceCodeCrosswalkDeclaration)
    assert first.member_name.value == "FAMSIT"
    assert first.supplied_period.value == "1990/91-1994"
    assert first.section_period is None
    assert [
        (item.role, item.name, item.code.value, item.code.raw_value)
        for item in first.operands
    ] == [
        ("input", "Indata Kod", "02", "02"),
        ("output", "Kodat till", "3", 3),
    ]
    assert first.description.value == "Annan familjesituation"
    assert first.locator.physical_record == "row:2"
    assert duplicate.locator.physical_record == "row:3"
    assert first.revision == cleaned.revision
    assert len(first.delivered_cells) == 5
    assert not cleaned.values and not cleaned.associations


def test_crosswalk_keeps_peer_namespaces_and_section_period_separate(
    tmp_path: Path,
) -> None:
    cleaned = _clean_code_rows(
        tmp_path,
        [
            ["Används i variabeln", "BHEM"],
            ["Tidsperiod", "SCBkod", "SiSkod", "Namn"],
            ["1994-", None, None, None],
            [None, 16, 313, "Håkanstorp"],
            ["2000-2001", None, 314, "Ensidig uppgift"],
        ],
    )
    first, second = cleaned.declarations
    assert isinstance(first, SourceCodeCrosswalkDeclaration)
    assert first.member_name is None
    assert first.supplied_period.status == "unknown"
    assert first.section_period.value == "1994-"
    assert first.section_locator.physical_record == "row:3"
    assert [(item.role, item.name, item.code.value) for item in first.operands] == [
        ("peer", "SCBkod", "16"),
        ("peer", "SiSkod", "313"),
    ]
    assert second.operands[0].code.status == "unknown"
    assert second.supplied_period.value == "2000-2001"
    assert second.section_period.value == "1994-"
    assert (
        cleaned.descriptors["sheet:Kodlista_Arbitrary"].member_hints[1].value == "BHEM"
    )
    assert not cleaned.values and not cleaned.associations


def test_code_directory_retains_bounds_without_inventing_item_ids_or_end_dates(
    tmp_path: Path,
) -> None:
    cleaned = _clean_code_rows(
        tmp_path,
        [
            [
                "Variabelnamn",
                "Från",
                "Till",
                "Sjukhuskod",
                "Region",
                "Sjukhusnamn",
                "Aktuella",
                "Kommentar",
            ],
            [
                "SJUKHUS",
                2023,
                None,
                "10011",
                "Stockholm",
                " Sjukhus ",
                "JA",
                "Startade 202304",
            ],
        ],
    )
    (association,) = cleaned.associations
    (validity,) = cleaned.validity
    assert (validity.valid_from, validity.valid_to, validity.item_id) == (
        "2023",
        None,
        None,
    )
    assert validity.locator == association.locator
    assert validity.locators[0] == cleaned.values[association.value_key].locators[0]
    assert validity.delivered_cells[2].present
    assert validity.delivered_cells[2].raw_type == "none"
    assert association.delivered_cells[-1].raw_value == "Startade 202304"
    assert association.supplied_period is None and association.section_period is None
    assert association.member_hints[0].value == "SJUKHUS"
    assert cleaned.values[association.value_key].normalized_content == (
        "10011",
        "Sjukhus",
    )


def test_formula_and_delivered_cache_survive_without_promoting_computed_facts(
    tmp_path: Path,
) -> None:
    import xml.etree.ElementTree as ET
    import zipfile

    import openpyxl

    path = tmp_path / "Metadata Test.xlsx"
    _write_source_workbook(path)
    workbook = openpyxl.load_workbook(path)
    workbook["Metadata - Variabelnivå"]["C2"] = '="Computed label"'
    del workbook["Kodlista_HDIA"]
    sheet = workbook.create_sheet("Kodlista_HOSPITAL")
    sheet.append(
        [
            "Variabelnamn",
            "Från",
            "Till",
            "Sjukhuskod",
            "Region",
            "Sjukhusnamn",
            "Aktuella",
            "Kommentar",
        ]
    )
    formula = '=IF(C2>1, " ","JA")'
    sheet.append(
        ["SJUKHUS", 1973, 1980, "10010", "Stockholm", "Sjukhus", formula, None]
    )
    workbook.save(path)

    # openpyxl writes formula text but cannot populate a delivered formula cache.
    namespace = {"s": "http://schemas.openxmlformats.org/spreadsheetml/2006/main"}
    with zipfile.ZipFile(path) as source:
        members = {name: source.read(name) for name in source.namelist()}
    for name, payload in members.items():
        if not name.startswith("xl/worksheets/") or not name.endswith(".xml"):
            continue
        document = ET.fromstring(payload)
        for cell in document.findall(".//s:c", namespace):
            if cell.find("s:f", namespace) is None:
                continue
            cell.set("t", "str")
            cached = cell.find("s:v", namespace)
            assert cached is not None
            cached.text = " " if cell.get("r") == "G2" else "Computed label"
        members[name] = ET.tostring(document)
    with zipfile.ZipFile(path, "w") as output:
        for name, payload in members.items():
            output.writestr(name, payload)

    cleaned = clean_sos_source(parse_register_file(path), _revision(path))
    (association,) = cleaned.associations
    delivered = association.delivered_cells[6]
    assert (
        delivered.raw_value,
        delivered.storage_type,
        delivered.interpreted_value,
    ) == (formula, "f", "")
    assert (delivered.cached_raw_value, delivered.cached_raw_type) == (" ", "str")
    assert cleaned.values[association.value_key].code == "10010"
    variable = next(
        record
        for record in cleaned.records
        if record.subject.member.name == "HDIA"
        and record.locators[0].physical_record == "row:2"
    )
    assert variable.fields.name is not None
    assert variable.fields.name.status == "unknown"
    assert variable.fields.name.raw_value == '="Computed label"'
    assert variable.subject.variable.name is None
    assert variable.delivered_cells[2].cached_raw_value == "Computed label"


@pytest.mark.parametrize(
    "headers, values",
    [
        (
            [
                "Variabler",
                "ICD 8",
                "ICD 9 ",
                "ICD 10",
                "Åtgärdskoder 1963-1996",
                "Åtgärdskoder 1997-",
            ],
            [
                "SECFORE=1 eller SECAVSL=1",
                None,
                None,
                "O82 (alla underdiagnoser)",
                "7010, 7030",
                "MAH03/MAC23  (ersatt av MAC23 2007-12-31)",
            ],
        ),
        (
            ["Variabler", "Villkor", "Algoritm"],
            [
                "KON, BVIKT, GRDBS, BORDF2",
                "GRVB mellan 22 och 45",
                "BVIKT < (-0.001 * GRDBS**4)\n  + 123",
            ],
        ),
    ],
)
def test_derivation_keeps_named_literal_clauses_without_evaluation(
    tmp_path: Path, headers, values
) -> None:
    cleaned = _clean_code_rows(
        tmp_path,
        [
            ["Variabelnamn", "Tidsperiod", "Beskrivning", *headers],
            ["RESULTAT", "1973-", "Definition", *values],
        ],
    )
    (declaration,) = cleaned.declarations
    assert isinstance(declaration, SourceDerivationDeclaration)
    assert declaration.member_name.value == "RESULTAT"
    assert declaration.supplied_period.value == "1973-"
    assert [clause.name for clause in declaration.clauses] == [
        header.strip() for header in headers
    ]
    assert [clause.content.raw_value for clause in declaration.clauses] == values
    assert [clause.content.value for clause in declaration.clauses] == values
    assert not cleaned.values and not cleaned.associations and not cleaned.validity


def test_code_patterns_and_adjacent_legend_stay_literal_and_separate(
    tmp_path: Path,
) -> None:
    cleaned = _clean_code_rows(
        tmp_path,
        [
            ["KOD ", "Beskrivning", "Följande sympoler används", None, None],
            ["1XXXX", "missbildning i CNS", "X: valfri siffra", None, None],
            ["11xxy", "Annan kategori", "x: 1 unilateral", "Räknas ej", "Osäker"],
            [None, None, 1, "A", "K"],
        ],
    )
    assert {value.normalized_content for value in cleaned.values.values()} == {
        ("1XXXX", "missbildning i CNS"),
        ("11xxy", "Annan kategori"),
    }
    assert len(cleaned.associations) == 2
    assert all(len(value.delivered_cells) == 2 for value in cleaned.values.values())
    table = next(
        table for table in cleaned.tables if table.name == "Kodlista_Arbitrary"
    )
    assert table.rows[1].cells[2].raw_value == "X: valfri siffra"
    assert table.rows[3].cells[2].raw_value == "1"
    assert table.rows[3].cells[3].raw_value == "A"
    assert not cleaned.declarations and not cleaned.validity
