"""Common parent resolution supplies topology to ordinary variable formation."""

from __future__ import annotations

from contextlib import closing
from dataclasses import replace
from typing import TYPE_CHECKING

import pytest
from _csv_fixtures import REGISTERINFORMATION_HEADER, _var_row
from reg_meta.source_evidence import SourceRevision
from reg_meta_build.catalog_resolution import resolve_parents
from reg_meta_build.db import open_built_db
from reg_meta_build.resolved_catalog import write_resolved_catalog
from reg_meta_build.source_coding import resolve_code_membership
from reg_meta_build.source_coordinates import (
    native_column_key,
    native_parent_key,
    native_variant_key,
    source_register_key,
)
from reg_meta_build.source_curation import (
    CheckedEditionRebind,
    CurationCase,
    OccurrenceCorrectionDecision,
    PeerGuard,
    capture_expectations,
)
from reg_meta_build.source_effects import apply_occurrence_cases, record_ref
from reg_meta_build.source_formation import form_native_variable
from reg_meta_build.source_naming import (
    NamingDeclaration,
    NativeNamingTarget,
    check_naming_target,
)
from reg_meta_build.source_occurrences import source_occurrence
from reg_meta_build.source_records import SourceFields, value_field
from reg_meta_build.sources.scb_records import clean_scb_row

from reg_meta_build.fqid_slugs import SlugEntry

if TYPE_CHECKING:
    from pathlib import Path

    from reg_meta_build.source_records import SourceRecord


_REVISION = SourceRevision.create(
    dataset="fixture",
    publisher="SCB",
    purpose="common parent integration",
    upstream_revision="1",
    artifact_path="Registerinformation.csv",
    artifact_size=1,
    artifact_sha256="a" * 64,
)


def _record(row: int = 2, **changes: str) -> SourceRecord:
    header = REGISTERINFORMATION_HEADER.split("|")
    raw = (
        dict(
            zip(
                header,
                _var_row(colname="VALUE", cvid=4, var_id=5).split("|"),
                strict=True,
            )
        )
        | changes
    )
    cells: dict[str, tuple[bool, str | None, str]] = {
        name: (True, value, value) for name, value in raw.items()
    }
    return clean_scb_row(header, row, cells, _REVISION).record


def _names(record: SourceRecord) -> tuple[NamingDeclaration, ...]:
    # Exact checked identities are inputs of parent resolution; applicability itself
    # is exercised by the naming and pipeline-boundary tests.
    result = []
    for parent in record.parent_facts:
        if parent.kind not in {"register", "variant"}:
            continue
        key = native_parent_key(record.source, record.subject.provider, parent)
        assert key is not None
        kind = "register" if parent.kind == "register" else "register_variant"
        result.append(
            NamingDeclaration(
                target=NativeNamingTarget(
                    kind=kind,
                    provider=record.subject.provider,
                    source_key=key,
                    register_key=source_register_key(record)
                    if parent.kind == "variant"
                    else None,
                ),
                naming=SlugEntry(
                    kind=kind,
                    provider=record.subject.provider,
                    source_id="1" if parent.kind == "register" else "1.10",
                    slug="example" if parent.kind == "register" else "people",
                ),
                contributors=(),
            )
        )
    return tuple(result)


def test_explicit_rebind_roots_variant_edition_and_population_parents() -> None:
    flow = _record(Registerversionnamn="2007", RegVerID="1")
    stock = _record(3, CVID="6", Registerversionnamn="2007-12-31", RegVerID="2")
    native = native_variant_key(stock)
    assert native is not None
    split = (*native, "edition-split", "stock")
    original = source_occurrence(stock)
    assert original.edition_key is not None
    moved = replace(
        original,
        variant_key=split,
        edition_key=(*split, *original.edition_key[len(native) :]),
    )
    split_name = NamingDeclaration(
        target=NativeNamingTarget(
            kind="register_variant",
            provider="scb",
            source_key=split,
            register_key=source_register_key(stock),
        ),
        naming=SlugEntry(
            kind="register_variant",
            provider="scb",
            source_id="1.10.stock",
            slug="stock",
        ),
        contributors=(),
    )
    parents = resolve_parents(
        (source_occurrence(flow), moved),
        (*_names(flow), split_name),
        rebinds={record_ref(stock): split},
    )
    assert set(parents.variants) == {native, split}
    assert moved.edition_key in parents.editions
    assert any(
        key[: len(split)] == split and "population" in key for key in parents.fields
    )


def test_moved_occurrence_population_key_resolves_to_parent() -> None:
    record = _record(Registerversionnamn="2007-12-31", RegVerID="2")
    population = next(
        parent for parent in record.parent_facts if parent.kind == "population"
    )
    record = record.model_copy(
        update={
            "subject": record.subject.model_copy(
                update={"population": population.coordinate}
            )
        }
    )
    native = native_variant_key(record)
    assert native is not None
    split = (*native, "edition-split", "stock")
    ref = record_ref(record)
    case = CurationCase(
        case_id="stock-population",
        targets=capture_expectations((record,), fields=()),
        peer_guards=(
            PeerGuard(
                guard_id="stock-population",
                source=record.source,
                coordinates=(
                    ("register", record.subject.register_name),
                    ("variant", record.subject.variant),
                ),
                expected_members=(ref,),
            ),
        ),
        decision=OccurrenceCorrectionDecision(
            reviewed=True,
            effects=(CheckedEditionRebind(ref=ref, variant_key=split),),
            reason="SCB population distinguishes stock",
            provenance="fixture",
        ),
    )
    (moved,) = apply_occurrence_cases((record,), (case,)).occurrences
    native_parent = native_parent_key(
        record.source, record.subject.provider, population
    )
    assert native_parent is not None
    expected = (*split, *native_parent[len(native) :])
    assert moved.population_key == expected

    split_name = NamingDeclaration(
        target=NativeNamingTarget(
            kind="register_variant",
            provider="scb",
            source_key=split,
            register_key=source_register_key(record),
        ),
        naming=SlugEntry(
            kind="register_variant",
            provider="scb",
            source_id="1.10.stock",
            slug="stock",
        ),
        contributors=(),
    )
    register_name = next(
        name for name in _names(record) if name.target.kind == "register"
    )
    parents = resolve_parents(
        (moved,), (register_name, split_name), rebinds={ref: split}
    )
    assert expected in parents.fields
    name = parents.fields[expected].name
    assert name is not None
    assert name.value == population.coordinate.name


def test_parent_facts_feed_ordinary_formation_and_direct_catalog(
    tmp_path: Path,
) -> None:
    record = _record()
    parents = resolve_parents((record, record), _names(record))
    assert parents.diagnostics == ()
    assert len(parents.registers) == len(parents.variants) == len(parents.editions) == 1
    register_key, code_key = source_register_key(record), native_column_key(record)
    assert register_key is not None and code_key is not None
    formed = form_native_variable(
        (record,),
        register=parents.registers[register_key],
        variants=parents.variants,
        slug="value",
        provider_key="5",
        coding={code_key: resolve_code_membership(())},
        flags=SourceFields(
            identifier=value_field(False), sensitivity=value_field(False)
        ),
    )
    assert formed.variable is not None and formed.diagnostics == ()
    output = tmp_path / "catalog.db"
    write_resolved_catalog(
        (formed.variable,),
        output,
        manifest={"fixture": "parent-resolution"},
        editions=tuple(parents.editions.values()),
    )
    with closing(open_built_db(output)) as conn:
        assert conn.execute("SELECT COUNT(*) FROM register_version").fetchone()[0] == 1
        assert (
            conn.execute("SELECT delivery_column_name FROM variable_state").fetchone()[
                0
            ]
            == "VALUE"
        )
        assert (
            conn.execute("SELECT name FROM population").fetchone()[0]
            == "Hela befolkningen"
        )


def test_parent_conflict_does_not_choose_first_prose_or_change_delivery_dates() -> None:
    first, other = _record(), _record(3, Registersyfte="Conflicting purpose")
    parents = resolve_parents((first, other), _names(first))
    assert len(parents.registers) == 1
    register = next(iter(parents.registers.values()))
    assert register.purpose is None
    assert len(parents.diagnostics) == 1
    assert parents.diagnostics[0].fields == ("purpose",)
    assert parents.diagnostics[0].withheld_output == ("register.purpose",)
    assert len(parents.editions) == 1
    assert parents == resolve_parents((other, first), _names(first))


def test_edition_populations_do_not_become_competing_variable_assignments() -> None:
    records = (
        _record(Populationnamn="Adults"),
        _record(3, Populationnamn="Children"),
    )
    parents = resolve_parents(records, _names(records[0]))
    register_key = source_register_key(records[0])
    column_key = native_column_key(records[0])
    assert register_key is not None and column_key is not None
    formed = form_native_variable(
        records,
        register=parents.registers[register_key],
        variants=parents.variants,
        slug="value",
        provider_key="5",
        flags=SourceFields(
            identifier=value_field(False), sensitivity=value_field(False)
        ),
        coding={column_key: resolve_code_membership(())},
    )
    assert formed.variable is not None and formed.diagnostics == ()
    assert len(formed.variable.states) == 1
    assert formed.occurrences == records
    assert {
        population.name
        for edition in parents.editions.values()
        for population in edition.populations
    } == {"Adults", "Children"}


def test_missing_naming_is_an_implementation_failure() -> None:
    with pytest.raises(ValueError, match="missing checked parent naming"):
        resolve_parents((_record(),), ())


def test_parent_translations_are_accounted_without_conflicting_with_requested_language() -> (
    None
):
    swedish = _record().model_copy(update={"language": "sv"})
    english = _record(3, Registersyfte="English purpose").model_copy(
        update={"language": "en"}
    )
    parents = resolve_parents((swedish, english), _names(swedish))
    assert parents.diagnostics == ()
    assert parents.registers == resolve_parents((swedish,), _names(swedish)).registers
    assert parents.other_language_refs == (record_ref(english),)
    translated = resolve_parents((swedish, english), _names(swedish), language="en")
    assert next(iter(translated.registers.values())).purpose == "English purpose"
    assert translated.other_language_refs == (record_ref(swedish),)


def _write_sos_name_workbook(path: Path, *, dataset: str, title: str) -> None:
    import openpyxl

    workbook = openpyxl.Workbook()
    general = workbook.active
    general.title = "Generell information"
    general.append(["", "Om datamängden version", None])
    general.append(["", "Datamängd", dataset])
    general.append(["", "Version", "2026:1"])
    dcat = workbook.create_sheet("Metadata-Datamängd (DCAT-AP)")
    dcat.append(["Attribut", "Definition", "Svenska", "Engelska"])
    dcat.append(["Titel", None, title, None])
    variables = workbook.create_sheet("Metadata - Variabelnivå")
    variables.append(
        [
            "Deldatamängdsnamn",
            "Variabelnamn",
            "Variabeletikett",
            "Variabelbeskrivning",
            "Värdemängd",
            "Datatyp",
            "Länk kodverk",
            "Data från",
            "Data till",
        ]
    )
    variables.append(
        ["SOL_A", "INSATS", "Insats", None, None, "Heltal", None, 2001, 2020]
    )
    workbook.save(path)


def test_sos_dataset_label_does_not_conflict_with_dcat_title(tmp_path: Path) -> None:
    import hashlib

    from reg_meta_build.sources.sos import parse_register_file
    from reg_meta_build.sources.sos_records import clean_sos_source

    dataset = "Socialtjänstinsatser till äldre och personer med funktionsnedsättning"
    title = "Registret över socialtjänstinsatser till äldre och personer med funktionsnedsättning"
    path = tmp_path / "Metadata SOL (SOL)_webb.xlsx"
    _write_sos_name_workbook(path, dataset=dataset, title=title)
    payload = path.read_bytes()
    revision = SourceRevision.create(
        dataset="sos-metadata",
        publisher="Socialstyrelsen",
        purpose="SOS name precedence fixture",
        upstream_revision="2026:1",
        artifact_path=path.name,
        artifact_size=len(payload),
        artifact_sha256=hashlib.sha256(payload).hexdigest(),
    )
    cleaned = clean_sos_source(parse_register_file(path), revision)
    parents = tuple(
        record
        for record in cleaned.records
        if record.parent_facts
        and record.parent_facts[0].kind == "register"
        and record.language in {None, "sv"}
    )
    assert len(parents) == 3
    names = tuple(declaration for record in parents for declaration in _names(record))
    resolved = resolve_parents(parents, names)
    assert not [
        diagnostic
        for diagnostic in resolved.diagnostics
        if diagnostic.code == "conflicting_parent_metadata"
    ]
    assert resolved.diagnostics == ()
    (register,) = resolved.registers.values()
    assert register.name == title
    (key,) = resolved.registers.keys()
    assert resolved.fields[key].dataset_label is not None
    assert resolved.fields[key].dataset_label.value == dataset


def _variant_key(record: SourceRecord) -> tuple[str | int, ...]:
    parent = next(item for item in record.parent_facts if item.kind == "variant")
    key = native_parent_key(record.source, record.subject.provider, parent)
    assert key is not None
    return key


def _register_declaration(record: SourceRecord) -> NamingDeclaration:
    return next(
        declaration
        for declaration in _names(record)
        if declaration.target.kind == "register"
    )


def _checked_variant_declaration(pinned: SourceRecord) -> NamingDeclaration:
    # A checked parent naming declaration pins its endorsed observation exactly;
    # the second physical row keeps its own semantic member, so the pin stays
    # applicable instead of going stale on the added alternative.
    key = _variant_key(pinned)
    register_key = source_register_key(pinned)
    assert register_key is not None
    return NamingDeclaration(
        target=NativeNamingTarget(
            kind="register_variant",
            provider=pinned.subject.provider,
            source_key=key,
            register_key=register_key,
            expectations=capture_expectations(
                (pinned,), fields=("column_name",), parents=True
            ),
        ),
        naming=SlugEntry(
            kind="register_variant",
            provider=pinned.subject.provider,
            source_id="1.10",
            slug="people",
        ),
        contributors=(),
    )


def _conflicting_variant_pair() -> tuple[SourceRecord, SourceRecord]:
    first = _record()
    second = _record(
        3,
        CVID="6",
        VarId="7",
        Registervariantnamn="Huvudtabell",
        Registervariantbeskrivning="Annan beskrivning",
    )
    assert _variant_key(first) == _variant_key(second)
    assert record_ref(first) != record_ref(second)
    return first, second


def _other_variant_record() -> SourceRecord:
    return _record(
        4,
        CVID="8",
        VarId="9",
        Registervariantnamn="Inaktuell etikett",
        Registervariantbeskrivning="Inaktuell beskrivning",
    )


def test_checked_naming_declaration_resolves_conflicting_parent_name() -> None:
    first, second = _conflicting_variant_pair()
    key = _variant_key(first)
    declaration = _checked_variant_declaration(second)
    naming = (_register_declaration(first), declaration)
    assert check_naming_target(declaration.target, (first, second)) == ()
    parents = resolve_parents((first, second), naming)
    assert parents.variants[key].name == "Huvudtabell"
    assert parents.variants[key].description is None
    assert len(parents.editions) == 1
    name = parents.fields[key].name
    assert name is not None and name.value == "Huvudtabell"
    (conflict,) = [
        diagnostic
        for diagnostic in parents.diagnostics
        if diagnostic.code == "conflicting_parent_metadata"
    ]
    assert conflict.fields == ("description",)
    assert conflict.withheld_output == ("variant.description",)
    assert not [
        diagnostic
        for diagnostic in parents.diagnostics
        if diagnostic.code == "unknown_parent_name"
    ]


def test_conflicting_parent_name_without_declaration_still_withholds() -> None:
    first, second = _conflicting_variant_pair()
    key = _variant_key(first)
    parents = resolve_parents((first, second), (_register_declaration(first),))
    assert key not in parents.variants
    assert {diagnostic.code for diagnostic in parents.diagnostics} == {
        "conflicting_parent_metadata",
        "unknown_parent_name",
    }
    unknown = next(
        diagnostic
        for diagnostic in parents.diagnostics
        if diagnostic.code == "unknown_parent_name"
    )
    assert unknown.withheld_output == ("variant",)


def test_stale_naming_declaration_still_withholds_conflicting_parent_name() -> None:
    first, second = _conflicting_variant_pair()
    key = _variant_key(first)
    stale = _other_variant_record()
    declaration = _checked_variant_declaration(stale)
    assert check_naming_target(declaration.target, (first, second)) != ()
    parents = resolve_parents(
        (first, second),
        (_register_declaration(first), declaration),
        withheld_naming=frozenset({key}),
    )
    assert key not in parents.variants
    assert {diagnostic.code for diagnostic in parents.diagnostics} == {
        "conflicting_parent_metadata",
        "unknown_parent_name",
    }


def test_naming_declaration_for_unobserved_name_never_invents_parent_name() -> None:
    first, second = _conflicting_variant_pair()
    key = _variant_key(first)
    other = _other_variant_record()
    declaration = _checked_variant_declaration(other)
    parents = resolve_parents(
        (first, second), (_register_declaration(first), declaration)
    )
    assert key not in parents.variants
    assert {diagnostic.code for diagnostic in parents.diagnostics} == {
        "conflicting_parent_metadata",
        "unknown_parent_name",
    }


def test_support_occurrence_retains_parents_without_materializing_delivery() -> None:
    record = _record(Registerversionnamn="2012, preliminär version")
    support = replace(source_occurrence(record), use="support")
    parents = resolve_parents((support,), _names(record))
    assert replace(parents, support_only_refs=()) == resolve_parents(
        (record,), _names(record)
    )
    assert parents.support_only_refs == (record_ref(record),)
    assert len(parents.editions) == 1
    with pytest.raises(ValueError, match="support-only"):
        form_native_variable(
            (support,),
            register=next(iter(parents.registers.values())),
            variants=parents.variants,
            slug="value",
            provider_key="5",
            flags=SourceFields(),
            coding={},
        )


def test_support_parent_conflict_retains_existing_withholding_guard() -> None:
    first, second = _conflicting_variant_pair()
    names = (_register_declaration(first), _checked_variant_declaration(first))
    ordinary = resolve_parents((first, second), names)
    support = replace(source_occurrence(second), use="support")
    resolved = resolve_parents((first, support), names)
    assert replace(resolved, support_only_refs=()) == ordinary
    assert resolved.support_only_refs == (record_ref(second),)
    assert any(d.code == "conflicting_parent_metadata" for d in resolved.diagnostics)


def test_support_lookup_parent_requires_admitted_variant_topology() -> None:
    record = _record(Registervariantnamn="AGARKAT")
    support = replace(source_occurrence(record), use="support")
    register_name = tuple(n for n in _names(record) if n.target.kind == "register")
    parents = resolve_parents((support,), register_name)
    assert len(parents.registers) == 1
    assert parents.variants == parents.editions == {}
    assert all("variant" not in key for key in parents.fields)
    assert parents.support_only_refs == (record_ref(record),)
    assert parents.diagnostics == ()
