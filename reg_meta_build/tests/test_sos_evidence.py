"""Parser-side evidence retention for Socialstyrelsen workbook structures."""

from __future__ import annotations

from typing import TYPE_CHECKING

from reg_meta_build.sources.sos import parse_register_file

if TYPE_CHECKING:
    from pathlib import Path

    import pytest
    from reg_meta_build.sources.sos import SosRegister


def _write_evidence_workbook(path: Path) -> None:
    import openpyxl

    workbook = openpyxl.Workbook()
    general = workbook.active
    general.title = "Generell information"
    general.append([None, "Om metadatamallen", None])
    general.append([None, "Version", "1.0"])
    general.append([None, "Version", "2.0"])
    general.append([None, "Om datamängden version", None])
    general.append([None, "Datamängd", "Testregistret"])
    general.append([None, "Okänd uppgift", "bevaras"])

    dcat = workbook.create_sheet("Metadata-Datamängd (DCAT-AP)")
    dcat.append(["Attribut SoS-metadata", "Definition", "Svenska", "Engelska"])
    dcat.append(["Titel", "Definition ett", "Svensk titel ett", "English title one"])
    dcat.append(["Titel", "Definition två", "Svensk titel två", "English title two"])
    dcat.append(["Okänt attribut", "Fri definition", "Svenskt värde", "English value"])

    subsets = workbook.create_sheet("Deldatamängder och datavyer")
    subsets.append(
        ["Deldatamängdsnamn", "Deldatamängdsetikett", "Data från", "Data till"]
    )
    subsets.append(["DEL_A", "Del A", 2001, 2020])
    subsets.append([None, "Etikett utan namn", 2010, 2015])

    variables = workbook.create_sheet("Metadata - Variabelnivå")
    variables.append(["Deldatamängdsnamn", "Variabelnamn", "Variabeletikett"])
    variables.append(["DEL_A", "VAR_A", "Variabel A", "unheaded source cell"])
    variables.append(["DEL_A", None, "Etikett utan variabelnamn"])

    codes = workbook.create_sheet("Kodlista_VAR_A")
    codes.append(["Kodverk", "Testkoder", None])
    codes.append(["Variabelnamn", "VAR_A", None])
    codes.append(["Bakgrund", "Bakgrundstext", None])
    codes.append(["Tidsperiod", "Kod", "Beskrivning (kliniknamn)"])
    codes.append(["2020", None, None])
    codes.append([None, "01", "Klinik ett"])
    codes.append([None, None, "Oförstådd efterrad"])

    raw_codes = workbook.create_sheet("Kodlista_RAW")
    raw_codes.append(["Sjukhuskatalog", "Region"])
    raw_codes.append(["Sjukhus ett", "Nord"])

    quality = workbook.create_sheet("Kvalitet_VAR_A")
    quality.append(["Rubrik", "Text"])
    quality.append(["Aktualitet", "Årlig"])
    quality.merge_cells("A3:B3")
    quality["A3"] = "Sammanslagen rubrik"

    workbook.save(path)


def _sheet(register: SosRegister, kind: str, name: str):
    return next(
        sheet
        for sheet in register.source_sheets
        if sheet.kind == kind and sheet.sheet_name == name
    )


def test_all_current_sos_sheet_structures_retain_ordered_original_rows(
    tmp_path: Path,
) -> None:
    path = tmp_path / "Metadata Test (TST)_webb.xlsx"
    _write_evidence_workbook(path)

    register = parse_register_file(path)

    assert [sheet.kind for sheet in register.source_sheets] == [
        "general",
        "dcat",
        "subsets",
        "variables",
        "codelist",
        "codelist",
        "quality",
    ]
    assert register.parse_issues == ()

    general = _sheet(register, "general", "Generell information")
    assert [row.role for row in general.rows] == [
        "section",
        "metadata",
        "metadata",
        "section",
        "metadata",
        "metadata",
    ]
    repeated_versions = [
        next(
            cell
            for cell in row.source_evidence.cells
            if cell.field_name == "template_version"
        )
        for row in general.rows
        if any(
            cell.field_name == "template_version" for cell in row.source_evidence.cells
        )
    ]
    assert [cell.raw_value for cell in repeated_versions] == ["1.0", "2.0"]
    unknown_general = general.rows[-1].source_evidence
    assert [cell.raw_value for cell in unknown_general.cells] == [
        None,
        "Okänd uppgift",
        "bevaras",
    ]
    assert [cell.field_name for cell in unknown_general.cells] == [
        None,
        "attribute",
        "value",
    ]

    dcat = _sheet(register, "dcat", "Metadata-Datamängd (DCAT-AP)")
    assert [row.role for row in dcat.rows] == [
        "header",
        "attribute",
        "attribute",
        "attribute",
    ]
    title_rows = dcat.rows[1:3]
    assert [
        next(
            cell for cell in row.source_evidence.cells if cell.field_name == "title_sv"
        ).raw_value
        for row in title_rows
    ] == ["Svensk titel ett", "Svensk titel två"]
    first_title = title_rows[0].source_evidence.cells
    assert {
        (cell.field_name, cell.language, cell.raw_value)
        for cell in first_title
        if cell.language is not None
    } == {
        ("title_sv", "sv", "Svensk titel ett"),
        ("title_en", "en", "English title one"),
    }
    unknown_dcat = dcat.rows[-1].source_evidence.cells
    assert {cell.field_name for cell in unknown_dcat} >= {
        "attribute",
        "value_sv",
        "value_en",
    }

    subsets = _sheet(register, "subsets", "Deldatamängder och datavyer")
    assert [row.role for row in subsets.rows] == ["header", "subset", "raw"]
    skipped_subset = subsets.rows[-1].source_evidence
    assert (
        next(
            cell for cell in skipped_subset.cells if cell.field_name == "label"
        ).raw_value
        == "Etikett utan namn"
    )

    variables = _sheet(register, "variables", "Metadata - Variabelnivå")
    assert [row.role for row in variables.rows] == ["header", "variable", "raw"]
    assert register.variables[0].source_evidence == variables.rows[1].source_evidence
    unheaded = variables.rows[1].source_evidence.cells[-1]
    assert (unheaded.header, unheaded.field_name, unheaded.raw_value) == (
        "D",
        None,
        "unheaded source cell",
    )

    codes = _sheet(register, "codelist", "Kodlista_VAR_A")
    assert [row.role for row in codes.rows] == [
        "preamble",
        "preamble",
        "preamble",
        "header",
        "period_section",
        "code",
        "raw",
    ]
    code = register.kodlistor[0].rows[0]
    assert code.tidsperiod == "2020"
    assert code.source_tidsperiod is None
    code_evidence = code.source_evidence
    assert code_evidence is not None
    assert code_evidence == codes.rows[5].source_evidence
    assert (
        next(
            cell for cell in code_evidence.cells if cell.field_name == "beskrivning"
        ).raw_value
        == "Klinik ett"
    )
    assert (
        next(
            cell
            for cell in codes.rows[4].source_evidence.cells
            if cell.field_name == "tidsperiod"
        ).raw_value
        == "2020"
    )

    raw_codes = _sheet(register, "codelist", "Kodlista_RAW")
    assert [row.role for row in raw_codes.rows] == ["raw", "raw"]
    assert register.kodlistor[1].rows == ()
    assert register.kodlistor[1].raw_rows[1][:2] == ("Sjukhus ett", "Nord")

    quality = _sheet(register, "quality", "Kvalitet_VAR_A")
    assert [row.role for row in quality.rows] == ["quality", "quality", "quality"]
    assert [cell.raw_value for cell in quality.rows[1].source_evidence.cells] == [
        "Aktualitet",
        "Årlig",
    ]
    merged = quality.rows[2].source_evidence.cells
    assert [(cell.coordinate, cell.raw_value) for cell in merged] == [
        ("A3", "Sammanslagen rubrik"),
        ("B3", None),
    ]


def test_code_parse_failure_retains_original_rows_and_explicit_issue(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    from reg_meta_build.sources import sos

    path = tmp_path / "Metadata Test (TST)_webb.xlsx"
    _write_evidence_workbook(path)

    def fail(*_args: object, **_kwargs: object) -> None:
        raise ValueError("forced code parse failure")

    monkeypatch.setattr(sos, "_parse_kodlista", fail)
    register = parse_register_file(path)

    assert len(register.parse_issues) == 2
    assert all(issue.kind == "code_list_parse_error" for issue in register.parse_issues)
    assert all(
        "forced code parse failure" in issue.detail for issue in register.parse_issues
    )
    assert register.kodlistor[0].raw_rows[0][:2] == ("Kodverk", "Testkoder")
    assert register.kodlistor[1].raw_rows[0][:2] == ("Sjukhuskatalog", "Region")
    assert not any(
        "<unparseable>" in str(value)
        for kodlista in register.kodlistor
        for row in kodlista.raw_rows
        for value in row
    )
    assert [
        row.role for row in _sheet(register, "codelist", "Kodlista_VAR_A").rows
    ] == [
        "preamble",
        "preamble",
        "preamble",
        "header",
        "period_section",
        "code",
        "raw",
    ]
