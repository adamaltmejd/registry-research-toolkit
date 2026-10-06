"""Crosswalk, directory and derivation declarations from Socialstyrelsen workbooks."""

from __future__ import annotations

from typing import TYPE_CHECKING

import pytest
from _sos_fixtures import (
    clean_code_rows as _clean_code_rows,
    source_revision as _revision,
    write_source_workbook as _write_source_workbook,
)
from reg_meta.documentary import (
    SourceCodeCrosswalkDeclaration,
    SourceDerivationDeclaration,
)
from reg_meta_build.sources.sos import parse_register_file
from reg_meta_build.sources.sos_records import (
    clean_sos_source,
)

if TYPE_CHECKING:
    from pathlib import Path


def test_crosswalk_retains_explicit_recode_direction_and_literal_period(
    tmp_path: Path,
) -> None:
    cleaned = _clean_code_rows(
        tmp_path,
        [
            ["Variabelnamn", "Tidsperiod", "Indata Kod", "Kodat till", "Beskrivning"],
            ["FAMSIT", "1990/91-1994", "02", 3, " Annan familjesituation "],
            ["FAMSIT", "1990/91-1994", "02", 3, " Annan familjesituation "],
        ],
    )
    first, duplicate = cleaned.declarations
    assert isinstance(first, SourceCodeCrosswalkDeclaration)
    assert first.member_name.value == "FAMSIT"
    assert first.supplied_period.value == "1990/91-1994"
    assert first.section_period is None
    assert [
        (item.role, item.name, item.code.value, item.code.raw_value)
        for item in first.operands
    ] == [
        ("input", "Indata Kod", "02", "02"),
        ("output", "Kodat till", "3", 3),
    ]
    assert first.description.value == "Annan familjesituation"
    assert first.locator.physical_record == "row:2"
    assert duplicate.locator.physical_record == "row:3"
    assert first.revision == cleaned.revision
    assert len(first.delivered_cells) == 5
    assert not cleaned.values and not cleaned.associations


def test_non_code_list_sheets_keep_evidence_without_binding_targets(
    tmp_path: Path,
) -> None:
    import openpyxl

    path = tmp_path / "Metadata Test.xlsx"
    _write_source_workbook(path)
    workbook = openpyxl.load_workbook(path)
    shapes = {
        "Kodlista_famsit": (
            ["Variabelnamn", "Tidsperiod", "Indata Kod", "Kodat till", "Beskrivning"],
            ["HDIA", "1990/91-1994", "02", "3", "Recoding"],
        ),
        "Kodlista_förlossningssätt": (
            [
                "Variabelnamn",
                "Tidsperiod",
                "Beskrivning",
                "Variabler",
                "ICD 8",
                "ICD 9",
                "ICD 10",
                "Åtgärdskoder 1963-1996",
                "Åtgärdskoder 1997-",
            ],
            ["HDIA", "1973-", "Delivery mode", "A", "B", "C", "D", "E", "F"],
        ),
        "Kodlista_msga_mlga": (
            [
                "Variabelnamn",
                "Tidsperiod",
                "Beskrivning",
                "Variabler",
                "Villkor",
                "Algoritm",
            ],
            ["HDIA", "1973-", "Growth", "A", "A > 0", "A * 2"],
        ),
    }
    for name, (header, row) in shapes.items():
        sheet = workbook.create_sheet(name)
        sheet.append(header)
        sheet.append(row)
    workbook.save(path)
    workbook.close()

    parsed = parse_register_file(path)
    cleaned = clean_sos_source(parsed, _revision(path))
    assert "sheet:Kodlista_HDIA" in cleaned.descriptors
    assert len(cleaned.declarations) == 3
    for name, (_, row) in shapes.items():
        kodlista = next(item for item in parsed.kodlistor if item.sheet_name == name)
        assert kodlista.rows == ()
        assert kodlista.raw_rows[1] == tuple(row)
        sheet = next(item for item in parsed.source_sheets if item.sheet_name == name)
        assert sheet.kind == "documentation"
        assert [cell.raw_value for cell in sheet.rows[1].source_evidence.cells] == row
        assert f"sheet:{name}" not in cleaned.descriptors
        table = next(item for item in cleaned.tables if item.name == name)
        assert [cell.raw_value for cell in table.rows[1].cells] == row


def test_crosswalk_keeps_peer_namespaces_and_section_period_separate(
    tmp_path: Path,
) -> None:
    cleaned = _clean_code_rows(
        tmp_path,
        [
            ["Används i variabeln", "BHEM"],
            ["Tidsperiod", "SCBkod", "SiSkod", "Namn"],
            ["1994-", None, None, None],
            [None, 16, 313, "Håkanstorp"],
            ["2000-2001", None, 314, "Ensidig uppgift"],
        ],
    )
    first, second = cleaned.declarations
    assert isinstance(first, SourceCodeCrosswalkDeclaration)
    assert first.member_name is None
    assert first.supplied_period.status == "unknown"
    assert first.section_period.value == "1994-"
    assert first.section_locator.physical_record == "row:3"
    assert [(item.role, item.name, item.code.value) for item in first.operands] == [
        ("peer", "SCBkod", "16"),
        ("peer", "SiSkod", "313"),
    ]
    assert second.operands[0].code.status == "unknown"
    assert second.supplied_period.value == "2000-2001"
    assert second.section_period.value == "1994-"
    assert "sheet:Kodlista_Arbitrary" not in cleaned.descriptors
    assert cleaned.declarations
    assert not cleaned.values and not cleaned.associations


def test_code_directory_retains_bounds_without_inventing_item_ids_or_end_dates(
    tmp_path: Path,
) -> None:
    cleaned = _clean_code_rows(
        tmp_path,
        [
            [
                "Variabelnamn",
                "Från",
                "Till",
                "Sjukhuskod",
                "Region",
                "Sjukhusnamn",
                "Aktuella",
                "Kommentar",
            ],
            [
                "SJUKHUS",
                2023,
                None,
                "10011",
                "Stockholm",
                " Sjukhus ",
                "JA",
                "Startade 202304",
            ],
        ],
    )
    (association,) = cleaned.associations
    (validity,) = cleaned.validity
    assert "sheet:Kodlista_Arbitrary" in cleaned.descriptors
    assert (validity.valid_from, validity.valid_to, validity.item_id) == (
        "2023",
        None,
        None,
    )
    assert validity.locator == association.locator
    assert validity.locators[0] == cleaned.values[association.value_key].locators[0]
    assert validity.delivered_cells[2].present
    assert validity.delivered_cells[2].raw_type == "none"
    assert association.delivered_cells[-1].raw_value == "Startade 202304"
    assert association.supplied_period is None and association.section_period is None
    assert association.member_hints[0].value == "SJUKHUS"
    assert cleaned.values[association.value_key].normalized_content == (
        "10011",
        "Sjukhus",
    )


def test_nonstandard_code_sheet_keeps_unknown_validity_for_binding(
    tmp_path: Path,
) -> None:
    cleaned = _clean_code_rows(
        tmp_path,
        [
            [
                "Variabelnamn",
                "Från",
                "Till",
                "Sjukhuskod",
                "Region",
                "Sjukhusnamn",
                "Aktuella",
                "Kommentar",
            ],
            ["METOD", "unparseable", None, "1", "Region", "Method", None, None],
        ],
    )
    assert "sheet:Kodlista_Arbitrary" in cleaned.descriptors
    (association,) = cleaned.associations
    (validity,) = cleaned.validity
    assert association.value_key in cleaned.values
    assert validity.window is not None
    assert validity.window.status == "unknown"


@pytest.mark.parametrize(
    "headers, values",
    [
        (
            [
                "Variabler",
                "ICD 8",
                "ICD 9 ",
                "ICD 10",
                "Åtgärdskoder 1963-1996",
                "Åtgärdskoder 1997-",
            ],
            [
                "SECFORE=1 eller SECAVSL=1",
                None,
                None,
                "O82 (alla underdiagnoser)",
                "7010, 7030",
                "MAH03/MAC23  (ersatt av MAC23 2007-12-31)",
            ],
        ),
        (
            ["Variabler", "Villkor", "Algoritm"],
            [
                "KON, BVIKT, GRDBS, BORDF2",
                "GRVB mellan 22 och 45",
                "BVIKT < (-0.001 * GRDBS**4)\n  + 123",
            ],
        ),
    ],
)
def test_derivation_keeps_named_literal_clauses_without_evaluation(
    tmp_path: Path, headers, values
) -> None:
    cleaned = _clean_code_rows(
        tmp_path,
        [
            ["Variabelnamn", "Tidsperiod", "Beskrivning", *headers],
            ["RESULTAT", "1973-", "Definition", *values],
        ],
    )
    (declaration,) = cleaned.declarations
    assert isinstance(declaration, SourceDerivationDeclaration)
    assert declaration.member_name.value == "RESULTAT"
    assert declaration.supplied_period.value == "1973-"
    assert [clause.name for clause in declaration.clauses] == [
        header.strip() for header in headers
    ]
    assert [clause.content.raw_value for clause in declaration.clauses] == values
    assert [clause.content.value for clause in declaration.clauses] == values
    assert not cleaned.values and not cleaned.associations and not cleaned.validity
