"""Focused tests for LISA/SCB source records and diagnostic target inspection."""

from __future__ import annotations

import hashlib
import json
import subprocess
from contextlib import contextmanager
from typing import TYPE_CHECKING

import pytest
from _csv_fixtures import (
    _var_row,
    sparsify_scb_values,
    write_input_bundle,
    write_scb_input,
    write_scb_snapshot,
)
from _lisa_fixtures import write_lisa_workbook
from openpyxl import load_workbook
from pydantic import ValidationError
from reg_meta.errors import EXIT_CONFIG, RegMetaError
from reg_meta_build.input_snapshot import (
    LISA_DATASET_ID,
    LisaWorkbookSelection,
    ScbSnapshotReader,
    SnapshotError,
    open_input_bundle,
    open_scb_snapshot,
)
from reg_meta_build.source_inspection import (
    SourceInspectionReport,
    compare_availability_records,
    inspect_bundle_source_records,
    report_semantic_sha256,
)
from reg_meta_build.source_records import (
    NativeCoordinates,
    RecordLocator,
    ScopeInterval,
    SourceCoordinate,
    SourceField,
    SourceFields,
    SourceRecord,
    SourceRevision,
    SourceSubject,
    TemporalScope,
    value_field,
)
from reg_meta_build.sources import lisa as lisa_module
from reg_meta_build.sources.lisa import LisaWorkbookError, read_lisa_source
from reg_meta_build.sources.scb_records import _scopes, read_scb_lisa_records

from reg_meta_build import cli as cli_module, source_inspection as inspection_module

if TYPE_CHECKING:
    from collections.abc import Iterator
    from pathlib import Path
    from typing import Any


def _revision(path: Path, dataset: str = LISA_DATASET_ID) -> SourceRevision:
    payload = path.read_bytes()
    return SourceRevision.create(
        dataset=dataset,
        publisher="SCB",
        purpose="fixture source records",
        upstream_revision="fixture-2024-2025",
        artifact_path=path.name,
        artifact_size=len(payload),
        artifact_sha256=hashlib.sha256(payload).hexdigest(),
    )


def _field_text(record: SourceRecord, field: str) -> str | None:
    observation = getattr(record.fields, field)
    if observation is None or observation.status != "value":
        return None
    assert isinstance(observation.value, str)
    return observation.value


_QUALIFICATION_CONTEXT = {
    (
        "worksheet-context Individ!A139: Befolkningens arbetsmarknadsstatus "
        "(BAS) är källa från 2022 om inget annat år anges"
    ),
    (
        "worksheet-context Företag!A48: Ekonomiska nyckeltal och ekonomisk "
        "grunddata finns för företag som ingår i Företagens ekonomi (FEK)."
    ),
    (
        "worksheet-context Företag!A49: FEK täcker näringslivet (exklusive de "
        "finansiella och offentliga sektorerna samt hushållens icke-vinstdrivande"
    ),
    "worksheet-context Företag!A50:  organisationer).",
    (
        "worksheet-context Företag!A91: Från 2024 inkluderas godkända "
        "resultaträkningar även om balansräkning är underkänd och tvärtom."
    ),
}


def _record(
    revision: SourceRevision,
    *,
    key: str,
    column: str | None,
    year: int,
    availability: SourceField | None = None,
    variant_id: int = 153,
    variable_id: int = 1,
    member_id: int = 1,
    workbook_table: str | None = None,
    population_name: str | None = None,
) -> SourceRecord:
    return SourceRecord.create(
        revision=revision,
        locator=RecordLocator(
            semantic_record_key=(key,),
            physical_file="fixture",
            physical_table="fixture",
            physical_record=key,
            physical_cells=(f"fixture:{key}",),
        ),
        subject=SourceSubject(
            provider="scb",
            register=SourceCoordinate(status="value", native_id=34, name="LISA"),
            variant=(
                SourceCoordinate(status="value", name=workbook_table)
                if workbook_table is not None
                else SourceCoordinate(status="value", native_id=variant_id)
            ),
            population=(
                SourceCoordinate(status="value", name=population_name)
                if population_name is not None
                else SourceCoordinate(status="unknown")
            ),
            member=SourceCoordinate(
                status="value",
                native_id=None if workbook_table is not None else member_id,
                name=column if workbook_table is not None else None,
            ),
            native=(
                NativeCoordinates()
                if workbook_table is not None
                else NativeCoordinates(
                    register_id=34,
                    register_variant_id=variant_id,
                    edition_id=year,
                    variable_id=variable_id,
                    member_id=member_id,
                )
            ),
        ),
        edition_scope=TemporalScope(
            kind="intervals",
            intervals=(ScopeInterval(start=str(year), end=str(year)),),
        ),
        reference_period_scope=TemporalScope(kind="not_applicable"),
        fields=SourceFields(
            availability=availability or value_field(True),
            column_name=(
                value_field(column)
                if column is not None
                else SourceField(status="unknown", raw_value="")
            ),
        ),
    )


def test_lisa_reader_preserves_four_layouts_sections_periods_and_occurrences(
    tmp_path: Path,
) -> None:
    path = write_lisa_workbook(tmp_path / "lisa.xlsx")
    source = read_lisa_source(path, _revision(path))
    records = source.records

    assert len(records) == 9
    assert {record.locator.physical_table for record in records} == {
        "Individ",
        "Individ årsoberoende",
        "Företag",
        "Arbetsställe",
    }
    person_numbers = [
        record for record in records if _field_text(record, "column_name") == "PersonNr"
    ]
    assert len(person_numbers) == 4
    assert len({record.record_id for record in person_numbers}) == 4
    assert {record.context[1] for record in person_numbers} == {
        "Födelseland",
        "Avlidna",
        "In- och Utvandring",
    }
    assert all(
        record.edition_scope.kind == "year_independent" for record in person_numbers
    )

    rekr = next(
        record for record in records if _field_text(record, "column_name") == "RekrBidr"
    )
    assert rekr.original_period_text == "2003-2008\n2013"
    assert [(item.start, item.end) for item in rekr.edition_scope.intervals] == [
        ("2003", "2008"),
        ("2013", "2013"),
    ]
    workplace = next(
        record
        for record in records
        if _field_text(record, "column_name") == "Ast_LoneSum"
    )
    assert [(item.start, item.end) for item in workplace.edition_scope.intervals] == [
        ("1990", "2021"),
        ("2024", "2024"),
    ]
    company = next(
        record
        for record in records
        if _field_text(record, "column_name") == "ForetagKod"
    )
    assert company.edition_scope.intervals == (ScopeInterval(start="2024", end="2024"),)
    assert company.subject.variant.name == "company"

    ampoltyp = next(
        record for record in records if _field_text(record, "column_name") == "AmPolTyp"
    )
    assert ampoltyp.locator.physical_record == "row:600"
    assert _field_text(ampoltyp, "description") == (
        "Typ av arbetsmarknadspolitisk åtgärd"
    )
    assert ampoltyp.fields.sensitivity == SourceField(
        status="value", value=False, raw_value="Nej"
    )
    assert ampoltyp.fields.availability == value_field(True)

    ku2 = next(
        record
        for record in records
        if _field_text(record, "column_name") == "KU2YrkStalln"
    )
    assert ku2.fields.sensitivity == SourceField(
        status="value", value="conditional", raw_value="I vissa fall"
    )
    assert ku2.context[-2:] == (
        (
            "continuation-note Individ!B323: Före 2010 finns inte kod 5 "
            "(företagare i eget AB). För åren 1993- finns istället variabeln "
            "KU1Faman som anges som 1 om personen är en företagare i eget AB."
        ),
        "continuation-note Individ!F323: RAMS-Jobb",
    )
    assert ku2.locator.physical_cells[-2:] == ("Individ!B323", "Individ!F323")
    recognized_context = tuple(
        item
        for item in source.worksheet_context
        if item.startswith("worksheet-context ")
    )
    expected_context = {
        f"worksheet-context {sheet}!A{row}: {text}"
        for sheet, spec in lisa_module._TABLES.items()
        for row, text in spec.text_rows.items()
    }
    assert set(recognized_context) == expected_context
    assert len(recognized_context) == len(expected_context)
    assert set(recognized_context) >= _QUALIFICATION_CONTEXT

    footnotes = tuple(
        item
        for item in source.worksheet_context
        if item.startswith("worksheet-footnote ")
    )
    assert len(footnotes) == 5
    assert footnotes[0] == ("worksheet-footnote Individ!B815: _ftnref2")
    assert footnotes[-1].startswith(
        "worksheet-footnote Individ!B819: 7 Från och med årgång 2020"
    )
    assert all(
        "Individ!B815" not in record.locator.physical_cells for record in records
    )
    assert all(
        all(
            not item.startswith(("worksheet-context ", "worksheet-footnote "))
            for item in record.context
        )
        for record in records
    )


@pytest.mark.parametrize(
    ("mutation", "expected"),
    (
        ("shift-header", r"Individ!A3:F3"),
        ("change-section", r"Individ årsoberoende!A30"),
        ("bad-period", r"Individ!C600"),
        ("extra-column", r"expected 6 columns, got 7"),
    ),
)
def test_lisa_reader_rejects_changed_actual_layout_with_coordinates(
    tmp_path: Path, mutation: str, expected: str
) -> None:
    path = write_lisa_workbook(tmp_path / "lisa.xlsx")
    workbook = load_workbook(path)
    if mutation == "shift-header":
        sheet = workbook["Individ"]
        for column in range(1, 7):
            sheet.cell(4, column, sheet.cell(3, column).value)
            sheet.cell(3, column).value = None
    elif mutation == "change-section":
        workbook["Individ årsoberoende"]["A30"] = "Ny tabell"
    elif mutation == "bad-period":
        workbook["Individ"]["C600"] = "1990 till 2024"
    else:
        workbook["Individ"]["G3"] = "Ny kolumn"
    workbook.save(path)
    workbook.close()

    with pytest.raises(LisaWorkbookError, match=expected):
        read_lisa_source(path, _revision(path))


@pytest.mark.parametrize(
    ("cell", "value", "expected"),
    (
        ("A322", "Other", r"Individ!A323"),
        ("B323", "changed continuation", r"Individ!A323:F323"),
        ("B816", "changed footer", r"Individ!A816:F816"),
        ("B324", "arbitrary unkeyed text", r"Individ!A324:F324"),
    ),
)
def test_lisa_reader_accepts_only_observed_unkeyed_context_rows(
    tmp_path: Path, cell: str, value: str, expected: str
) -> None:
    path = write_lisa_workbook(tmp_path / "lisa.xlsx")
    workbook = load_workbook(path)
    workbook["Individ"][cell] = value
    workbook.save(path)
    workbook.close()

    with pytest.raises(LisaWorkbookError, match=expected):
        read_lisa_source(path, _revision(path))


def test_lisa_reader_loads_the_small_selected_workbook_in_memory(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    path = write_lisa_workbook(tmp_path / "lisa.xlsx")
    read_only_values: list[bool | None] = []
    original = lisa_module.load_workbook

    def tracked_load_workbook(*args: Any, **kwargs: Any) -> Any:
        value = kwargs.get("read_only")
        assert value is None or isinstance(value, bool)
        read_only_values.append(value)
        return original(*args, **kwargs)

    monkeypatch.setattr(lisa_module, "load_workbook", tracked_load_workbook)

    read_lisa_source(path, _revision(path))

    assert read_only_values == [False]


def test_source_fields_distinguish_missing_unknown_negative_and_sensitivity() -> None:
    missing = SourceFields()
    unknown = SourceField(status="unknown", raw_value="")
    negative = SourceField(status="negative", raw_value="Nej")

    assert missing.column_name is None
    assert unknown.status == "unknown"
    assert negative.status == "negative"
    assert SourceFields(availability=negative).availability == negative
    sensitivity = SourceFields(sensitivity=value_field(False)).sensitivity
    assert sensitivity is not None
    assert sensitivity.status == "value"
    with pytest.raises(ValidationError, match="negative is supported only"):
        SourceFields(sensitivity=negative)
    assert SourceFields(
        sensitivity=value_field("conditional", raw="I vissa fall")
    ).sensitivity == SourceField(
        status="value", value="conditional", raw_value="I vissa fall"
    )
    with pytest.raises(ValidationError, match="use negative status"):
        SourceFields(availability=value_field(False))


def test_prepare_cli_requires_complete_lisa_selection(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    def reject_prepare(*_args: object, **_kwargs: object) -> None:
        raise AssertionError("incomplete LISA selection reached preparation")

    monkeypatch.setattr(cli_module, "prepare_input_bundle", reject_prepare)
    args = cli_module._build_parser().parse_args(
        [
            "prepare-input-bundle",
            "--input-dir",
            "inputs",
            "--scb-snapshot",
            "snapshot",
            "--scb-input-commit",
            "a" * 40,
            "--scb-manifest-sha256",
            "b" * 64,
            "--output-dir",
            "candidate",
            "--lisa-workbook",
            "lisa.xlsx",
        ]
    )

    with pytest.raises(RegMetaError) as exc_info:
        cli_module._cmd_prepare_input_bundle(args)
    assert exc_info.value.code == "lisa_source_selection_incomplete"


@pytest.mark.parametrize("handler_error", (False, True), ids=("success", "error"))
def test_prepare_cli_never_writes_output_over_selected_external_lisa(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
    capsys: pytest.CaptureFixture[str],
    handler_error: bool,
) -> None:
    input_dir = tmp_path / "source"
    write_scb_input(input_dir)
    snapshot = write_scb_snapshot(tmp_path / "accepted", input_dir / "SCB")
    workbook = write_lisa_workbook(tmp_path / "official-source" / "lisa.xlsx")
    source_bytes = workbook.read_bytes()
    handler_called = False

    def prospective_handler(_args: object) -> tuple[dict[str, object], int]:
        nonlocal handler_called
        handler_called = True
        if handler_error:
            raise SnapshotError("prospective handler failure")
        return {"data": {"status": "unexpected"}}, 0

    monkeypatch.setitem(
        cli_module.COMMAND_DISPATCH, "prepare-input-bundle", prospective_handler
    )
    exit_code = cli_module.run(
        [
            "--output",
            str(workbook),
            "prepare-input-bundle",
            "--input-dir",
            str(input_dir),
            "--scb-snapshot",
            str(snapshot.path),
            "--scb-input-commit",
            snapshot.input_commit,
            "--scb-manifest-sha256",
            snapshot.manifest_sha256,
            "--output-dir",
            str(snapshot.path.parent / "candidate"),
            "--lisa-workbook",
            str(workbook),
            "--lisa-revision",
            "2024-2025",
            "--lisa-workbook-sha256",
            hashlib.sha256(source_bytes).hexdigest(),
        ]
    )

    assert exit_code == EXIT_CONFIG
    assert not handler_called
    error = json.loads(capsys.readouterr().out)["error"]
    assert error["code"] == "catalog_input_output_conflict"
    assert workbook.read_bytes() == source_bytes


def test_inspection_cli_rejects_bundle_where_lisa_was_not_selected(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
    capsys: pytest.CaptureFixture[str],
) -> None:
    input_dir = tmp_path / "source"
    write_scb_input(input_dir)
    selection = write_input_bundle(tmp_path / "accepted", input_dir)
    monkeypatch.setattr(cli_module, "source_interpreter_commit", lambda: "c" * 40)

    exit_code = cli_module.run(
        [
            "inspect-source-records",
            "--input-bundle",
            str(selection.path),
            "--input-commit",
            selection.input_commit,
            "--input-manifest-sha256",
            selection.manifest_sha256,
        ]
    )

    assert exit_code != 0
    error = json.loads(capsys.readouterr().out)["error"]
    assert error["code"] == "source_record_inspection_invalid"
    assert "was not selected" in error["message"]
    assert "explicitly captures" in error["remediation"]


def test_raw_scb_reader_preserves_native_instances_fields_and_period_limits(
    tmp_path: Path,
) -> None:
    rows = [
        _var_row(
            colname="AmPolTyp",
            cvid=1,
            var_id=31619,
            year="2020",
            regver_id=200,
            register=("LISA", 34, 153),
            vardef="Definition 2020",
            varopdef="Operational 2020",
        ),
        _var_row(
            colname="",
            cvid=2,
            var_id=31619,
            year="2021",
            regver_id=201,
            register=("LISA", 34, 153),
            data_type="text",
            data_length="3",
        ),
        _var_row(
            colname="Pooled",
            cvid=3,
            var_id=3,
            year="2018",
            versionname="2018-2019",
            regver_id=202,
            register=("LISA", 34, 1335),
        ),
        _var_row(
            colname="UnknownPeriod",
            cvid=4,
            var_id=4,
            year="2020",
            versionname="okänd utgåva",
            regver_id=203,
            register=("LISA", 34, 153),
        ),
    ]
    scb_dir = write_scb_input(tmp_path / "source", registerinformation_rows=rows)
    selection = write_scb_snapshot(tmp_path / "accepted", scb_dir)
    snapshot = open_scb_snapshot(selection)
    item = next(
        item
        for item in snapshot.manifest.files
        if item.name == "Registerinformation.csv"
    )
    assert item.raw_size is not None and item.raw_sha256 is not None
    revision = SourceRevision.create(
        dataset="scb-registerinformation",
        publisher="SCB",
        purpose="fixture raw rows",
        upstream_revision=snapshot.manifest.edition,
        artifact_path="Registerinformation.csv",
        artifact_size=item.raw_size,
        artifact_sha256=item.raw_sha256,
    )

    records, issues = read_scb_lisa_records(snapshot, revision)

    assert len(records) == 4
    blank = next(record for record in records if record.subject.native.member_id == 2)
    assert blank.fields.column_name == SourceField(status="unknown", raw_value="")
    assert blank.fields.availability == value_field(True)
    assert blank.subject.native.variable_id == 31619
    assert blank.subject.native.register_variant_id == 153
    assert _field_text(blank, "data_type") == "text"
    assert _field_text(blank, "data_length") == "3"
    pooled = next(record for record in records if record.subject.native.member_id == 3)
    assert pooled.edition_scope == TemporalScope(kind="pooled", label="2018-2019")
    unknown = next(record for record in records if record.subject.native.member_id == 4)
    assert unknown.edition_scope.kind == "unknown"
    assert [(issue.kind, issue.record_id) for issue in issues] == [
        ("pooled_period", pooled.record_id),
        ("unparseable_period", unknown.record_id),
    ]


@pytest.mark.parametrize(
    ("version_name", "expected_kind", "expected_issue"),
    (
        ("1990, 2000", "unknown", "unparseable_period"),
        ("LISA 2011 och 2019", "unknown", "unparseable_period"),
        ("1990-2000", "pooled", "pooled_period"),
        ("LISA 2011", "intervals", None),
    ),
)
def test_scb_scope_rejects_multi_year_tokens_only_on_single_claim_fallback(
    version_name: str, expected_kind: str, expected_issue: str | None
) -> None:
    edition_scope, reference_scope, issue = _scopes(34, version_name)

    assert edition_scope.kind == expected_kind
    assert reference_scope.kind == expected_kind
    assert issue == expected_issue


def test_compact_comparison_retains_witnesses_conflicts_ids_and_spelling() -> None:
    revision = SourceRevision.create(
        dataset="fixture",
        publisher="SCB",
        purpose="comparison",
        upstream_revision="1",
        artifact_path="fixture",
        artifact_size=0,
        artifact_sha256="0" * 64,
    )
    workbook = (
        _record(
            revision,
            key="workbook",
            column="AmPolTyp",
            year=2020,
            workbook_table="individual",
        ),
        _record(
            revision,
            key="negative",
            column="Negative",
            year=2020,
            availability=SourceField(status="negative"),
            workbook_table="individual",
        ),
        _record(
            revision,
            key="suffix",
            column="Thing_LISA",
            year=2020,
            workbook_table="individual",
        ),
        _record(
            revision,
            key="unknown-availability",
            column="Unknown",
            year=2020,
            availability=SourceField(status="unknown", raw_value="Okänt"),
            workbook_table="individual",
        ),
    )
    scb = (
        _record(
            revision,
            key="named-a",
            column="AmPolTyp",
            year=2020,
            variable_id=31619,
            member_id=10,
        ),
        _record(
            revision,
            key="named-b",
            column="AmPolTyp",
            year=2020,
            variant_id=1335,
            variable_id=999,
            member_id=11,
        ),
        _record(
            revision,
            key="positive",
            column="Negative",
            year=2020,
            member_id=12,
        ),
        _record(
            revision,
            key="distinct-suffix",
            column="Thing_MiDAS",
            year=2020,
            member_id=13,
        ),
        _record(
            revision,
            key="known-availability",
            column="Unknown",
            year=2020,
            member_id=14,
        ),
    )

    outcomes = compare_availability_records(workbook, scb)
    by_column = {
        outcome.column_name: outcome
        for outcome in outcomes
        if outcome.workbook_record_ids
    }

    assert by_column["AmPolTyp"].status == "agreement"
    assert by_column["AmPolTyp"].variable_ids == (31619,)
    assert len(by_column["AmPolTyp"].scb_record_ids) == 1
    assert by_column["AmPolTyp"].assumption_ids == (
        "lisa-individual-2010-2024-to-scb-variant-153",
    )
    assert any(
        outcome.status == "source_only_observation" and outcome.variable_ids == (999,)
        for outcome in outcomes
    )
    assert by_column["Negative"].status == "conflict"
    assert by_column["Unknown"].status == "unknown_applicability"
    assert by_column["Thing_LISA"].status == "unobserved_counterpart"
    assert not by_column["Thing_LISA"].scb_record_ids
    assert all(outcome.assumption_ids for outcome in outcomes)


@pytest.mark.parametrize(
    (
        "scb_column",
        "scb_variant",
        "expected_status",
        "edition_present",
        "expected_witness",
    ),
    (
        ("X", 153, "unknown_applicability", None, True),
        ("Other", 153, "missing_edition", False, False),
        ("X", 152, "agreement", True, True),
        ("Other", 152, "unobserved_counterpart", True, False),
    ),
    ids=(
        "cross-variant-exact-is-unknown",
        "cross-variant-edition-is-missing",
        "scoped-exact-agrees",
        "scoped-edition-is-present",
    ),
)
def test_company_comparison_uses_only_its_finite_native_variant_assumption(
    scb_column: str,
    scb_variant: int,
    expected_status: str,
    edition_present: bool | None,
    expected_witness: bool,
) -> None:
    revision = SourceRevision.create(
        dataset="fixture",
        publisher="SCB",
        purpose="scoped comparison",
        upstream_revision="1",
        artifact_path="fixture",
        artifact_size=0,
        artifact_sha256="0" * 64,
    )
    workbook = _record(
        revision,
        key="company-workbook",
        column="X",
        year=2020,
        workbook_table="company",
        population_name=inspection_module._ENTERPRISE_POPULATION,
    )
    scb_record = _record(
        revision,
        key="scb",
        column=scb_column,
        year=2020,
        variant_id=scb_variant,
        population_name=("individuals 15+" if scb_variant == 153 else "companies"),
    )

    outcomes = compare_availability_records((workbook,), (scb_record,))
    compared = next(outcome for outcome in outcomes if outcome.workbook_record_ids)

    assert compared.status == expected_status
    assert compared.scb_edition_present is edition_present
    assert bool(compared.scb_record_ids) is expected_witness
    assert compared.assumption_ids[0] == ("lisa-company-1990-2024-to-scb-variant-152")
    if scb_variant == 153 and scb_column == "X":
        assert compared.assumption_ids == (
            "lisa-company-1990-2024-to-scb-variant-152",
            "unestablished-workbook-to-scb-variant-scope",
        )
        assert any(
            outcome.status == "source_only_observation"
            and outcome.scb_record_ids == (scb_record.record_id,)
            for outcome in outcomes
        )


@pytest.mark.parametrize(
    ("anchor_year", "blank_year", "expected_target_count"),
    (
        (2019, 2020, 1),
        (2020, 2021, 0),
    ),
    ids=("anchor-outside-workbook-target-inside", "target-outside-workbook"),
)
def test_filtered_preview_uses_finite_anchors_and_workbook_scoped_targets(
    tmp_path: Path,
    anchor_year: int,
    blank_year: int,
    expected_target_count: int,
) -> None:
    input_dir = tmp_path / "source"
    rows = [
        _var_row(
            colname="X",
            cvid=1,
            var_id=9,
            year=str(anchor_year),
            regver_id=200,
            register=("LISA", 34, 152),
        ),
        _var_row(
            colname="",
            cvid=2,
            var_id=9,
            year=str(blank_year),
            regver_id=201,
            register=("LISA", 34, 152),
        ),
    ]
    write_scb_input(input_dir, registerinformation_rows=rows)
    workbook = write_lisa_workbook(input_dir / "docs" / "lisa.xlsx")
    source = load_workbook(workbook)
    source["Företag"]["A9"] = "X"
    source["Företag"]["C9"] = 2020
    source.save(workbook)
    source.close()
    selection = write_input_bundle(
        tmp_path / "accepted",
        input_dir,
        lisa_workbook=LisaWorkbookSelection(
            path=workbook,
            upstream_revision="2024-2025",
            sha256=hashlib.sha256(workbook.read_bytes()).hexdigest(),
        ),
    )
    bundle = open_input_bundle(selection)

    full = inspect_bundle_source_records(bundle, code_commit="c" * 40)
    filtered = inspect_bundle_source_records(
        bundle, code_commit="c" * 40, exact_column="X"
    )

    assert filtered.target_preview is not None
    scb_by_member = {
        record.subject.native.member_id: record
        for record in filtered.source_records
        if record.subject.native.member_id is not None
    }
    anchor = scb_by_member[1]
    blank = scb_by_member[2]
    assert anchor.record_id in filtered.target_preview.existing_named_record_ids
    assert filtered.target_preview.candidate_variable_ids == (9,)
    assert filtered.target_preview.candidate_variant_ids == (152,)
    assert filtered.target_preview.expected_source_target_count == expected_target_count
    assert (blank.record_id in filtered.target_preview.target_record_ids) is bool(
        expected_target_count
    )
    for report in (full, filtered):
        assert blank.record_id in {record.record_id for record in report.source_records}
        assert any(
            blank.record_id in outcome.scb_record_ids
            and outcome.status
            == (
                "unknown_spelling" if expected_target_count else "unknown_applicability"
            )
            for outcome in report.comparison_outcomes
        )


def test_competing_finite_native_spelling_makes_blank_target_ambiguous(
    tmp_path: Path,
) -> None:
    input_dir = tmp_path / "source"
    rows = [
        _var_row(
            colname=column,
            cvid=member_id,
            var_id=9,
            year=str(year),
            regver_id=200 + member_id,
            register=("LISA", 34, 152),
        )
        for column, year, member_id in (
            ("X", 2018, 1),
            ("Y", 2019, 2),
            ("", 2020, 3),
        )
    ]
    write_scb_input(input_dir, registerinformation_rows=rows)
    workbook = write_lisa_workbook(input_dir / "docs" / "lisa.xlsx")
    source = load_workbook(workbook)
    source["Företag"]["A9"] = "X"
    source["Företag"]["C9"] = 2020
    source.save(workbook)
    source.close()
    selection = write_input_bundle(
        tmp_path / "accepted",
        input_dir,
        lisa_workbook=LisaWorkbookSelection(
            path=workbook,
            upstream_revision="2024-2025",
            sha256=hashlib.sha256(workbook.read_bytes()).hexdigest(),
        ),
    )
    bundle = open_input_bundle(selection)

    full = inspect_bundle_source_records(bundle, code_commit="c" * 40)
    filtered = inspect_bundle_source_records(
        bundle, code_commit="c" * 40, exact_column="X"
    )

    assert filtered.target_preview is not None
    assert filtered.target_preview.expected_source_target_count == 0
    assert filtered.target_preview.target_record_ids == ()
    scb_by_member = {
        record.subject.native.member_id: record
        for record in filtered.source_records
        if record.subject.native.member_id is not None
    }
    assert set(scb_by_member) == {1, 2, 3}
    competing = scb_by_member[2]
    blank = scb_by_member[3]
    for report in (full, filtered):
        ambiguous = next(
            outcome
            for outcome in report.comparison_outcomes
            if outcome.column_name == "X"
            and outcome.workbook_record_ids
            and outcome.status == "ambiguous_match"
        )
        assert set(ambiguous.scb_record_ids) == {
            competing.record_id,
            blank.record_id,
        }
        assert not any(
            outcome.status == "unknown_spelling"
            and blank.record_id in outcome.scb_record_ids
            for outcome in report.comparison_outcomes
        )


def test_filtered_report_retains_unparseable_period_as_incomplete(
    tmp_path: Path,
) -> None:
    input_dir = tmp_path / "source"
    rows = [
        _var_row(
            colname="AmPolTyp",
            cvid=1,
            var_id=31619,
            year="2020",
            versionname="okänd utgåva",
            regver_id=200,
            register=("LISA", 34, 153),
        )
    ]
    write_scb_input(input_dir, registerinformation_rows=rows)
    workbook = write_lisa_workbook(input_dir / "docs" / "lisa.xlsx")
    selection = write_input_bundle(
        tmp_path / "accepted",
        input_dir,
        lisa_workbook=LisaWorkbookSelection(
            path=workbook,
            upstream_revision="2024-2025",
            sha256=hashlib.sha256(workbook.read_bytes()).hexdigest(),
        ),
    )

    report = inspect_bundle_source_records(
        open_input_bundle(selection),
        code_commit="c" * 40,
        exact_column="AmPolTyp",
    )

    assert report.complete is False
    assert len(report.interpretation_issues) == 1
    assert report.interpretation_issues[0].kind == "unparseable_period"
    assert report.interpretation_issues[0].incomplete is True
    assert any(
        outcome.status == "unknown_applicability"
        and outcome.scb_edition_present is None
        and outcome.scb_record_ids
        for outcome in report.comparison_outcomes
    )


@pytest.mark.parametrize(
    ("version_name", "scope_kind", "issue_kind", "expected_complete"),
    (
        ("okänd utgåva", "unknown", "unparseable_period", False),
        ("2020-2021", "pooled", "pooled_period", True),
    ),
    ids=("unparseable", "pooled"),
)
def test_filtered_report_retains_unscoped_same_variable_witness_but_not_target(
    tmp_path: Path,
    version_name: str,
    scope_kind: str,
    issue_kind: str,
    expected_complete: bool,
) -> None:
    input_dir = tmp_path / "source"
    rows = [
        _var_row(
            colname="AmPolTyp",
            cvid=1,
            var_id=31619,
            year="2020",
            regver_id=200,
            register=("LISA", 34, 153),
        ),
        _var_row(
            colname="",
            cvid=2,
            var_id=31619,
            year="2021",
            versionname=version_name,
            regver_id=201,
            register=("LISA", 34, 153),
        ),
    ]
    write_scb_input(input_dir, registerinformation_rows=rows)
    workbook = write_lisa_workbook(input_dir / "docs" / "lisa.xlsx")
    selection = write_input_bundle(
        tmp_path / "accepted",
        input_dir,
        lisa_workbook=LisaWorkbookSelection(
            path=workbook,
            upstream_revision="2024-2025",
            sha256=hashlib.sha256(workbook.read_bytes()).hexdigest(),
        ),
    )
    bundle = open_input_bundle(selection)

    full = inspect_bundle_source_records(bundle, code_commit="c" * 40)
    filtered = inspect_bundle_source_records(
        bundle, code_commit="c" * 40, exact_column="AmPolTyp"
    )

    assert full.complete is expected_complete
    assert filtered.complete is expected_complete
    assert filtered.target_preview is not None
    unscoped = next(
        record
        for record in filtered.source_records
        if record.subject.native.member_id == 2
    )
    assert unscoped.edition_scope.kind == scope_kind
    assert unscoped.record_id not in filtered.target_preview.target_record_ids
    assert filtered.target_preview.expected_source_target_count == 0
    assert any(
        issue.source_record_ids == (unscoped.record_id,) and issue.kind == issue_kind
        for issue in filtered.interpretation_issues
    )
    assert any(
        outcome.status == "unknown_applicability"
        and outcome.scb_record_ids == (unscoped.record_id,)
        and outcome.edition_scope.kind == scope_kind
        for outcome in filtered.comparison_outcomes
    )


@pytest.mark.parametrize(
    ("version_name", "issue_kind", "expected_complete"),
    (
        ("okänd utgåva", "unparseable_period", False),
        ("2018-2019", "pooled_period", True),
    ),
    ids=("unparseable", "pooled"),
)
def test_source_only_unknown_period_survives_full_and_filtered_inspection(
    tmp_path: Path,
    version_name: str,
    issue_kind: str,
    expected_complete: bool,
) -> None:
    input_dir = tmp_path / "source"
    rows = [
        _var_row(
            colname="OnlyScb",
            cvid=1,
            var_id=9,
            year="2019",
            regver_id=200,
            register=("LISA", 34, 152),
        ),
        _var_row(
            colname="",
            cvid=2,
            var_id=9,
            year="2020",
            versionname=version_name,
            regver_id=201,
            register=("LISA", 34, 152),
        ),
    ]
    write_scb_input(input_dir, registerinformation_rows=rows)
    workbook = write_lisa_workbook(input_dir / "docs" / "lisa.xlsx")
    selection = write_input_bundle(
        tmp_path / "accepted",
        input_dir,
        lisa_workbook=LisaWorkbookSelection(
            path=workbook,
            upstream_revision="2024-2025",
            sha256=hashlib.sha256(workbook.read_bytes()).hexdigest(),
        ),
    )
    bundle = open_input_bundle(selection)

    full = inspect_bundle_source_records(bundle, code_commit="c" * 40)
    filtered = inspect_bundle_source_records(
        bundle, code_commit="c" * 40, exact_column="OnlyScb"
    )

    named = next(
        record
        for record in filtered.source_records
        if record.subject.native.member_id == 1
    )
    unknown = next(
        record
        for record in filtered.source_records
        if record.subject.native.member_id == 2
    )
    assert filtered.complete is expected_complete
    assert filtered.target_preview is not None
    assert filtered.target_preview.expected_source_target_count == 0
    assert filtered.target_preview.target_record_ids == ()
    assert filtered.target_preview.existing_named_record_ids == (named.record_id,)
    assert filtered.target_preview.assumption_ids == (
        "unestablished-workbook-to-scb-variant-scope",
    )
    for report in (full, filtered):
        assert report.complete is expected_complete
        assert {named.record_id, unknown.record_id} <= {
            record.record_id for record in report.source_records
        }
        assert any(
            outcome.status == "unknown_applicability"
            and outcome.workbook_record_ids == ()
            and outcome.scb_record_ids == (unknown.record_id,)
            and outcome.assumption_ids
            == (
                "unestablished-workbook-to-scb-variant-scope",
                "same-native-variable-spelling-continuity",
            )
            for outcome in report.comparison_outcomes
        )
        assert any(
            issue.kind == issue_kind
            and issue.source_record_ids == (unknown.record_id,)
            for issue in report.interpretation_issues
        )


def test_year_independent_declarations_without_counterparts_have_zero_target_preview(
    tmp_path: Path,
) -> None:
    input_dir = tmp_path / "source"
    write_scb_input(
        input_dir,
        registerinformation_rows=[
            _var_row(
                colname="Other",
                cvid=1,
                var_id=1,
                year="2020",
                regver_id=200,
                register=("LISA", 34, 153),
            )
        ],
    )
    workbook = write_lisa_workbook(input_dir / "docs" / "lisa.xlsx")
    source = load_workbook(workbook)
    for sheet_name, row, values in (
        (
            "Individ",
            42,
            (
                "Inv_UtvGrEg5",
                "In- och utvandring efter grund för egen uppgift",
                None,
                "LISA",
                "Nej",
                "RTB",
            ),
        ),
        (
            "Individ årsoberoende",
            43,
            (
                "Inv_UtvGrEg5",
                "In- och utvandring efter grund för egen uppgift",
                "LISA",
                "RTB",
            ),
        ),
    ):
        sheet = source[sheet_name]
        for column, value in enumerate(values, start=1):
            sheet.cell(row, column, value)
    source.save(workbook)
    source.close()
    selection = write_input_bundle(
        tmp_path / "accepted",
        input_dir,
        lisa_workbook=LisaWorkbookSelection(
            path=workbook,
            upstream_revision="2024-2025",
            sha256=hashlib.sha256(workbook.read_bytes()).hexdigest(),
        ),
    )
    bundle = open_input_bundle(selection)

    full = inspect_bundle_source_records(bundle, code_commit="c" * 40)
    filtered = inspect_bundle_source_records(
        bundle, code_commit="c" * 40, exact_column="Inv_UtvGrEg5"
    )

    assert filtered.summary.workbook_selected_occurrences == 2
    assert filtered.summary.workbook_year_independent_occurrences == 2
    assert filtered.target_preview is not None
    assert filtered.target_preview.documented_edition_scope is None
    assert filtered.target_preview.expected_source_target_count == 0
    assert filtered.target_preview.target_record_ids == ()
    assert filtered.target_preview.assumption_ids == (
        "unestablished-workbook-to-scb-variant-scope",
    )
    assert {
        (record.locator.physical_table, record.locator.physical_record)
        for record in filtered.source_records
    } == {
        ("Individ", "row:42"),
        ("Individ årsoberoende", "row:43"),
    }
    for report in (full, filtered):
        matching = [
            outcome
            for outcome in report.comparison_outcomes
            if outcome.column_name == "Inv_UtvGrEg5"
        ]
        assert len(matching) == 2
        assert all(outcome.status == "unknown_applicability" for outcome in matching)
        assert all(
            outcome.assumption_ids == ("unestablished-workbook-to-scb-variant-scope",)
            for outcome in matching
        )


def test_ampoltyp_preview_keeps_nine_targets_and_a_separate_2024_gap(
    tmp_path: Path,
) -> None:
    input_dir = tmp_path / "source"
    rows = [
        _var_row(
            colname="" if 2011 <= year <= 2019 else "AmPolTyp",
            cvid=year,
            var_id=31619,
            year=str(year),
            regver_id=year,
            register=("LISA", 34, 1335 if year <= 2009 else 153),
        )
        for year in range(1990, 2024)
    ]
    write_scb_input(input_dir, registerinformation_rows=rows)
    workbook = write_lisa_workbook(input_dir / "docs" / "lisa.xlsx")
    selection = write_input_bundle(
        tmp_path / "accepted",
        input_dir,
        lisa_workbook=LisaWorkbookSelection(
            path=workbook,
            upstream_revision="2024-2025",
            sha256=hashlib.sha256(workbook.read_bytes()).hexdigest(),
        ),
    )

    report = inspect_bundle_source_records(
        open_input_bundle(selection),
        code_commit="c" * 40,
        exact_column="AmPolTyp",
    )

    assert report.target_preview is not None
    assert report.target_preview.expected_source_target_count == 9
    assert len(report.target_preview.target_record_ids) == 9
    assert len(report.target_preview.existing_named_record_ids) == 25
    assert report.target_preview.candidate_variable_ids == (31619,)
    assert report.target_preview.candidate_variant_ids == (153, 1335)
    assert report.target_preview.unobserved_documented_editions == (2024,)
    target_records = {
        record.record_id: record
        for record in report.source_records
        if record.record_id in report.target_preview.target_record_ids
    }
    assert {
        record.subject.native.member_id for record in target_records.values()
    } == set(range(2011, 2020))
    assert all(
        record.fields.column_name == SourceField(status="unknown", raw_value="")
        for record in target_records.values()
    )
    assert any(
        outcome.status == "missing_edition"
        and outcome.edition_scope.intervals
        == (ScopeInterval(start="2024", end="2024"),)
        and not outcome.scb_record_ids
        for outcome in report.comparison_outcomes
    )


def test_column_filter_retains_source_only_exact_observations(
    tmp_path: Path,
) -> None:
    input_dir = tmp_path / "source"
    rows = [
        _var_row(
            colname="AmPolTyp",
            cvid=1,
            var_id=31619,
            year="2020",
            regver_id=200,
            register=("LISA", 34, 153),
        ),
        _var_row(
            colname="AmPolTyp",
            cvid=2,
            var_id=31619,
            year="2025",
            regver_id=201,
            register=("LISA", 34, 153),
        ),
        _var_row(
            colname="OnlyScb",
            cvid=3,
            var_id=40000,
            year="2026",
            regver_id=202,
            register=("LISA", 34, 153),
        ),
    ]
    write_scb_input(input_dir, registerinformation_rows=rows)
    workbook = write_lisa_workbook(input_dir / "docs" / "lisa.xlsx")
    selection = write_input_bundle(
        tmp_path / "accepted",
        input_dir,
        lisa_workbook=LisaWorkbookSelection(
            path=workbook,
            upstream_revision="2024-2025",
            sha256=hashlib.sha256(workbook.read_bytes()).hexdigest(),
        ),
    )
    bundle = open_input_bundle(selection)

    ampoltyp = inspect_bundle_source_records(
        bundle, code_commit="c" * 40, exact_column="AmPolTyp"
    )
    assert {record.subject.native.member_id for record in ampoltyp.source_records} >= {
        1,
        2,
    }
    assert any(
        outcome.status == "source_only_observation"
        and outcome.edition_scope.intervals
        == (ScopeInterval(start="2025", end="2025"),)
        and not outcome.workbook_record_ids
        for outcome in ampoltyp.comparison_outcomes
    )

    scb_only = inspect_bundle_source_records(
        bundle, code_commit="c" * 40, exact_column="OnlyScb"
    )
    assert scb_only.summary.workbook_selected_occurrences == 0
    assert set(scb_only.workbook_context) >= _QUALIFICATION_CONTEXT
    assert all(
        scb_only.workbook_context.count(item) == 1 for item in _QUALIFICATION_CONTEXT
    )
    assert [record.subject.native.member_id for record in scb_only.source_records] == [
        3
    ]
    assert all(
        item not in record.context
        for record in scb_only.source_records
        for item in _QUALIFICATION_CONTEXT
    )
    assert len(scb_only.comparison_outcomes) == 1
    assert scb_only.comparison_outcomes[0].status == "source_only_observation"
    assert scb_only.comparison_outcomes[0].workbook_record_ids == ()


def test_pinned_bundle_cli_reports_deterministic_source_targets_without_cold_values(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
    capsys: pytest.CaptureFixture[str],
) -> None:
    rows = [
        _var_row(
            colname="AmPolTyp",
            cvid=1,
            var_id=31619,
            year="2020",
            regver_id=200,
            register=("LISA", 34, 153),
            vardef="Kept definition",
        ),
        _var_row(
            colname="",
            cvid=2,
            var_id=31619,
            year="2021",
            regver_id=201,
            register=("LISA", 34, 153),
            data_type="text",
        ),
        _var_row(
            colname="ampoltyp",
            cvid=3,
            var_id=99,
            year="2022",
            regver_id=202,
            register=("LISA", 34, 1335),
        ),
        _var_row(
            colname="Other",
            cvid=4,
            var_id=4,
            year="2023",
            regver_id=203,
            register=("LISA", 34, 153),
        ),
    ]
    input_dir = tmp_path / "source"
    write_scb_input(input_dir, registerinformation_rows=rows)
    workbook = write_lisa_workbook(input_dir / "docs" / "lisa.xlsx")
    workbook_sha = hashlib.sha256(workbook.read_bytes()).hexdigest()
    selection = write_input_bundle(
        tmp_path / "accepted",
        input_dir,
        lisa_workbook=LisaWorkbookSelection(
            path=workbook,
            upstream_revision="2024-2025",
            sha256=workbook_sha,
        ),
    )
    bundle = open_input_bundle(selection)
    tracked_before = (bundle.root / "catalog-bundle.json").read_bytes()
    workbook_before = (
        bundle.root / "supplemental" / LISA_DATASET_ID / "source.xlsx"
    ).read_bytes()
    sparsify_scb_values(selection)

    opened: list[str] = []
    original_open_csv = ScbSnapshotReader.open_csv

    @contextmanager
    def guarded_open_csv(
        reader: ScbSnapshotReader, name: str
    ) -> Iterator[tuple[list[str | None], Iterator[list[str | None]]]]:
        opened.append(name)
        assert name == "Registerinformation.csv"
        with original_open_csv(reader, name) as value:
            yield value

    monkeypatch.setattr(ScbSnapshotReader, "open_csv", guarded_open_csv)
    monkeypatch.setattr(cli_module, "source_interpreter_commit", lambda: "c" * 40)
    outputs = [tmp_path / "report-a.json", tmp_path / "report-b.json"]
    reports: list[dict[str, Any]] = []
    for output in outputs:
        exit_code = cli_module.run(
            [
                "--output",
                str(output),
                "inspect-source-records",
                "--input-bundle",
                str(selection.path),
                "--input-commit",
                selection.input_commit,
                "--input-manifest-sha256",
                selection.manifest_sha256,
                "--column",
                "AmPolTyp",
            ]
        )
        assert exit_code == 0
        assert capsys.readouterr().out == ""
        reports.append(json.loads(output.read_text(encoding="utf-8")))

    assert reports[0] == reports[1]
    report = reports[0]
    assert report["schema_version"] == 2
    assert report["diagnostic_only"] is True
    assert report["preview_level"] == "source_target_only"
    assert report["complete"] is True
    assert report["pins"] == {
        "bundle_id": "scb-mikrometadata",
        "bundle_manifest_sha256": selection.manifest_sha256,
        "code_commit": "c" * 40,
        "input_repository_commit": selection.input_commit,
        "interpretation_id": "scb-lisa-source-record-inspection-v2",
        "scb_snapshot_manifest_sha256": bundle.manifest.scb_manifest_sha256,
    }
    assert report["scope"]["selected_sources"] == [
        LISA_DATASET_ID,
        "scb-registerinformation",
    ]
    assert "SCB Vardemangder value streams" in report["scope"]["excluded_inputs"]
    assert {revision["artifact_sha256"] for revision in report["source_revisions"]} == {
        workbook_sha,
        bundle.snapshot.raw_sha256("Registerinformation.csv"),
    }
    assert report["target_preview"]["expected_source_target_count"] == 1
    assert report["target_preview"]["candidate_variable_ids"] == [31619]
    assert report["target_preview"]["candidate_variant_ids"] == [153]
    assert set(report["target_preview"]["assumption_ids"]) >= {
        "lisa-individual-1990-2009-to-scb-variant-1335",
        "lisa-individual-2010-2024-to-scb-variant-153",
        "same-native-variable-spelling-continuity",
    }
    assert 2024 in report["target_preview"]["unobserved_documented_editions"]
    assert len(report["target_preview"]["target_record_ids"]) == 1
    assert report["summary"]["outcome_counts"]["unknown_spelling"] == 1
    assert report["summary"]["outcome_counts"]["unknown_applicability"] >= 1
    assert report["summary"]["outcome_counts"]["unobserved_counterpart"] == 1
    assert report["summary"]["outcome_counts"]["missing_edition"] >= 1
    agreement = next(
        outcome
        for outcome in report["comparison_outcomes"]
        if outcome["status"] == "agreement"
    )
    assert len(agreement["workbook_record_ids"]) == 1
    assert len(agreement["scb_record_ids"]) == 1
    assert agreement["assumption_ids"] == [
        "lisa-individual-2010-2024-to-scb-variant-153"
    ]
    cross_variant_casefold = next(
        outcome
        for outcome in report["comparison_outcomes"]
        if outcome["status"] == "unknown_applicability"
        and outcome["register_variant_ids"] == [1335]
    )
    assert "casefold-spelling-candidate" in cross_variant_casefold["assumption_ids"]
    assert (
        "unestablished-workbook-to-scb-variant-scope"
        in cross_variant_casefold["assumption_ids"]
    )
    target_id = report["target_preview"]["target_record_ids"][0]
    target = next(
        record
        for record in report["source_records"]
        if record["record_id"] == target_id
    )
    assert target["subject"]["native"] == {
        "edition_id": 201,
        "member_id": 2,
        "register_id": 34,
        "register_variant_id": 153,
        "variable_id": 31619,
    }
    assert target["fields"]["column_name"] == {
        "raw_value": "",
        "status": "unknown",
    }
    assert target["fields"]["data_type"]["value"] == "text"
    assert (bundle.root / "catalog-bundle.json").read_bytes() == tracked_before
    assert (
        bundle.root / "supplemental" / LISA_DATASET_ID / "source.xlsx"
    ).read_bytes() == workbook_before
    semantic_payload = dict(report)
    semantic_sha256 = semantic_payload.pop("semantic_sha256")
    validated = SourceInspectionReport.model_validate_json(
        json.dumps(semantic_payload, ensure_ascii=False)
    )
    assert semantic_sha256 == report_semantic_sha256(validated)

    full_a = inspect_bundle_source_records(
        open_input_bundle(selection), code_commit="c" * 40
    )
    full_b = inspect_bundle_source_records(
        open_input_bundle(selection), code_commit="c" * 40
    )
    assert full_a == full_b
    assert full_a.target_preview is None
    assert full_a.summary.workbook_selected_occurrences == 9
    assert opened == ["Registerinformation.csv"] * 4


def test_source_interpreter_pin_rejects_a_loaded_dependency_from_another_checkout(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    other_repo = tmp_path / "other-checkout"
    other_repo.mkdir()
    subprocess.run(["git", "-C", str(other_repo), "init", "-q"], check=True)
    other_queries = other_repo / "queries.py"
    other_queries.write_text(
        "def extract_year(_value): return 2021\n", encoding="utf-8"
    )
    monkeypatch.setattr(
        inspection_module.reg_meta_queries, "__file__", str(other_queries)
    )

    with pytest.raises(SnapshotError, match="must come from the same Git checkout"):
        inspection_module.source_interpreter_commit()
