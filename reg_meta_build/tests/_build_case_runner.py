"""Run one `cases/build/` case: readable sources and curation through the real build.

`cases/build/README.md` is the case format. This module turns a case directory into
the boundary outputs its `expected.json` names: it materializes the source spec
(`_csv_fixtures` SCB rows, `_sos_fixtures` workbooks), prepares and accepts it
through the real input pipeline, renders the curation tree's authoring
placeholders from the accepted prepared records with the public capture helpers a
curator uses, runs `build_catalog`, and reads the result, the report ledger and the
built artifact through a fixed set of named projections.

Prepared inputs are cached by the content hash of their source spec, so cases that
share a source set prepare it once per session. An entry is written to a private
staging directory and renamed into place: concurrent xdist workers never write the
same path, and a failed preparation never leaves a partial entry. Entries
persist across sessions and worktrees in a generation of the shared fixture cache
(the reader cache's location), keyed by everything that shapes a prepared input.
"""

from __future__ import annotations

import functools
import hashlib
import json
import os
import re
import shutil
import sqlite3
import subprocess
import sys
import tempfile
import zipfile
from contextlib import closing
from dataclasses import dataclass
from datetime import datetime
from pathlib import Path
from typing import TYPE_CHECKING

import pytest
from _csv_fixtures import (
    replace_registerinformation_cell,
    summary_rows,
    var_row,
    write_scb_input,
)
from _pipeline_catalog_support import prepare_accepted, report_events
from _sos_fixtures import (
    SosRegister,
    SosSubset,
    SosVariable,
    write_sos_input,
)
from openpyxl import Workbook, load_workbook
from reader_artifacts import build_inputs_digest, generation_dir
from reg_meta_build.errors import EXIT_CONFIG, RegMetaError
from reg_meta_build.pipeline import build_catalog, check_curation
from reg_meta_build.prepared_catalog import (
    ReferenceEvidence,
    open_prepared_catalog_sources,
)
from reg_meta_build.source_coding import (
    coding_source_sha256,
    copied_coding_fingerprints,
)
from reg_meta_build.source_coordinates import native_variant_key
from reg_meta_build.source_curation import (
    acknowledgement_evidence_sha256,
    capture_expectations,
)
from reg_meta_build.source_evidence import canonical_sha256
from reg_meta_build.source_naming import authored_naming_id
from reg_meta_build.source_occurrences import source_occurrence
from reg_meta_build.source_records import SourceFields
from reg_meta_build.source_value_bindings import (
    ValueBindingResult,
    bind_code_lists,
    marker_binding_fingerprints,
    open_value_bindings,
)

if TYPE_CHECKING:
    from collections.abc import Callable

    from reg_meta_build.source_records import SourceRecord

CASES = Path(__file__).with_name("cases") / "build"
SOURCES = CASES / "_sources"


def case_dirs() -> list[Path]:
    """Every case directory, sorted; `_`-prefixed directories hold shared data."""
    return sorted(
        path
        for path in CASES.iterdir()
        if path.is_dir() and not path.name.startswith("_")
    )


def case_steps(case: Path) -> list[Path]:
    """A case is one step (its own `request.json`) or ordered step subdirectories."""
    if (case / "request.json").is_file():
        return [case]
    return sorted(path for path in case.iterdir() if path.is_dir())


# -- Sources -------------------------------------------------------------------


def source_spec(reference: str | dict) -> dict:
    """A source spec: inline, or the name of a shared `_sources/<name>.json`."""
    if isinstance(reference, dict):
        return reference
    return json.loads((SOURCES / f"{reference}.json").read_text(encoding="utf-8"))


def scb_row(row: dict) -> str:
    """One Registerinformation row: `var_row` arguments, then raw `cells` overrides."""
    text = var_row(
        **{
            key: tuple(value) if key == "register" else value
            for key, value in row.items()
            if key != "cells"
        }
    )
    for name, value in row.get("cells", {}).items():
        text = replace_registerinformation_cell(text, name, value)
    return text


def write_sources(spec: dict, source: Path) -> None:
    """Materialize a source spec as provider deliveries under ``source``."""
    scb = spec["scb"]
    rows = [scb_row(row) for row in scb["registerinformation"]]
    values = scb.get("vardemangder", [])
    # `"unika": null` delivers no Unika summary file at all.
    unika = scb.get("unika", summary_rows(rows))
    write_scb_input(
        source,
        registerinformation_rows=rows,
        vardemangder_rows=values,
        unika_rows=unika or [],
        # Default: every value item is valid over the whole fixture window.
        valid_dates_rows=scb.get(
            "valid_dates",
            sorted({f"{row.split('|')[-1]}|2000-01-01|2030-12-31" for row in values}),
        ),
        include=("registerinformation",)
        + (("unika",) if unika is not None else ())
        + (("vardemangder", "valid_dates") if values else ()),
    )
    if "join_keys" in scb:
        # ID-kolumner.xlsx: one worksheet of `Tabell | ID-kolumn | Beskrivning` rows.
        workbook = Workbook()
        sheet = workbook.active
        sheet.append(["Tabell", "ID-kolumn", "Beskrivning"])
        for row in scb["join_keys"]:
            sheet.append(row)
        path = source / "SCB/ID-kolumner.xlsx"
        workbook.properties.created = workbook.properties.modified = _XLSX_EPOCH
        workbook.save(path)
        _fix_zip_times(path)
    for relative, text in spec.get("files", {}).items():
        path = source / relative
        path.parent.mkdir(parents=True, exist_ok=True)
        if not isinstance(text, str):
            text = "\n".join(text) + "\n"
        path.write_text(text, encoding="utf-8")
    for register in spec.get("sos", ()):
        workbook_spec = SosRegister(
            abbrev=register["abbrev"],
            title_sv=register["title"],
            description_sv=register.get("description"),
            variables=tuple(
                SosVariable(**{k: v for k, v in row.items() if k != "linkage"})
                for row in register["variables"]
            ),
            deldatamangder=tuple(
                SosSubset(**row) for row in register.get("subsets", ())
            ),
        )
        sos_dir = write_sos_input(source, registers=(workbook_spec,))
        path = (
            sos_dir / f"Metadata {register['title']} ({register['abbrev']})_webb.xlsx"
        )
        # A delivered blank `Kopplingsvariabel` column is SOS's explicit "not an
        # identifier" claim; without it every variable's flag is unknown and withheld.
        workbook = load_workbook(path)
        sheet = workbook["Metadata - Variabelnivå"]
        linkage = sheet.max_column + 1
        sheet.cell(row=1, column=linkage, value="Kopplingsvariabel")
        # A variable's `linkage` text fills its cell: a declared linkage variable.
        for row, variable in enumerate(register["variables"], start=2):
            if variable.get("linkage"):
                sheet.cell(row=row, column=linkage, value=variable["linkage"])
        # A delivered Kodlista names its variable in a `Variabelnamn` preamble row;
        # without it the build cannot bind the list (`unresolved_list_reference`).
        code_lists = {
            f"Kodlista_{name}": [
                ["Variabelnamn", name],
                ["Tidsperiod", "Kod", "Beskrivning"],
                *rows,
            ]
            for name, rows in register.get("code_lists", {}).items()
        }
        for name, sheet_rows in {**code_lists, **register.get("sheets", {})}.items():
            extra = workbook.create_sheet(name)
            for row in sheet_rows:
                extra.append(row)
        if register.get("blank_dataset"):
            # The `Datamängd` cell names the register; blank, the workbook has none.
            workbook["Generell information"]["C4"] = None
        workbook.properties.created = workbook.properties.modified = _XLSX_EPOCH
        workbook.save(path)
        workbook.close()
        _fix_zip_times(path)


# openpyxl stamps the save time into docProps/core.xml and every zip entry, so two
# prepares of one spec would deliver different workbook bytes (and revisions), and a
# drift case could go stale on the timestamp alone.
_XLSX_EPOCH = datetime(2000, 1, 1)  # noqa: DTZ001 - openpyxl writes naive times as UTC


def _fix_zip_times(path: Path) -> None:
    with zipfile.ZipFile(path) as archive:
        entries = [(info, archive.read(info)) for info in archive.infolist()]
    with zipfile.ZipFile(path, "w") as archive:
        for info, data in entries:
            if info.filename == "docProps/core.xml":
                # `save` restamps `modified` with the current time.
                data = re.sub(
                    rb"(<dcterms:modified[^>]*>)[^<]*",
                    rb"\g<1>2000-01-01T00:00:00Z",
                    data,
                )
            fixed = zipfile.ZipInfo(info.filename, date_time=(1980, 1, 1, 0, 0, 0))
            fixed.compress_type = info.compress_type
            archive.writestr(fixed, data)


@dataclass(frozen=True)
class PreparedSet:
    """One accepted prepared input: what `build_catalog` and authoring read."""

    prepared: Path
    commit: str
    digest: str

    def records(self) -> tuple[SourceRecord, ...]:
        return _records(self.prepared, self.commit, self.digest)

    def bindings(self, record: SourceRecord) -> ValueBindingResult:
        """The source lists the prepared value store binds to one record."""
        with open_value_bindings(self.opened().value_sources) as sessions:
            return bind_code_lists(record, sessions)

    def claims(self, records: tuple[SourceRecord, ...]) -> list:
        """The code-list claims the prepared value store binds to ``records``."""
        return [claim for record in records for claim in self.bindings(record).claims]

    def coding_sha256(self, records: tuple[SourceRecord, ...]) -> list[str]:
        """The bound physical code-list evidence of ``records``, as curators pin it."""
        return [coding_source_sha256(claim) for claim in self.claims(records)]

    def opened(self):
        return _opened(self.prepared, self.commit, self.digest)

    def revision(self, record: SourceRecord) -> dict:
        return next(
            revision.model_dump(mode="json")
            for revision in self.opened().records.manifest.revisions
            if revision.revision_id == record.source_revision_id
        )

    def relationship(self, table: str):
        """The one literal relationship declaration delivered in ``table``."""
        (declaration,) = (
            evidence.declaration
            for evidence in self.opened().iter_evidence()
            if isinstance(evidence, ReferenceEvidence)
            and evidence.declaration.locator.physical_table == table
        )
        return declaration

    def table_sha256(self, table: str) -> str:
        """The content hash of the one prepared evidence table named ``table``."""
        (found,) = (t for t in self.opened().records.iter_tables() if t.name == table)
        return canonical_sha256(found.model_dump(mode="json"))


# Authoring reads open a cache entry's accepted sources once per process: an entry
# is immutable once renamed into place, and the build under test still opens and
# checks it itself.
@functools.cache
def _opened(prepared: Path, commit: str, digest: str):
    return open_prepared_catalog_sources(
        prepared, expected_sha256=digest, input_commit=commit
    )


@functools.cache
def _records(prepared: Path, commit: str, digest: str) -> tuple[SourceRecord, ...]:
    return tuple(_opened(prepared, commit, digest).records.records)


class PreparedCache:
    """Accepted prepared inputs keyed by the content hash of their source spec."""

    def __init__(self, root: Path) -> None:
        self.root = root
        root.mkdir(parents=True, exist_ok=True)

    def get(self, spec: dict) -> PreparedSet:
        key = hashlib.sha256(
            json.dumps(spec, sort_keys=True, ensure_ascii=False).encode()
        ).hexdigest()
        entry = self.root / key
        if not (entry / "selection.json").is_file():
            staging = Path(tempfile.mkdtemp(prefix=".staging-", dir=self.root))
            try:
                write_sources(spec, staging / "source")
                _, commit, digest = prepare_accepted(staging, staging / "source")
                # Only the accepted prepared repository is read after preparation.
                shutil.rmtree(staging / "source")
                shutil.rmtree(staging / "inputs")
                (staging / "selection.json").write_text(
                    json.dumps({"commit": commit, "digest": digest}), encoding="utf-8"
                )
                try:
                    staging.rename(entry)
                except OSError:
                    if not (entry / "selection.json").is_file():
                        raise
            finally:
                shutil.rmtree(staging, ignore_errors=True)
        selection = json.loads((entry / "selection.json").read_text(encoding="utf-8"))
        return PreparedSet(
            entry / "prepared" / "catalog", selection["commit"], selection["digest"]
        )


@pytest.fixture(scope="session")
def prepared_cache() -> PreparedCache:
    """The prepared inputs shared by the build cases and `catalog`.

    They live in a generation of the shared fixture cache, under the reader
    cache's one location policy (`reader_artifacts.fixture_cache_dir`):
    `$REG_FIXTURE_CACHE`, else the system temp directory. Entries are reused across
    sessions, worktrees and xdist workers; a fresh CI runner starts cold.
    """
    return PreparedCache(fixture_generation() / "build-case-prepared")


@functools.cache
def fixture_generation() -> Path:
    """This checkout's generation of the shared fixture cache.

    Keyed by `_prepared_inputs_digest`, so any cache under it (the prepared inputs,
    the `cases/cli/` artifact) is invalidated by a builder or fixture-module edit.
    """
    return generation_dir(_prepared_inputs_digest())


# The test modules that write and accept a source spec's provider deliveries; an
# edit to any of them changes what a prepared entry holds.
_PREPARING_MODULES = (
    "_build_case_runner.py",
    "_csv_fixtures.py",
    "_sos_fixtures.py",
    "_prepared_fixtures.py",
    "_pipeline_catalog_support.py",
)


def _prepared_inputs_digest() -> str:
    """Everything a prepared entry depends on besides its source spec.

    The reader cache's build inputs (the `reg_meta_build` sources, the native
    extension sources, the installed
    distributions, Python and SQLite), the modules that prepare a spec, and the Git
    that commits and checks the accepted repository.
    """
    here = Path(__file__).parent
    return canonical_sha256(
        {
            "build_inputs": build_inputs_digest(),
            "modules": {
                name: hashlib.sha256((here / name).read_bytes()).hexdigest()
                for name in _PREPARING_MODULES
            },
            "git": subprocess.run(
                ["git", "--version"], capture_output=True, text=True, check=True
            ).stdout.strip(),
        }
    )


# -- Curation authoring placeholders --------------------------------------------

PROSE = ("name", "definition", "description", "operational_definition")
_PLACEHOLDER = re.compile(r"\{\{\s*([a-z0-9_]+)((?:\s+[a-z_]+=[^\s}]*)*)\s*\}\}")


def toml_inline(value) -> str:
    """A JSON-shaped value as one inline TOML value; ``None`` members are omitted."""
    if isinstance(value, dict):
        items = ", ".join(
            f"{json.dumps(key)} = {toml_inline(item)}"
            for key, item in value.items()
            if item is not None
        )
        return "{ " + items + " }"
    if isinstance(value, (list, tuple)):
        return "[" + ", ".join(toml_inline(item) for item in value) + "]"
    if isinstance(value, bool):
        return "true" if value else "false"
    if isinstance(value, str):
        return json.dumps(value, ensure_ascii=False)
    return str(value)


def _column(record: SourceRecord) -> str:
    column = record.fields.column_name
    return (column and column.value) or ""


_FILTERS: dict[str, Callable[[SourceRecord], object]] = {
    "source": lambda r: r.source,
    "key": lambda r: r.locators[0].semantic_record_key[-1],
    "column": _column,
    "member": lambda r: r.subject.member.name,
    "variable": lambda r: str(
        r.subject.native.variable_id
        if r.subject.native.variable_id is not None
        else r.subject.variable.native_id
    ),
    "edition": lambda r: str(r.subject.native.edition_id),
}


def _select(records: tuple[SourceRecord, ...], args: dict[str, str]) -> tuple:
    """Records matching every filter argument; a value may list alternatives."""
    filters = {key: value.split(",") for key, value in args.items() if key in _FILTERS}
    selected = tuple(
        record
        for record in records
        if all(
            str(_FILTERS[key](record)).startswith(tuple(values))
            if key == "source"
            else str(_FILTERS[key](record)) in values
            for key, values in filters.items()
        )
    )
    assert selected, f"placeholder selects no prepared record: {args}"
    return selected


def _one(records: tuple[SourceRecord, ...], args: dict[str, str]) -> SourceRecord:
    selected = _select(records, args)
    assert len(selected) == 1, f"placeholder needs exactly one record: {args}"
    return selected[0]


def _fields(args: dict[str, str]) -> tuple[str, ...]:
    named = args.get("fields", "all")
    return (
        tuple(SourceFields.model_fields)
        if named == "all"
        else tuple(PROSE if named == "prose" else named.split(","))
    )


def _marker_bindings(s: PreparedSet, record: SourceRecord, args: dict[str, str]):
    """The fingerprints of the record's marker bindings in ``from``..``to``."""
    scope = source_occurrence(record).edition_period_scope
    markers = marker_binding_fingerprints(
        ((scope, binding) for binding in s.bindings(record).bindings),
        args["from"],
        args["to"],
    )
    assert markers is not None, f"placeholder binds no marker: {args}"
    return list(markers)


def _association(s: PreparedSet, args: dict[str, str]):
    """The one source-list association behind member ``code=``/``label=`` (blank if
    omitted) of the claims bound to the selected records."""
    found = [
        association
        for claim in s.claims(_select(s.records(), args))
        for member in claim.members
        if member.code == args["code"] and (member.label or "") == args.get("label", "")
        for association in member.associations
    ]
    assert len(found) == 1, f"placeholder needs exactly one association: {args}"
    return found[0]


def _scope(record: SourceRecord, args: dict[str, str], attribute: str) -> dict:
    scope = getattr(record, attribute).model_dump(mode="json")
    if "end" in args:
        scope["intervals"][0]["end"] = args["end"]
    return scope


_DIRECTIVES: dict[str, Callable[[PreparedSet, dict[str, str]], object]] = {
    "evidence_sha256": lambda s, a: acknowledgement_evidence_sha256(
        _select(s.records(), a)
    ),
    "coding_evidence_sha256": lambda s, a: acknowledgement_evidence_sha256(
        selected := _select(s.records(), a), s.coding_sha256(selected)
    ),
    "expected_records": lambda s, a: [
        expectation.model_dump(mode="json")
        for expectation in capture_expectations(
            _select(s.records(), a), fields=_fields(a), parents=True, coding=True
        )
    ],
    "expected_fields": lambda s, a: [
        field.model_dump(mode="json")
        for field in capture_expectations((_one(s.records(), a),), fields=_fields(a))[0]
        .alternatives[0]
        .fields
    ],
    "period_text": lambda s, a: _one(s.records(), a).original_period_text,
    "edition_scope": lambda s, a: _scope(_one(s.records(), a), a, "edition_scope"),
    "period_scope": lambda s, a: _scope(
        _one(s.records(), a), a, "edition_period_scope"
    ),
    "revision": lambda s, a: s.revision(_one(s.records(), a)),
    "locators": lambda s, a: [
        locator.model_dump(mode="json")
        for record in _select(s.records(), a)
        for locator in record.locators
    ],
    "marker_bindings": lambda s, a: _marker_bindings(s, _one(s.records(), a), a),
    "raw_codings": lambda s, a: sorted(set(s.coding_sha256(_select(s.records(), a)))),
    "source_codings": lambda s, a: list(
        copied_coding_fingerprints(s.claims(_select(s.records(), a)))
    ),
    "association": lambda s, a: _association(s, a).locator,
    "association_sha256": lambda s, a: coding_source_sha256(_association(s, a)),
    "relationship_row": lambda s, a: s.relationship(a["table"]).locator.physical_record,
    "relationship_sha256": lambda s, a: canonical_sha256(
        s.relationship(a["table"]).model_dump(mode="json")
    ),
    "table_sha256": lambda s, a: s.table_sha256(a["table"]),
    "naming_id": lambda s, a: authored_naming_id(
        a["kind"],
        provider=a["provider"],
        register_key=a["register_key"],
        member_key=a.get("member_key"),
    ),
    "variant_key": lambda s, a: _variant_key(_one(s.records(), a)),
}


def _variant_key(record: SourceRecord) -> list:
    key = native_variant_key(record)
    assert key is not None, f"record has no native variant: {record.record_id}"
    return list(key)


def render_curation(source: Path, target: Path, authored: PreparedSet) -> Path:
    """Copy a case's curation tree, resolving each `{{directive arg=value}}`."""

    def resolve(match: re.Match[str]) -> str:
        name, raw = match.group(1), match.group(2).split()
        args = dict(item.split("=", 1) for item in raw)
        assert name in _DIRECTIVES, f"unknown curation placeholder: {name}"
        return toml_inline(_DIRECTIVES[name](authored, args))

    for path in source.rglob("*"):
        if path.is_file():
            out = target / path.relative_to(source)
            out.parent.mkdir(parents=True, exist_ok=True)
            out.write_text(
                _PLACEHOLDER.sub(resolve, path.read_text(encoding="utf-8")),
                encoding="utf-8",
            )
    (target / "classifications").mkdir(parents=True, exist_ok=True)
    return target


# -- Projections ----------------------------------------------------------------


def _locator(detail: str) -> str | None:
    return detail.split(":", 1)[0] if detail.startswith("curation/") else None


@dataclass(frozen=True)
class Outcome:
    result: dict
    events: list[dict]
    db: Path
    built: PreparedSet

    def table(self, name: str) -> list[dict]:
        return _TABLES[name](self)

    def _sql(self, sql: str) -> list[dict]:
        with closing(sqlite3.connect(self.db)) as conn:
            conn.row_factory = sqlite3.Row
            return [dict(row) for row in conn.execute(sql)]


def _refs(refs) -> list[str]:
    """Source record refs as `<source>#<semantic key parts joined by />`."""
    return [
        f"{ref['source']}#{'/'.join(ref['semantic_record_key'])}" for ref in refs or ()
    ]


def _issue_row(event: dict) -> dict:
    return {
        "code": event["code"],
        "severity": event["severity"],
        "subject": str(event["subject"]),
        "case_id": event.get("case_id"),
        "locator": _locator(str(event.get("detail") or "")),
        "detail": event.get("detail"),
        "acknowledged_by": event.get("acknowledged_by"),
        "valid_from": event.get("valid_from"),
        "valid_to": event.get("valid_to"),
        "fields": event.get("fields") or [],
        "withheld_output": event.get("withheld_output") or [],
    }


def _issues(outcome: Outcome) -> list[dict]:
    return [_issue_row(event) for event in outcome.events if event["kind"] == "issue"]


def _issue_refs(outcome: Outcome) -> list[dict]:
    """One row per source record an issue cites: the issue's fields, then the ref,
    as its parts and in the `_refs` string form other tables' `refs` use."""
    return [
        {
            **_issue_row(event),
            "source": ref["source"],
            "key": ref["semantic_record_key"],
            "ref": _refs([ref])[0],
        }
        for event in outcome.events
        if event["kind"] == "issue"
        for ref in event.get("refs") or ()
    ]


def _source_issues(outcome: Outcome) -> list[dict]:
    """Support- and value-source issues: the source evidence behind ledger issues."""
    return [
        {
            "kind": event["kind"],
            "severity": event.get("severity"),
            "descriptor_key": event.get("descriptor_key"),
            "physical_associations": event.get("physical_associations"),
            "refs": _refs(event.get("refs")),
        }
        for event in outcome.events
        if event["kind"] in {"support_source_issue", "value_source_issue"}
    ]


def _dispositions(outcome: Outcome) -> list[tuple[dict, list[str]]]:
    """Each ledger disposition of a prepared source record occurrence as a `uses`
    row, with the curation case ids the disposition names."""
    by_id = {
        event["record_id"]: event["dispositions"]
        for event in outcome.events
        if event["kind"] == "source_occurrence"
    }
    rows = []
    for record in outcome.built.records():
        fields = {
            name: value.value
            if (value := getattr(record.fields, name)) is not None
            else None
            for name in ("column_name", "data_type", "name", "description")
        }
        for disposition in by_id.get(record.record_id, ()):
            row = {
                "source": record.source,
                "native_variable": _FILTERS["variable"](record),
                "key": record.locators[0].semantic_record_key[-1],
                **fields,
                "use": disposition["use"],
                "variable": disposition["variable"],
            }
            rows.append((row, disposition["cases"]))
    return rows


def _case_uses(outcome: Outcome) -> list[dict]:
    """One row per curation case a source record's ledger disposition names."""
    return [
        {"case_id": case_id, **{key: row[key] for key in _CASE_USE_FIELDS}}
        for row, cases in _dispositions(outcome)
        for case_id in cases
    ]


_CASE_USE_FIELDS = ("source", "key", "use", "variable")


def _variables(outcome: Outcome) -> list[dict]:
    """Each built variable's distinct delivery columns; one null row when it has none."""
    rows = outcome._sql(
        "SELECT DISTINCT r.slug AS register, v.slug AS variable, "
        "s.delivery_column_name AS column, v.provider_key, v.name, v.definition, "
        "v.description, v.measurement_unit, v.operational_definition, "
        "v.is_identifier, v.is_sensitive, v.deprecated, "
        "sr.slug AS source_register, v.source_label, v.source_register_text "
        "FROM variable v JOIN register r USING (register_id) "
        "LEFT JOIN register sr ON sr.register_id = v.source_register_id "
        "LEFT JOIN variable_state s USING (variable_id)"
    )
    named = {
        (row["register"], row["variable"]) for row in rows if row["column"] is not None
    }
    return [
        row
        for row in rows
        if row["column"] is not None or (row["register"], row["variable"]) not in named
    ]


def _concept_groups(outcome: Outcome) -> list[dict]:
    """Each concept group as the sorted slugs of its member variables."""
    groups: dict[int, list[str]] = {}
    for row in outcome._sql(
        "SELECT gv.group_id, v.slug FROM concept_group_variable gv "
        "JOIN variable v USING (variable_id)"
    ):
        groups.setdefault(row["group_id"], []).append(row["slug"])
    return [{"variables": sorted(slugs)} for slugs in groups.values()]


def _group_members(outcome: Outcome) -> list[dict]:
    """Each concept-group member with its literal column and its facets."""
    facets: dict[int, list[list]] = {}
    for row in outcome._sql(
        "SELECT member_id, axis, value, label FROM concept_group_variable_facet"
    ):
        facets.setdefault(row["member_id"], []).append(
            [row["axis"], row["value"], row["label"]]
        )
    rows = outcome._sql(
        "SELECT m.member_id, g.group_key, r.slug AS register, "
        "v.slug AS variable, m.delivery_column_name AS column "
        "FROM concept_group_variable m JOIN concept_group g USING (group_id) "
        "JOIN variable v USING (variable_id) "
        "JOIN register r ON r.register_id = v.register_id"
    )
    for row in rows:
        row["facets"] = sorted(facets.get(row.pop("member_id"), []))
    return rows


def _editions(outcome: Outcome) -> list[dict]:
    """Each register version with its prose, populations and object types."""
    populations: dict[int, list[list]] = {}
    for row in outcome._sql(
        "SELECT regver_id, name, definition, comment, date_range FROM population"
    ):
        populations.setdefault(row.pop("regver_id"), []).append(list(row.values()))
    objects: dict[int, list[list]] = {}
    for row in outcome._sql("SELECT regver_id, name, definition FROM object_type"):
        objects.setdefault(row.pop("regver_id"), []).append(list(row.values()))
    rows = outcome._sql(
        "SELECT e.regver_id, r.slug AS register, rv.slug AS variant, "
        "e.registerversionnamn AS name, "
        "e.registerversionbeskrivning AS description, "
        "e.registerversionmatinformation AS measurement_information, "
        "e.registerversion_docstaus AS documentation_status, "
        "e.registerversion_forstagodkannandedatum AS first_approved_at, "
        "e.registerversion_senastgodkanddatum AS last_approved_at "
        "FROM register_version e JOIN register_variant rv USING "
        "(register_variant_id) JOIN register r USING (register_id)"
    )
    for row in rows:
        edition = row.pop("regver_id")
        row["populations"] = _sorted(populations.get(edition, []))
        row["object_types"] = _sorted(objects.get(edition, []))
    return rows


def _value_sets(outcome: Outcome) -> list[dict]:
    """Each stored value set: its id, its sorted members and the sorted FQIDs of the
    variables whose states or alias windows carry it."""
    sets: dict[int, dict] = {}
    for row in outcome._sql(
        "SELECT value_set_id, code, label FROM value_set "
        "LEFT JOIN value_set_member USING (value_set_id) "
        "LEFT JOIN value_code USING (code_id)"
    ):
        entry = sets.setdefault(
            row["value_set_id"],
            {"id": str(row["value_set_id"]), "members": [], "variables": set()},
        )
        if row["code"] is not None:
            entry["members"].append([row["code"], row["label"]])
    for row in outcome._sql(
        "SELECT u.value_set_id, p.slug || '/' || r.slug || '/' || v.slug AS fqid "
        "FROM (SELECT variable_id, value_set_id FROM variable_state "
        "UNION SELECT variable_id, value_set_id FROM variable_alias_window) u "
        "JOIN variable v USING (variable_id) "
        "JOIN register r ON r.register_id = v.register_id "
        "JOIN provider p ON p.provider_id = r.provider_id "
        "WHERE u.value_set_id IS NOT NULL"
    ):
        sets[row["value_set_id"]]["variables"].add(row["fqid"])
    return [
        {
            **entry,
            "members": sorted(entry["members"]),
            "variables": sorted(entry["variables"]),
        }
        for entry in sets.values()
    ]


def _stored(row_value, json_value):
    """A coordinate both the `data_warning` row and its `warning_json` hold: the
    value when they agree, else both, so a disagreement fails any case naming it."""
    if row_value == json_value:
        return row_value
    return {"row": row_value, "json": json_value}


def _sha256(text: str) -> str:
    return hashlib.sha256(text.encode("utf-8")).hexdigest()


def _warnings(outcome: Outcome) -> list[dict]:
    issue_hashes = {
        (event["code"], str(event["subject"]), _sha256(event.get("detail") or ""))
        for event in outcome.events
        if event["kind"] == "issue"
    }
    rows = []
    for row in outcome._sql(
        "SELECT r.slug AS register, v.slug AS variable, rv.slug AS variant, "
        "w.delivery_column_name, w.valid_from, w.valid_to, w.warning_json "
        "FROM data_warning w JOIN register r USING (register_id) "
        "LEFT JOIN variable v ON v.variable_id = w.variable_id "
        "LEFT JOIN register_variant rv "
        "ON rv.register_variant_id = w.register_variant_id"
    ):
        payload = json.loads(row["warning_json"])
        variable = payload.get("variable_fqid")
        digest = payload["diagnostic_detail_sha256"]
        rows.append(
            {
                "register": _stored(
                    row["register"], payload["register_fqid"].split("/")[-1]
                ),
                "variable": _stored(
                    row["variable"], variable.split("/")[-1] if variable else None
                ),
                "variant": _stored(row["variant"], payload.get("variant")),
                "column": _stored(
                    row["delivery_column_name"], payload.get("delivery_column_name")
                ),
                "valid_from": _stored(row["valid_from"], payload.get("valid_from")),
                "valid_to": _stored(row["valid_to"], payload.get("valid_to")),
                "code": payload["code"],
                "severity": payload["severity"],
                "detail": payload["detail"],
                "summary": payload.get("summary"),
                "detail_hash_of": "issue"
                if (payload["code"], payload["source_subject"], digest) in issue_hashes
                else "detail"
                if _sha256(payload["detail"]) == digest
                else None,
                "fields": payload.get("fields"),
                "refs": _refs(payload.get("refs")),
                "withheld_output": payload.get("withheld_output"),
                "acknowledged_by": payload.get("acknowledged_by"),
                "source_subject": payload["source_subject"],
                "case_id": payload.get("case_id"),
            }
        )
    return rows


def _alias_windows(outcome: Outcome) -> list[dict]:
    """Each alias window of a delivery column, with its per-column facts, the
    members of its per-column value set and the books bound to it."""
    key = "w.variable_id, w.register_variant_id, w.delivery_column_name, w.valid_from"
    codes: dict[tuple, list[list[str]]] = {}
    for row in outcome._sql(
        f"SELECT {key}, c.code, c.label FROM variable_alias_window w "
        "JOIN value_set_member m ON m.value_set_id = w.value_set_id "
        "JOIN value_code c ON c.code_id = m.code_id"
    ):
        codes.setdefault(tuple(row.values())[:4], []).append(
            [row["code"], row["label"]]
        )
    books: dict[tuple, list[str]] = {}
    for row in outcome._sql(
        f"SELECT {key}, c.slug FROM alias_window_classification w "
        "JOIN classification c ON c.id = w.classification_id"
    ):
        books.setdefault(tuple(row.values())[:4], []).append(row["slug"])
    rows = []
    for row in outcome._sql(
        f"SELECT {key}, r.slug AS register, v.slug AS variable, rv.slug AS variant, "
        "w.delivery_column_name AS column, w.valid_to, w.column_metadata, "
        "w.provenance, w.coding_metadata, w.data_type, w.data_length, w.definition, "
        "w.measurement_unit, w.name, w.description, w.operational_definition, "
        "w.source_register_text FROM variable_alias_window w JOIN variable v USING (variable_id) "
        "JOIN register r ON r.register_id = v.register_id "
        "JOIN register_variant rv ON rv.register_variant_id = w.register_variant_id"
    ):
        ident = tuple(row.values())[:4]
        row = {
            k: v
            for k, v in row.items()
            if k not in {"variable_id", "register_variant_id", "delivery_column_name"}
        }
        rows.append(
            {
                **row,
                "codes": sorted(codes.get(ident, [])),
                "classifications": sorted(books.get(ident, [])),
            }
        )
    return rows


_ALIAS_CONFORMANCE_SQL = (
    "SELECT r.slug AS register, v.slug AS variable, "
    "w.delivery_column_name AS column, w.valid_from, w.valid_to, "
    "c.slug AS classification, a.conformance "
    "FROM alias_window_classification a "
    "JOIN variable_alias_window w USING "
    "(variable_id, register_variant_id, delivery_column_name, valid_from) "
    "JOIN classification c ON c.id = a.classification_id "
    "JOIN variable v ON v.variable_id = a.variable_id "
    "JOIN register r ON r.register_id = v.register_id "
    "WHERE a.conformance IS NOT NULL"
)


def _alias_conformance(outcome: Outcome) -> list[tuple[dict, dict]]:
    """Each alias window's book binding with its stored conformance evidence."""
    return [
        (
            {key: row[key] for key in row.keys() - {"conformance"}},
            json.loads(row["conformance"]),
        )
        for row in outcome._sql(_ALIAS_CONFORMANCE_SQL)
    ]


def _alias_extensions(evidence: dict) -> list[tuple[str, str, str]]:
    return [
        (code, label, kind)
        for kind, key in (
            ("nonstandard", "nonconforming_members"),
            ("sentinel", "sentinel_members"),
        )
        for code, label in evidence[key]
    ]


def _conformance(outcome: Outcome) -> list[dict]:
    """Each state's and alias window's conformance to one bound book.

    An alias window stores its evidence as JSON; its counts are read the way the
    reader counts them (distinct checked codes, distinct codes outside the book).
    """
    rows = [
        {**row, "window": "state"}
        for row in outcome._sql(
            "SELECT r.slug AS register, v.slug AS variable, "
            "s.delivery_column_name AS column, s.valid_from, s.valid_to, "
            "c.slug AS classification, cc.status, cc.checked_code_count AS checked, "
            "cc.matched_code_count AS matched, "
            "cc.nonconforming_code_count AS nonconforming, cc.overlap "
            "FROM classification_conformance cc "
            "JOIN classification c ON c.id = cc.declared_classification_id "
            "JOIN variable_state s ON s.state_id = cc.state_id "
            "JOIN variable v ON v.variable_id = s.variable_id "
            "JOIN register r ON r.register_id = v.register_id"
        )
    ]
    for row, evidence in _alias_conformance(outcome):
        checked = len(set(evidence["checked_codes"]))
        outside = len({code for code, _, _ in _alias_extensions(evidence)})
        rows.append(
            {
                **row,
                "window": "alias",
                "status": evidence["status"],
                "checked": checked,
                "matched": checked - outside,
                "nonconforming": outside,
                "overlap": (checked - outside) / checked if checked else 1.0,
            }
        )
    return rows


def _conformance_codes(outcome: Outcome) -> list[dict]:
    """Each source member a state's or alias window's book conformance records
    outside the book.

    A scoped sentinel certificate is projected as its window only; its fingerprints
    restate the code under test. An alias window stores no sentinel meaning.
    """
    rows = [
        {
            **{key: row[key] for key in row.keys() - {"scoped_sentinels"}},
            "window": "state",
            "scoped_windows": [
                [certificate["valid_from"], certificate["valid_to"]]
                for certificate in json.loads(row["scoped_sentinels"])
            ],
        }
        for row in outcome._sql(
            "SELECT r.slug AS register, v.slug AS variable, "
            "s.delivery_column_name AS column, s.valid_from, s.valid_to, "
            "c.slug AS classification, vc.code, vc.label, cc.member_kind, "
            "cc.sentinel_meaning, cc.scoped_sentinels "
            "FROM classification_conformance_code cc "
            "JOIN classification c ON c.id = cc.declared_classification_id "
            "JOIN value_code vc ON vc.code_id = cc.code_id "
            "JOIN variable_state s ON s.state_id = cc.state_id "
            "JOIN variable v ON v.variable_id = s.variable_id "
            "JOIN register r ON r.register_id = v.register_id"
        )
    ]
    for row, evidence in _alias_conformance(outcome):
        rows.extend(
            {
                **row,
                "window": "alias",
                "code": code,
                "label": label,
                "member_kind": kind,
                "sentinel_meaning": None,
                "scoped_windows": [
                    [certificate["valid_from"], certificate["valid_to"]]
                    for certificate in evidence["scoped_sentinels"]
                    if [code, label] in certificate["members"]
                ]
                if kind == "sentinel"
                else [],
            }
            for code, label, kind in _alias_extensions(evidence)
        )
    return rows


_STATE_JOIN = (
    "FROM variable_state s JOIN variable v USING (variable_id) "
    "JOIN register r ON r.register_id = v.register_id "
    "JOIN register_variant rv ON rv.register_variant_id = s.register_variant_id "
)
# Each state's id and coordinates; the lineage tables join it once per edge end.
_STATE_COORDINATES = (
    "SELECT s.state_id, r.slug AS register, v.slug AS variable, rv.slug AS variant, "
    "s.delivery_column_name AS column, s.valid_from, s.valid_to " + _STATE_JOIN
)
# A succession endpoint's register or variable FQID, from a `{side}_provider`,
# `{side}_register` (and `{side}_variable`) column triple.
_REGISTER = "{0}_provider || '/' || {0}_register"
_VARIABLE = "{0}_provider || '/' || {0}_register || '/' || {0}_variable"
_SUCCESSION_FACTS = "effective_year, note, beskrivning AS description"
_TABLES: dict[str, Callable[[Outcome], list[dict]]] = {
    "issues": _issues,
    "issue_refs": _issue_refs,
    "cases": lambda o: [
        {"case_id": event["case_id"], "status": event["status"]}
        for event in o.events
        if event["kind"] == "case"
    ],
    "uses": lambda o: [row for row, _ in _dispositions(o)],
    "case_uses": _case_uses,
    "states": lambda o: o._sql(
        "SELECT r.slug AS register, v.slug AS variable, rv.slug AS variant, "
        "s.delivery_column_name AS column, s.valid_from, s.valid_to, s.data_type, "
        "v.name, s.name AS state_name, s.provenance, s.pooled, s.data_length, "
        "s.definition, s.measurement_unit, s.description, s.operational_definition, "
        "s.source_register_text, s.value_set_version_label " + _STATE_JOIN
    ),
    "state_codes": lambda o: o._sql(
        "SELECT r.slug AS register, v.slug AS variable, rv.slug AS variant, "
        "s.delivery_column_name AS column, s.valid_from, s.valid_to, c.code, c.label "
        + _STATE_JOIN
        + "LEFT JOIN value_set_member m ON m.value_set_id = s.value_set_id "
        "LEFT JOIN value_code c ON c.code_id = m.code_id"
    ),
    "variables": _variables,
    "variants": lambda o: o._sql(
        "SELECT r.slug AS register, rv.slug AS variant, rv.name, "
        "rv.panel_entity_key, rv.panel_time_key, rv.panel_time_grain, "
        "rv.description, rv.display_group "
        "FROM register_variant rv JOIN register r USING (register_id)"
    ),
    "registers": lambda o: o._sql(
        "SELECT slug AS register, name, purpose FROM register"
    ),
    "editions": _editions,
    # One row per tag member as its FQID; a tag without members is one null row.
    "tags": lambda o: o._sql(
        "SELECT t.slug, CASE WHEN m.variable_id IS NOT NULL THEN "
        "vp.slug || '/' || vr.slug || '/' || v.slug "
        "WHEN m.register_id IS NOT NULL THEN mp.slug || '/' || mr.slug END "
        "AS member FROM tag t LEFT JOIN tag_member m USING (tag_id) "
        "LEFT JOIN variable v ON v.variable_id = m.variable_id "
        "LEFT JOIN register vr ON vr.register_id = v.register_id "
        "LEFT JOIN provider vp ON vp.provider_id = vr.provider_id "
        "LEFT JOIN register mr ON mr.register_id = m.register_id "
        "LEFT JOIN provider mp ON mp.provider_id = mr.provider_id"
    ),
    "aliases": lambda o: o._sql(
        "SELECT r.slug AS register, v.slug AS variable, rv.slug AS variant, "
        "a.delivery_column_name AS column FROM variable_alias a "
        "JOIN variable v USING (variable_id) "
        "JOIN register r ON r.register_id = v.register_id "
        "JOIN register_variant rv ON rv.register_variant_id = a.register_variant_id"
    ),
    "alias_windows": _alias_windows,
    "concept_groups": _concept_groups,
    "group_axes": lambda o: o._sql(
        "SELECT g.group_key, a.axis, a.ordinal, a.label FROM concept_group_axis a "
        "JOIN concept_group g USING (group_id)"
    ),
    "group_members": _group_members,
    "warnings": _warnings,
    "search_pins": lambda o: o._sql(
        "SELECT key AS query, type, position, entity FROM search_pin"
    ),
    # The unfolded text `variable_fts` indexes, one row per built variable.
    "search_text": lambda o: [
        {**row, "delivery_column_names": json.loads(row["delivery_column_names"])}
        for row in o._sql(
            "SELECT r.slug AS register, v.slug AS variable, t.name, t.definition, "
            "t.description, t.delivery_column_names FROM variable_search_text t "
            "JOIN variable v USING (variable_id) "
            "JOIN register r ON r.register_id = v.register_id"
        )
    ],
    "manifest": lambda o: o._sql("SELECT key, value FROM import_manifest"),
    "edges": lambda o: o._sql(
        "SELECT 'same_as' AS type, "
        "a_provider || '/' || a_register || '/' || a_variable AS a, "
        "b_provider || '/' || b_register || '/' || b_variable AS b "
        "FROM variable_same_as UNION ALL SELECT 'replaced_by', "
        "predecessor_provider || '/' || predecessor_register || '/' "
        "|| predecessor_variable, "
        "successor_provider || '/' || successor_register || '/' || successor_variable "
        "FROM variable_replaced_by"
    ),
    # Every replaced_by edge below the classification grain, in one row shape.
    "successions": lambda o: o._sql(
        "SELECT 'register' AS grain, "
        f"{_REGISTER.format('predecessor')} AS predecessor, "
        f"{_REGISTER.format('successor')} AS successor, NULL AS predecessor_variant, "
        "NULL AS successor_variant, NULL AS predecessor_column, "
        "NULL AS successor_column, NULL AS variant, "
        f"{_SUCCESSION_FACTS} FROM register_replaced_by UNION ALL SELECT 'variable', "
        f"{_VARIABLE.format('predecessor')}, {_VARIABLE.format('successor')}, "
        f"NULL, NULL, NULL, NULL, NULL, {_SUCCESSION_FACTS} FROM variable_replaced_by "
        f"UNION ALL SELECT 'variant', {_REGISTER.format('predecessor')}, "
        f"{_REGISTER.format('successor')}, predecessor_variant, successor_variant, "
        f"NULL, NULL, NULL, {_SUCCESSION_FACTS} FROM variant_replaced_by "
        f"UNION ALL SELECT 'representation', {_VARIABLE.format('predecessor')}, "
        f"{_VARIABLE.format('successor')}, NULL, NULL, predecessor_column, "
        f"successor_column, variant, {_SUCCESSION_FACTS} "
        "FROM representation_replaced_by"
    ),
    "classification_derivations": lambda o: o._sql(
        "SELECT derived_slug AS derived, source_slug AS source, note "
        "FROM classification_derived_from"
    ),
    "timeseries_events": lambda o: o._sql(
        "SELECT namn AS name, handelse AS event, beskrivning AS description, "
        "entitet AS entity, id1 AS first_token, id2 AS second_token, "
        "fil_id AS file_token FROM timeseries_event"
    ),
    "source_columns": lambda o: o._sql(
        "SELECT table_name, column_name, sql_type, nullable FROM source_column_type"
    ),
    "join_keys": lambda o: o._sql(
        "SELECT table_name, column_name, description FROM source_join_key"
    ),
    "identifiers": lambda o: o._sql(
        "SELECT var_id AS native_variable, variabelnamn AS name, "
        "variabeldefinition AS definition FROM identifier_semantics"
    ),
    "lineage": lambda o: o._sql(
        "SELECT c.register, c.variable, c.variant, c.column, "
        "s.register AS source_register, s.variable AS source_variable, "
        "s.variant AS source_variant, s.column AS source_column, "
        "l.valid_from, l.valid_to FROM variable_state_lineage l "
        f"JOIN ({_STATE_COORDINATES}) c ON c.state_id = l.consumer_state_id "
        f"JOIN ({_STATE_COORDINATES}) s ON s.state_id = l.source_state_id"
    ),
    "lineage_warnings": lambda o: o._sql(
        "SELECT c.register, c.variable, c.variant, c.column, c.valid_from, "
        "c.valid_to, w.warning_kind AS kind, w.message "
        "FROM variable_state_lineage_warning w "
        f"JOIN ({_STATE_COORDINATES}) c ON c.state_id = w.consumer_state_id"
    ),
    "state_classifications": lambda o: o._sql(
        "SELECT r.slug AS register, v.slug AS variable, "
        "s.delivery_column_name AS column, s.valid_from, s.valid_to, "
        "c.slug AS classification, sc.provenance "
        "FROM state_classification sc "
        "JOIN classification c ON c.id = sc.classification_id "
        "JOIN variable_state s ON s.state_id = sc.state_id "
        "JOIN variable v ON v.variable_id = s.variable_id "
        "JOIN register r ON r.register_id = v.register_id"
    ),
    "conformance": _conformance,
    "conformance_codes": _conformance_codes,
    # The reader's code-search index: each code a variable carries through a
    # state's or an alias window's value set, with the code's variable count.
    "code_index": lambda o: o._sql(
        "SELECT r.slug AS register, v.slug AS variable, vc.code, vc.label, "
        "vc.mapping_count FROM code_variable_map m "
        "JOIN value_code vc USING (code_id) JOIN variable v USING (variable_id) "
        "JOIN register r ON r.register_id = v.register_id"
    ),
    "value_sets": _value_sets,
    "classifications": lambda o: o._sql(
        "SELECT c.slug, c.short_name, c.name, c.name_en, c.publisher, c.valid_from, "
        "c.valid_to, c.description, c.url, c.code_count, c.valid_code_count, "
        "p.slug AS supersedes FROM classification c "
        "LEFT JOIN classification p ON p.id = c.supersedes_id"
    ),
    "classification_successions": lambda o: o._sql(
        "SELECT predecessor_slug AS predecessor, successor_slug AS successor, "
        "effective_year, note FROM classification_replaced_by"
    ),
    "classification_codes": lambda o: o._sql(
        "SELECT c.slug, v.code, v.label, cc.level, cc.is_valid "
        "FROM classification_code cc "
        "JOIN classification c ON c.id = cc.classification_id "
        "JOIN value_code v ON v.code_id = cc.code_id"
    ),
    "relationships": lambda o: o._sql(
        "SELECT kind, binding_status, source_dataset, v.slug AS owner, "
        "(SELECT COUNT(*) FROM source_relationship_variable e "
        "WHERE e.relationship_id = r.relationship_id) AS endpoints "
        "FROM source_relationship r "
        "LEFT JOIN variable v ON v.variable_id = r.owner_variable_id"
    ),
    "evidence": lambda o: [
        {"kind": event["evidence_kind"], "disposition": event["disposition"]}
        for event in o.events
        if event["kind"] == "prepared_evidence"
    ],
    "source_issues": _source_issues,
}


# The fields of each table, as cases/build/README.md documents them.
FIELDS: dict[str, frozenset[str]] = {
    name: frozenset(fields.split())
    for name, fields in {
        "issues": "code severity subject case_id locator detail acknowledged_by "
        "valid_from valid_to fields withheld_output",
        "issue_refs": "code severity subject case_id locator detail acknowledged_by "
        "valid_from valid_to fields withheld_output source key ref",
        "cases": "case_id status",
        "uses": "source native_variable key column_name data_type name description "
        "use variable",
        "case_uses": "case_id source key use variable",
        "states": "register variable variant column valid_from valid_to data_type "
        "name state_name provenance pooled data_length definition measurement_unit "
        "description operational_definition source_register_text "
        "value_set_version_label",
        "state_codes": "register variable variant column valid_from valid_to code label",
        "variables": "register variable column provider_key name definition "
        "description measurement_unit operational_definition is_identifier "
        "is_sensitive deprecated source_register source_label source_register_text",
        "variants": "register variant name panel_entity_key panel_time_key "
        "panel_time_grain description display_group",
        "registers": "register name purpose",
        "editions": "register variant name description measurement_information "
        "documentation_status first_approved_at last_approved_at populations "
        "object_types",
        "tags": "slug member",
        "aliases": "register variable variant column",
        "alias_windows": "register variable variant column valid_from valid_to provenance "
        "column_metadata coding_metadata data_type data_length definition "
        "measurement_unit name description operational_definition "
        "source_register_text codes classifications",
        "concept_groups": "variables",
        "group_axes": "group_key axis ordinal label",
        "group_members": "group_key register variable column facets",
        "warnings": "register variable variant column valid_from valid_to code "
        "severity detail summary detail_hash_of fields refs withheld_output "
        "acknowledged_by source_subject case_id",
        "search_pins": "query type position entity",
        "search_text": "register variable name definition description "
        "delivery_column_names",
        "manifest": "key value",
        "edges": "type a b",
        "successions": "grain predecessor successor predecessor_variant "
        "successor_variant predecessor_column successor_column variant "
        "effective_year note description",
        "classification_derivations": "derived source note",
        "timeseries_events": "name event description entity first_token "
        "second_token file_token",
        "source_columns": "table_name column_name sql_type nullable",
        "join_keys": "table_name column_name description",
        "identifiers": "native_variable name definition",
        "lineage": "register variable variant column source_register "
        "source_variable source_variant source_column valid_from valid_to",
        "lineage_warnings": "register variable variant column valid_from valid_to "
        "kind message",
        "state_classifications": "register variable column valid_from valid_to "
        "classification provenance",
        "conformance": "register variable column valid_from valid_to window "
        "classification status checked matched nonconforming overlap",
        "conformance_codes": "register variable column valid_from valid_to window "
        "classification code label member_kind sentinel_meaning scoped_windows",
        "code_index": "register variable code label mapping_count",
        "value_sets": "id members variables",
        "classifications": "slug short_name name name_en publisher valid_from "
        "valid_to description url code_count valid_code_count supersedes",
        "classification_successions": "predecessor successor effective_year note",
        "classification_codes": "slug code label level is_valid",
        "relationships": "kind binding_status source_dataset owner endpoints",
        "evidence": "kind disposition",
        "source_issues": "kind severity descriptor_key physical_associations refs",
    }.items()
}


def _matches(row: dict, where: dict) -> bool:
    for key, wanted in where.items():
        value = row[key]
        if isinstance(wanted, dict):
            if wanted.keys() != {"contains"}:
                raise ValueError(f"unknown where operator: {wanted}")
            texts = wanted["contains"]
            if value is None or not all(
                text in str(value)
                for text in ([texts] if isinstance(texts, str) else texts)
            ):
                return False
        elif isinstance(wanted, list):
            if value not in wanted:
                return False
        elif value != wanted:
            return False
    return True


def _sorted(rows: list[list]) -> list[list]:
    return sorted(rows, key=lambda row: json.dumps(row, ensure_ascii=False))


def _cited_refs(outcome: Outcome, selector: dict) -> list[str]:
    """Every source ref the rows of ``selector`` cite, in the `_refs` string form."""
    known = FIELDS[selector["table"]]
    where = selector.get("where", {})
    if selector.keys() - {"table", "where"} or not set(where) <= known:
        raise ValueError(f"unknown field in refs_of {selector}")
    if not known & {"ref", "refs"}:
        raise ValueError(f"refs_of names a table without refs: {selector}")
    return [
        ref
        for row in outcome.table(selector["table"])
        if _matches(row, where)
        for ref in (row["refs"] if "refs" in known else [row["ref"]])
    ]


def expected_rows(
    outcome: Outcome, spec: dict, earlier: dict[str, Outcome] | None = None
) -> list[list]:
    """A projection's expected rows: literal, the refs another projection cites, or
    the same projection of an earlier step of the case.

    `{"refs_of": {"table": ..., "where": ...}}` states a relation instead of values:
    the projection's one `ref` field must list exactly the refs those rows cite. Both
    sides are read from the same build, so no source ref is written as a literal.
    `{"step": "<earlier step>"}` expects the rows this projection (its table, `where`
    and `fields`) reads from that step's build, so a value a case must not write as a
    literal, such as a stored id, is still compared across builds.
    """
    rows = spec["rows"]
    if isinstance(rows, list):
        return rows
    if rows.keys() == {"step"}:
        if rows["step"] not in (earlier or {}):
            raise ValueError(f"no earlier step {rows['step']!r} in projection {spec}")
        found = project(earlier[rows["step"]], {**spec, "match": "exact"})["rows"]
        if not found and spec.get("match") == "includes":
            # Nothing to include would compare nothing.
            raise ValueError(f"earlier step reads no rows for projection {spec}")
        return found
    if rows.keys() != {"refs_of"} or spec["fields"] != ["ref"]:
        raise ValueError(f"unknown expected rows in projection {spec}")
    return [[ref] for ref in _cited_refs(outcome, rows["refs_of"])]


def project(outcome: Outcome, spec: dict) -> dict:
    """The actual value of one expected projection, shaped like ``spec``."""
    known = FIELDS[spec["table"]]
    where = spec.get("where", {})
    if not set(spec["fields"]) | set(where) <= known:
        raise ValueError(f"unknown field in projection {spec}")
    if any(isinstance(w, dict) and w.keys() != {"contains"} for w in where.values()):
        raise ValueError(f"unknown where operator in projection {spec}")
    if spec.get("match", "exact") not in {"exact", "set", "includes"}:
        raise ValueError(f"unknown match in projection {spec}")
    table = outcome.table(spec["table"])
    # Every row must carry exactly the documented fields, so an empty table
    # cannot hide a misspelled field.
    assert all(row.keys() == known for row in table), spec["table"]
    rows = [
        [row[field] for field in spec["fields"]]
        for row in table
        if _matches(row, spec.get("where", {}))
    ]
    match = spec.get("match", "exact")
    if match == "set":
        unique: list[list] = []
        for row in rows:
            if row not in unique:
                unique.append(row)
        rows = unique
    elif match == "includes":
        rows = [row for row in spec["rows"] if row in rows]
    elif match != "exact":
        raise ValueError(f"unknown match: {match}")
    return {**spec, "rows": _sorted(rows)}


# -- Running a step ---------------------------------------------------------------


def _subset(actual, expected):
    """``actual`` cut to the shape of ``expected``: only the keys it names, nested.

    An expected list names members that must be present, in its own order.
    """
    if isinstance(expected, dict) and isinstance(actual, dict):
        return {key: _subset(actual.get(key), value) for key, value in expected.items()}
    if isinstance(expected, list) and isinstance(actual, list):
        return [item for item in expected if item in actual]
    return actual


def _refusal(error: Exception, expected: dict) -> dict:
    """A refusal as `reg-meta-build build-db` and `check-curation` report it, cut to the
    keys ``expected`` names.

    Both commands wrap a `ValueError`, `OSError` or `KeyError` from the pipeline in the
    configuration error `pipeline_build_failed` (`cli._cmd_build_db`,
    `cli._cmd_check_curation`).
    """
    if isinstance(error, RegMetaError):
        refusal = {
            "code": error.code,
            "exit_code": error.exit_code,
            "message": error.message,
        }
    else:
        refusal = {
            "code": "pipeline_build_failed",
            "exit_code": EXIT_CONFIG,
            "message": str(error),
        }
    refusal["message_contains"] = [
        text for text in expected["message_contains"] if text in refusal["message"]
    ]
    return {key: refusal[key] for key in expected if key in refusal}


_REBUILD = """
import json, sys
from pathlib import Path
from reg_meta_build.pipeline import build_catalog
a = json.loads(sys.argv[1])
build_catalog(
    Path(a["prepared"]), a["commit"], a["digest"], Path(a["output"]),
    Path(a["report"]), curation_dir=Path(a["curation_dir"]),
    diagnostic=a["diagnostic"], registers=tuple(a["registers"]),
    dump_decisions=Path(a["dump_decisions"]),
)
"""


def _rebuild_in_fresh_process(args: dict) -> None:
    """Build again in a new interpreter under a different string-hash seed.

    The first build ran under this process's seed: random unless `PYTHONHASHSEED`
    fixes it. The rebuild takes seed 0, or 1 when this process already runs under 0,
    so the two builds never share a fixed seed and an output that depends on set or
    dict iteration order over strings shows up as a byte difference.
    """
    seed = "1" if os.environ.get("PYTHONHASHSEED") == "0" else "0"
    subprocess.run(
        [sys.executable, "-c", _REBUILD, json.dumps(args)],
        env={**os.environ, "PYTHONHASHSEED": seed},
        check=True,
        capture_output=True,
    )


def _tree_bytes(path: Path) -> dict[str, bytes]:
    """Every file under ``path`` (or ``path`` itself), by relative name.

    A missing or empty output is refused: two builds that both skip an output would
    otherwise compare equal and pass the rebuild check.
    """
    if path.is_file():
        return {"": path.read_bytes()}
    files = {
        str(file.relative_to(path)): file.read_bytes()
        for file in sorted(path.rglob("*"))
        if file.is_file()
    }
    if not files:
        raise AssertionError(f"rebuilt_identical: {path} holds no output to compare")
    return files


def run_step(
    step: Path,
    cache: PreparedCache,
    scratch: Path,
    earlier: dict[str, Outcome] | None = None,
) -> tuple[dict, dict]:
    """Run one case step; return ``(actual, expected)`` in the same shape.

    ``earlier`` maps the case's already-run step names to their builds. A step that
    builds a catalog joins it, so a later step's `{"step": ...}` rows can read it.
    """
    request = json.loads((step / "request.json").read_text(encoding="utf-8"))
    # A case states the product change that makes it fail, so a reviewer can check
    # that its oracle can fail at all (cases/build/README.md).
    fails_if = request.get("fails_if")
    if not isinstance(fails_if, str) or not fails_if.strip():
        raise ValueError(f"{step}/request.json: fails_if must be a non-empty string")
    expected = json.loads((step / "expected.json").read_text(encoding="utf-8"))
    built = cache.get(source_spec(request["sources"]))
    authored = cache.get(source_spec(request.get("authored_from", request["sources"])))
    curation = render_curation(step / "curation", scratch / "curation", authored)
    output, report = scratch / "reg_meta.db", scratch / "report"
    registers = tuple(request.get("registers", ()))
    rebuild = bool(expected.get("rebuilt_identical"))

    def build(output: Path, report: Path, decisions: Path | None) -> dict:
        return build_catalog(
            built.prepared,
            built.commit,
            built.digest,
            output,
            report,
            curation_dir=curation,
            diagnostic=request.get("diagnostic", True),
            registers=registers,
            dump_decisions=decisions,
        )

    try:
        if request.get("mode", "build") == "check":
            result = check_curation(
                built.prepared,
                built.commit,
                built.digest,
                report,
                curation_dir=curation,
                registers=registers,
            )
        else:
            result = build(output, report, scratch / "decisions" if rebuild else None)
    except (RegMetaError, ValueError, OSError, KeyError) as error:
        if "error" not in expected:
            raise
        return {"error": _refusal(error, expected["error"])}, expected
    actual: dict = {}
    if rebuild:
        again = scratch / "rebuild"
        again.mkdir()
        _rebuild_in_fresh_process(
            {
                "prepared": str(built.prepared),
                "commit": built.commit,
                "digest": built.digest,
                "output": str(again / "reg_meta.db"),
                "report": str(again / "report"),
                "curation_dir": str(curation),
                "diagnostic": request.get("diagnostic", True),
                "registers": list(registers),
                "dump_decisions": str(again / "decisions"),
            }
        )
        actual["rebuilt_identical"] = all(
            _tree_bytes(scratch / name) == _tree_bytes(again / name)
            for name in ("reg_meta.db", "report/events.jsonl.gz", "decisions")
        )
    if "error" in expected:
        actual["error"] = None
    if "status" in expected:
        actual["status"] = result["status"]
    if "result" in expected:
        actual["result"] = _subset(result, expected["result"])
    outcome = Outcome(result, report_events(report), output, built)
    if earlier is not None and request.get("mode", "build") == "build":
        earlier[step.name] = outcome
    if "projections" in expected:
        specs = [
            {**spec, "rows": expected_rows(outcome, spec, earlier)}
            for spec in expected["projections"]
        ]
        actual["projections"] = [project(outcome, spec) for spec in specs]
        expected = {
            **expected,
            "projections": [{**spec, "rows": _sorted(spec["rows"])} for spec in specs],
        }
    return actual, expected
