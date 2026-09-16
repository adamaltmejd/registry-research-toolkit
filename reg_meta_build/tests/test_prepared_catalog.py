"""Complete source preparation, exact accounting and a single warm trust boundary."""

from __future__ import annotations

import hashlib
import json
import subprocess
from typing import TYPE_CHECKING

import pytest
from _csv_fixtures import (
    repin_input_bundle,
    sparsify_scb_values,
    write_input_bundle,
    write_scb_input,
)
from _lisa_fixtures import write_lisa_workbook
from _prepared_fixtures import accept_prepared
from _sos_fixtures import write_sos_input
from openpyxl import Workbook, load_workbook
from reg_meta_build.input_snapshot import (
    CatalogBundleSelection,
    LisaWorkbookSelection,
    SnapshotMaterializationError,
)
from reg_meta_build.prepared_catalog import (
    PreparedCatalogError,
    open_prepared_catalog_sources,
    prepare_catalog_sources,
)
from reg_meta_build.source_records import SourceRecord

from reg_meta_build import _accepted_prepared, prepared_catalog

if TYPE_CHECKING:
    from pathlib import Path


def _selection(
    tmp_path: Path, *, all_roles: bool = False, validity: bool = True
) -> CatalogBundleSelection:
    source = tmp_path / "source"
    write_scb_input(
        source,
        include=(
            "registerinformation",
            "unika",
            "identifierare",
            "timeseries",
            *(("vardemangder", "valid_dates") if validity else ()),
        ),
    )
    curation = tmp_path / "curation"
    curation.mkdir()
    (curation / "tags.toml").write_text("")
    slugs = tmp_path / "slugs"
    slugs.mkdir()
    (slugs / "scb.toml").write_text("")
    lisa = None
    if all_roles:
        write_sos_input(source)
        workbook = write_lisa_workbook(tmp_path / "lisa.xlsx")
        lisa = LisaWorkbookSelection(
            workbook, "2024-2025", hashlib.sha256(workbook.read_bytes()).hexdigest()
        )
        thin = source / "Folkhalsomyndigheten"
        thin.mkdir()
        (thin / "fohm.toml").write_text(
            '[[register]]\nkey="sample"\nname="Sample"\nvalid_from="2000-01-01"\n[[register.variable]]\nname="Number"\ncolumn="N"\ndata_type="text"\n'
        )
        canonical = source / "scb_canonical"
        canonical.mkdir()
        (canonical / "scb_canonical.toml").write_text(
            '[[register]]\nkey="sample"\nname="Sample"\nvalid_from="2000-01-01"\n[[register.variable]]\nname="Code"\ncolumn="CODE"\ndata_type="text"\nvalue_set="sample-codes"\n'
        )
        (canonical / "sample-codes.csv").write_text(
            "code,label\n001,First\n001,First\n002,Other\n"
        )
        (source / "SCB/Tabelldefinitioner.sql").write_text(
            "CREATE TABLE [dbo].[Example](\n[A] [int] NULL\n) ON [PRIMARY]\nGO\n"
        )
        identifiers = Workbook()
        sheet = identifiers.active
        assert sheet is not None
        sheet.title = "Blad1"
        sheet.append(["Tabell", "ID-kolumn", "Beskrivning"])
        sheet.append(["Example", "A", "Declared identifier"])
        identifiers.save(source / "SCB/ID-kolumner.xlsx")
        identifiers.close()
    return write_input_bundle(
        tmp_path / "accepted",
        source,
        curation_dir=curation,
        slug_dir=slugs,
        lisa_workbook=lisa,
    )


def test_all_selected_source_roles_and_evidence_roundtrip(tmp_path: Path) -> None:
    selection = _selection(tmp_path, all_roles=True)
    destination = tmp_path / "prepared" / "catalog"
    manifest = prepare_catalog_sources(selection, destination)
    commit = accept_prepared(destination)
    prepared = open_prepared_catalog_sources(
        destination, expected_sha256=manifest.sha256, input_commit=commit
    )
    accounting = {(item.origin, item.path): item for item in manifest.inputs}
    original = json.loads((selection.path / "catalog-bundle.json").read_text())
    assert {path for origin, path in accounting if origin == "bundle"} == {
        item["path"] for item in original["files"]
    }
    assert {
        item.role for item in manifest.inputs if item.disposition == "prepared"
    } == {
        "scb_records",
        "scb_auxiliary",
        "scb_events",
        "scb_values",
        "scb_validity",
        "column_types",
        "join_keys",
        "sos_workbook",
        "lisa_workbook",
        "thin_provider",
        "code_list",
    }
    assert (
        accounting[("bundle", "curation/tags.toml")].disposition == "excluded_curation"
    )
    assert (
        accounting[("bundle", "fqid_slugs/scb.toml")].disposition == "excluded_naming"
    )
    assert (
        accounting[("bundle", "catalog/Forsakringskassan/fk.toml")].disposition
        == "absent"
    )
    assert {
        item.role for item in manifest.inputs if item.record_usage == "occurrence"
    } == {
        "scb_records",
        "sos_workbook",
        "thin_provider",
    }
    assert {
        item.role for item in manifest.inputs if item.record_usage == "support"
    } == {
        "scb_auxiliary",
        "lisa_workbook",
    }
    assert all(
        item.record_usage == "none" for item in manifest.inputs if not item.present
    )
    assert sum(item.counts.records for item in manifest.inputs) == len(
        list(prepared.records.records)
    )
    assert sum(item.counts.tables for item in manifest.inputs) == len(
        list(prepared.records.iter_tables())
    )
    evidence = list(prepared.iter_evidence())
    assert {item.kind for item in evidence} >= {
        "header",
        "reference",
        "worksheet_context",
    }
    assert {item.declaration.kind for item in evidence if item.kind == "reference"} >= {
        "event",
        "column_type",
        "join_key",
    }
    code_source = next(
        value
        for value in prepared.value_sources
        if value.manifest.revision.artifact_path.endswith("sample-codes.csv")
    )
    assert len(list(code_source.associations())) == 3
    assert [(value.code, value.label) for value in code_source.values()] == [
        ("001", "First"),
        ("002", "Other"),
    ]
    assert all(
        item.revision and len(item.revision.artifact_sha256) == 64
        for item in manifest.inputs
        if item.present
    )
    assert not (destination / "tables.jsonl").exists()


def test_sos_declarations_validity_support_and_formula_evidence_survive_preparation(
    tmp_path: Path,
) -> None:
    selection = _selection(tmp_path, all_roles=True)
    path = next((selection.path / "catalog/Socialstyrelsen").glob("*SYN*.xlsx"))
    workbook = load_workbook(path)
    crosswalk = workbook.create_sheet("Kodlista_CROSSWALK")
    crosswalk.append(
        ["Variabelnamn", "Tidsperiod", "Indata Kod", "Kodat till", "Beskrivning"]
    )
    crosswalk.append(["CODE", "1990/91-1994", "02", "03", "Recode declaration"])
    derivation = workbook.create_sheet("Kodlista_DERIVATION")
    derivation.append(
        [
            "Variabelnamn",
            "Tidsperiod",
            "Beskrivning",
            "Variabler",
            "Villkor",
            "Algoritm",
        ]
    )
    derivation.append(["RESULT", "1973-", "Documentation", "A, B", "A > 0", "B**2"])
    hospital = workbook.create_sheet("Kodlista_HOSPITAL")
    hospital.append(
        [
            "Variabelnamn",
            "Från",
            "Till",
            "Sjukhuskod",
            "Region",
            "Sjukhusnamn",
            "Aktuella",
            "Kommentar",
        ]
    )
    formula = '=IF(C2>1, " ","JA")'
    hospital.append(
        ["HOSPITAL", 1973, None, "001", "Region", "Hospital", formula, "Source note"]
    )
    support = workbook.create_sheet("Ej relevant_listor")
    support.sheet_state = "veryHidden"
    support.append(["Binär", "Datatyp"])
    workbook.save(path)
    workbook.close()
    bundle_path = selection.path / "catalog-bundle.json"
    document = json.loads(bundle_path.read_text())
    selected_file = next(
        item
        for item in document["files"]
        if item["path"] == path.relative_to(selection.path).as_posix()
    )
    selected_file.update(
        size=path.stat().st_size, sha256=hashlib.sha256(path.read_bytes()).hexdigest()
    )
    bundle_path.write_text(json.dumps(document))
    selection = repin_input_bundle(selection)

    destination = tmp_path / "prepared" / "catalog"
    manifest = prepare_catalog_sources(selection, destination)
    prepared = open_prepared_catalog_sources(
        destination,
        expected_sha256=manifest.sha256,
        input_commit=accept_prepared(destination),
    )
    entry = next(item for item in manifest.inputs if item.path == selected_file["path"])
    assert entry.counts.declarations == 2 and entry.counts.validity == 1
    references = [
        item.declaration
        for item in prepared.iter_evidence()
        if item.kind == "reference" and item.revision_id == entry.revision.revision_id
    ]
    assert {item.kind for item in references} == {"code_crosswalk", "derivation"}
    crosswalk = next(item for item in references if item.kind == "code_crosswalk")
    assert [(item.role, item.code.value) for item in crosswalk.operands] == [
        ("input", "02"),
        ("output", "03"),
    ]
    assert crosswalk.supplied_period.value == "1990/91-1994"
    derivation = next(item for item in references if item.kind == "derivation")
    assert derivation.clauses[-1].content.value == "B**2"
    values = next(
        item
        for item in prepared.value_sources
        if item.manifest.revision == entry.revision
    )
    (validity,) = values.validity()
    association = next(
        item
        for item in values.associations()
        if item.source_table == "Kodlista_HOSPITAL"
    )
    assert (validity.valid_from, validity.valid_to, validity.item_id) == (
        "1973",
        None,
        None,
    )
    assert validity.locator == association.locator
    cell = association.delivered_cells[6]
    assert (
        cell.raw_value,
        cell.storage_type,
        cell.cached_raw_value,
        cell.cached_raw_type,
    ) == (formula, "f", "", "none")
    tables = [
        table
        for table in prepared.records.iter_tables()
        if table.source_revision_id == entry.revision.revision_id
    ]
    source_table = next(table for table in tables if table.name == "Kodlista_HOSPITAL")
    assert source_table.rows[1].cells[6] == cell
    source_table = next(table for table in tables if table.name == "Ej relevant_listor")
    assert source_table.rows[0].role == "unparsed"
    assert [cell.raw_value for cell in source_table.rows[0].cells] == [
        "Binär",
        "Datatyp",
    ]


def test_warm_open_has_one_trust_check_and_never_reparses_inputs(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    selection = _selection(tmp_path)
    destination = tmp_path / "prepared" / "catalog"
    manifest = prepare_catalog_sources(selection, destination)
    commit = accept_prepared(destination)
    calls = []
    accepted = _accepted_prepared.clean_git_commit

    def checked(path):
        calls.append(path)
        return accepted(path)

    def forbidden(*_args, **_kwargs):
        raise AssertionError("warm opening cannot parse/hash original evidence")

    monkeypatch.setattr(_accepted_prepared, "clean_git_commit", checked)
    for name in (
        "open_input_bundle",
        "clean_scb_values",
        "iter_scb_observations",
        "read_lisa_source",
        "read_curated_source",
        "read_code_list",
        "parse_register_file",
        "_file_sha256",
    ):
        monkeypatch.setattr(prepared_catalog, name, forbidden)
    monkeypatch.setattr(SourceRecord, "create", forbidden)
    monkeypatch.setattr(SourceRecord, "_record_id", forbidden)
    prepared = open_prepared_catalog_sources(
        destination, expected_sha256=manifest.sha256, input_commit=commit
    )
    assert len(calls) == 1
    assert next(prepared.records.records).locators
    assert list(prepared.value_sources[0].associations())
    assert list(prepared.iter_evidence())


def test_preparation_replays_identical_bytes_and_does_not_rehash_child_data(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    selection = _selection(tmp_path)
    first, second = tmp_path / "first", tmp_path / "second"
    hashed = []
    original = prepared_catalog._file_sha256

    def small_only(path):
        hashed.append(path.name)
        assert path.name in {"manifest.json", "evidence.jsonl"}
        return original(path)

    monkeypatch.setattr(prepared_catalog, "_file_sha256", small_only)
    before = prepare_catalog_sources(selection, first)
    after = prepare_catalog_sources(selection, second)
    assert before == after and hashed
    assert {file.path: (first / file.path).read_bytes() for file in before.files} == {
        file.path: (second / file.path).read_bytes() for file in after.files
    }


def test_unsupported_or_incoherent_manifest_fails_before_opening_children(
    tmp_path: Path,
) -> None:
    selection = _selection(tmp_path)
    destination = tmp_path / "prepared" / "catalog"
    prepare_catalog_sources(selection, destination)
    path = destination / "manifest.json"
    document = json.loads(path.read_text())
    next(item for item in document["inputs"] if item["role"] == "scb_records")[
        "counts"
    ]["records"] += 1
    path.write_text(json.dumps(document))
    commit = accept_prepared(destination)
    with pytest.raises(PreparedCatalogError, match="counts differ"):
        open_prepared_catalog_sources(
            destination,
            expected_sha256=hashlib.sha256(path.read_bytes()).hexdigest(),
            input_commit=commit,
        )
    document["schema_version"] = 99
    path.write_text(json.dumps(document))
    commit = accept_prepared(destination)
    with pytest.raises(ValueError, match="schema_version"):
        open_prepared_catalog_sources(
            destination,
            expected_sha256=hashlib.sha256(path.read_bytes()).hexdigest(),
            input_commit=commit,
        )


@pytest.mark.parametrize("child", ["records", "values"])
def test_child_manifest_hash_must_match_its_outer_preparation_proof(
    tmp_path: Path, child: str
) -> None:
    selection = _selection(tmp_path)
    destination = tmp_path / "prepared" / "catalog"
    prepare_catalog_sources(selection, destination)
    path = destination / "manifest.json"
    document = json.loads(path.read_text())
    child_path = (
        "files/records" if child == "records" else document["values"][0]["path"]
    )
    next(
        item
        for item in document["files"]
        if item["path"] == f"{child_path}/manifest.json"
    )["sha256"] = "0" * 64
    path.write_text(json.dumps(document))
    commit = accept_prepared(destination)
    with pytest.raises(ValueError, match="child manifest SHA-256 differs"):
        open_prepared_catalog_sources(
            destination,
            expected_sha256=hashlib.sha256(path.read_bytes()).hexdigest(),
            input_commit=commit,
        )


def test_absent_validity_is_accounted_without_inventing_validity(
    tmp_path: Path,
) -> None:
    selection = _selection(tmp_path, validity=False)
    destination = tmp_path / "prepared" / "catalog"
    manifest = prepare_catalog_sources(selection, destination)
    absent = next(
        item for item in manifest.inputs if item.path == "VardemangderValidDates.csv"
    )
    assert absent.disposition == "absent" and absent.revision is None
    commit = accept_prepared(destination)
    prepared = open_prepared_catalog_sources(
        destination, expected_sha256=manifest.sha256, input_commit=commit
    )
    assert prepared.value_sources == ()


def test_preparation_rejects_unknown_selected_file_before_output(
    tmp_path: Path,
) -> None:
    selection = _selection(tmp_path)
    unknown = selection.path / "catalog/SCB/new-source.csv"
    unknown.parent.mkdir(parents=True, exist_ok=True)
    unknown.write_text("unexpected\n")
    path = selection.path / "catalog-bundle.json"
    document = json.loads(path.read_text())
    document["files"].append(
        {
            "path": "catalog/SCB/new-source.csv",
            "present": True,
            "size": unknown.stat().st_size,
            "sha256": hashlib.sha256(unknown.read_bytes()).hexdigest(),
        }
    )
    document["files"].sort(key=lambda item: item["path"])
    path.write_text(json.dumps(document))
    selection = repin_input_bundle(selection)
    destination = tmp_path / "prepared"
    with pytest.raises(PreparedCatalogError, match="no stage-1 input role"):
        prepare_catalog_sources(selection, destination)
    assert not destination.exists()


def test_selected_classification_lists_keep_their_provider_subdirectory(
    tmp_path: Path,
) -> None:
    selection = _selection(tmp_path)
    relative = "catalog/classifications/sos/atc.csv"
    source = selection.path / relative
    source.parent.mkdir(parents=True)
    source.write_text(
        "code,label,label_en,parent_code,valid_from,valid_to\n001,Etikett,Label,00,2001,2003\n"
    )
    path = selection.path / "catalog-bundle.json"
    document = json.loads(path.read_text())
    document["files"].append(
        {
            "path": relative,
            "present": True,
            "size": source.stat().st_size,
            "sha256": hashlib.sha256(source.read_bytes()).hexdigest(),
        }
    )
    document["files"].sort(key=lambda item: item["path"])
    path.write_text(json.dumps(document))
    selection = repin_input_bundle(selection)

    output = tmp_path / "prepared" / "catalog"
    manifest = prepare_catalog_sources(selection, output)
    entry = next(item for item in manifest.inputs if item.path == relative)
    assert entry.role == "code_list" and entry.counts.associations == 1
    prepared = open_prepared_catalog_sources(
        output, expected_sha256=manifest.sha256, input_commit=accept_prepared(output)
    )
    child = next(
        item
        for item in prepared.value_sources
        if item.manifest.revision == entry.revision
    )
    (value,) = child.values()
    assert (value.code, value.label) == ("001", "Etikett")
    assert value.raw_cells == ("001", "Etikett", "Label", "00", "2001", "2003")


def test_output_conflicts_and_cold_materialization_fail_without_overwrite(
    tmp_path: Path,
) -> None:
    selection = _selection(tmp_path)
    destination = tmp_path / "prepared"
    destination.mkdir()
    marker = destination / "keep"
    marker.write_text("original")
    with pytest.raises(PreparedCatalogError, match="already exists"):
        prepare_catalog_sources(selection, destination)
    assert marker.read_text() == "original"
    sparsify_scb_values(selection)
    missing = tmp_path / "cold"
    with pytest.raises(SnapshotMaterializationError):
        prepare_catalog_sources(selection, missing)
    assert not missing.exists()


def test_failed_source_preparation_cleans_staging_and_preserves_accepted_inputs(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    selection = _selection(tmp_path)
    original = (selection.path / "catalog-bundle.json").read_bytes()

    def invalid(*_args, **_kwargs):
        raise ValueError("malformed source contract")

    monkeypatch.setattr(prepared_catalog, "read_scb_events", invalid)
    destination = tmp_path / "new" / "prepared"
    with pytest.raises(ValueError, match="malformed source contract"):
        prepare_catalog_sources(selection, destination)
    assert not destination.exists() and list(destination.parent.iterdir()) == []
    assert (selection.path / "catalog-bundle.json").read_bytes() == original


def test_source_change_during_preparation_does_not_publish(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    selection = _selection(tmp_path)
    destination = tmp_path / "prepared"
    original = prepared_catalog.read_scb_events

    def change_source(*args, **kwargs):
        result = original(*args, **kwargs)
        (selection.path / "curation/tags.toml").write_text(
            "# changed after selection\n"
        )
        return result

    monkeypatch.setattr(prepared_catalog, "read_scb_events", change_source)
    with pytest.raises(PreparedCatalogError, match="changed during preparation"):
        prepare_catalog_sources(selection, destination)
    assert not destination.exists()


@pytest.mark.parametrize("flag", ["--assume-unchanged", "--skip-worktree"])
def test_source_index_flags_cannot_hide_changes_during_preparation(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch, flag: str
) -> None:
    selection = _selection(tmp_path)
    destination = tmp_path / "prepared"
    original = prepared_catalog.read_scb_events

    def hide_source_change(*args, **kwargs):
        result = original(*args, **kwargs)
        subprocess.run(
            [
                "git",
                "-C",
                str(selection.path),
                "update-index",
                flag,
                "curation/tags.toml",
            ],
            check=True,
        )
        (selection.path / "curation/tags.toml").write_text(
            "# changed but hidden from status\n"
        )
        return result

    monkeypatch.setattr(prepared_catalog, "read_scb_events", hide_source_change)
    with pytest.raises(PreparedCatalogError, match="source index flags changed"):
        prepare_catalog_sources(selection, destination)
    assert not destination.exists()
    assert not list(destination.parent.glob(".prepared.*"))


@pytest.mark.parametrize("change", ["untracked", "missing", "changed"])
def test_warm_inventory_rejects_unaccepted_or_stale_payloads(
    tmp_path: Path, change: str
) -> None:
    selection = _selection(tmp_path)
    destination = tmp_path / "prepared" / "catalog"
    manifest = prepare_catalog_sources(selection, destination)
    commit = accept_prepared(destination)
    payload = destination / "files/evidence.jsonl"
    if change == "untracked":
        (destination / "files/untracked").write_text("unaccepted")
    elif change == "missing":
        payload.unlink()
    else:
        payload.write_bytes(payload.read_bytes().replace(b"SCB", b"scb", 1))
        commit = accept_prepared(destination)
    with pytest.raises(ValueError):
        open_prepared_catalog_sources(
            destination, expected_sha256=manifest.sha256, input_commit=commit
        )
