"""Variable coverage periods read from Socialstyrelsen workbooks."""

from __future__ import annotations

from datetime import datetime
from typing import TYPE_CHECKING

import pytest
from _sos_fixtures import (
    CLASSIFICATION_URL as _CLASSIFICATION_URL,
    source_revision as _revision,
    write_source_workbook as _write_source_workbook,
)
from reg_meta_build.source_records import (
    value_field,
)
from reg_meta_build.sources.sos import parse_register_file
from reg_meta_build.sources.sos_records import (
    clean_sos_source,
    clean_sos_variable,
)

if TYPE_CHECKING:
    from pathlib import Path


@pytest.mark.parametrize(
    ("start", "end", "expected"),
    [
        # Excel calendar cells have no timezone.
        (datetime(2014, 2, 3), datetime(2015, 4, 5), ("2014-02-03", "2015-04-05")),  # noqa: DTZ001
        ("201402", "201503", ("2014-02-01", "2015-03-31")),
        ("20140203", "20150405", ("2014-02-03", "2015-04-05")),
        ("2014-02-03", "2015-04-05", ("2014-02-03", "2015-04-05")),
        ("2020-02-03", "2020", ("2020-02-03", "2020-12-31")),
        ("202002", "2020", ("2020-02-01", "2020-12-31")),
        ("2020", "20200203", ("2020-01-01", "2020-02-03")),
        ("2020", "2020", ("2020", "2020")),
        ("20140203", None, ("2014-02-03", None)),
        ("20140230", "20150405", None),
        ("20150405", "20140203", None),
        ("=2001", "2020", None),
    ],
)
def test_variable_coverage_preserves_explicit_date_bounds(
    tmp_path: Path, start: object, end: object, expected: tuple[str, str] | None
) -> None:
    import openpyxl

    path = tmp_path / "Metadata Patientregistret (PAR)_webb.xlsx"
    _write_source_workbook(path)
    workbook = openpyxl.load_workbook(path)
    sheet = workbook["Metadata - Variabelnivå"]
    sheet["H2"], sheet["I2"] = start, end
    workbook.save(path)
    register = parse_register_file(path)
    record = clean_sos_variable(register, register.variables[0], _revision(path))
    if expected is None:
        assert record.edition_scope.kind == "unknown"
    else:
        assert record.edition_scope.kind == "intervals"
        assert tuple((i.start, i.end) for i in record.edition_scope.intervals) == (
            expected,
        )
    assert record.fields.classification_declared == value_field(_CLASSIFICATION_URL)
    assert next(
        cell for cell in record.delivered_cells if cell.name == "Data från"
    ).raw_value == str(start)
    assert next(
        cell for cell in record.delivered_cells if cell.name == "Data till"
    ).raw_value == (str(end) if end is not None else "")


def _write_period_context_workbook(
    path: Path,
    *,
    register_period: str | None = "2005-07-01-",
    subset_end: object = None,
) -> None:
    import openpyxl

    _write_source_workbook(path)
    workbook = openpyxl.load_workbook(path)
    if register_period is not None:
        dcat = workbook.create_sheet("Metadata-Datamängd (DCAT-AP)")
        dcat.append(["Attribut", "Definition", "Svenska", "Engelska"])
        dcat.append(
            ["Tidsperiod", "Tidsmässig täckning", register_period, register_period]
        )
    subsets = workbook.create_sheet("Deldatamängder")
    subsets.append(
        ["Deldatamängdsnamn", "Deldatamängdsetikett", "Data från", "Data till"]
    )
    subsets.append(["PAR_OV", "Öppenvård", 2005, subset_end])
    workbook.save(path)


def _partial_record(path: Path):
    register = parse_register_file(path)
    variable = next(
        variable for variable in register.variables if variable.name == "PARTIELL"
    )
    return clean_sos_variable(register, variable, _revision(path))


def test_blank_end_reads_open_without_register_or_subset_context(
    tmp_path: Path,
) -> None:
    from reg_meta_build.source_intervals import scope_bounds

    # A supplied Data från with a delivered blank Data till reads as an open
    # end on its own: no register-level Tidsperiod and no Deldatamängder row
    # are required.
    path = tmp_path / "Metadata Patientregistret (PAR)_webb.xlsx"
    _write_period_context_workbook(path)
    record = _partial_record(path)
    assert record.edition_scope.kind == "intervals"
    assert tuple((i.start, i.end) for i in record.edition_scope.intervals) == (
        ("2010", None),
    )
    assert scope_bounds(record.edition_scope) is not None
    assert (
        next(
            cell for cell in record.delivered_cells if cell.name == "Data till"
        ).raw_value
        == ""
    )

    # A closed register period still reads the row's own open end.
    closed = tmp_path / "closed.xlsx"
    _write_period_context_workbook(closed, register_period="2005-07-01-2024-12-31")
    assert tuple(
        (i.start, i.end) for i in _partial_record(closed).edition_scope.intervals
    ) == (("2010", None),)

    # An absent register period still reads the row's own open end.
    absent = tmp_path / "absent.xlsx"
    _write_period_context_workbook(absent, register_period=None)
    assert tuple(
        (i.start, i.end) for i in _partial_record(absent).edition_scope.intervals
    ) == (("2010", None),)

    # A closed enclosing subset still reads the variable row's own open end.
    shut_subset = tmp_path / "shut.xlsx"
    _write_period_context_workbook(shut_subset, subset_end=2020)
    assert tuple(
        (i.start, i.end) for i in _partial_record(shut_subset).edition_scope.intervals
    ) == (("2010", None),)


def test_blank_end_open_without_any_period_context(tmp_path: Path) -> None:
    import openpyxl

    # Year-only start with a blank end and no register/subset context.
    path = tmp_path / "Metadata Patientregistret (PAR)_webb.xlsx"
    _write_source_workbook(path)
    record = _partial_record(path)
    assert record.edition_scope.kind == "intervals"
    assert tuple((i.start, i.end) for i in record.edition_scope.intervals) == (
        ("2010", None),
    )

    # Compact-date start with a blank end reads the parsed ISO start.
    workbook = openpyxl.load_workbook(path)
    workbook["Metadata - Variabelnivå"]["H6"] = "20050701"
    workbook["Metadata - Variabelnivå"]["I6"] = None
    workbook.save(path)
    compact = _partial_record(path)
    assert compact.edition_scope.kind == "intervals"
    assert tuple((i.start, i.end) for i in compact.edition_scope.intervals) == (
        ("2005-07-01", None),
    )

    # A supplied Data till that does not parse stays unknown.
    workbook = openpyxl.load_workbook(path)
    workbook["Metadata - Variabelnivå"]["H6"] = 2010
    workbook["Metadata - Variabelnivå"]["I6"] = "2017 (Malmö)"
    workbook.save(path)
    assert _partial_record(path).edition_scope.kind == "unknown"


def test_blank_end_open_rule_never_widens_start_or_end(tmp_path: Path) -> None:
    import openpyxl

    path = tmp_path / "Metadata Patientregistret (PAR)_webb.xlsx"
    _write_period_context_workbook(path)
    register = parse_register_file(path)
    revision = _revision(path)

    # An explicit end is unchanged under the open context.
    explicit = clean_sos_variable(register, register.variables[0], revision)
    assert tuple((i.start, i.end) for i in explicit.edition_scope.intervals) == (
        ("2001", "2020"),
    )

    # A blank start is not backfilled from the open context.
    workbook = openpyxl.load_workbook(path)
    workbook["Metadata - Variabelnivå"]["H6"] = None
    workbook.save(path)
    assert _partial_record(path).edition_scope.kind == "unknown"

    # A malformed end is not read as an open bound.
    workbook = openpyxl.load_workbook(path)
    workbook["Metadata - Variabelnivå"]["H6"] = 2010
    workbook["Metadata - Variabelnivå"]["I6"] = "+2001"
    workbook.save(path)
    assert _partial_record(path).edition_scope.kind == "unknown"


def test_trailing_dash_start_reads_explicit_open_end(tmp_path: Path) -> None:
    import openpyxl

    # No open register period and no Deldatamängder sheet: the dash alone
    # carries the open end.
    path = tmp_path / "Metadata Patientregistret (PAR)_webb.xlsx"
    _write_source_workbook(path)
    workbook = openpyxl.load_workbook(path)
    workbook["Metadata - Variabelnivå"]["H2"] = "1999-"
    workbook["Metadata - Variabelnivå"]["I2"] = None
    workbook.save(path)
    register = parse_register_file(path)
    record = clean_sos_variable(register, register.variables[0], _revision(path))
    assert record.edition_scope.kind == "intervals"
    assert tuple((i.start, i.end) for i in record.edition_scope.intervals) == (
        ("1999", None),
    )


def test_trailing_dash_subset_start_reads_open(tmp_path: Path) -> None:
    import openpyxl

    path = tmp_path / "Metadata Patientregistret (PAR)_webb.xlsx"
    _write_period_context_workbook(path)
    workbook = openpyxl.load_workbook(path)
    workbook["Deldatamängder"]["C2"] = "2007-"
    workbook["Deldatamängder"]["D2"] = None
    workbook.save(path)
    cleaned = clean_sos_source(parse_register_file(path), _revision(path))
    subset = next(
        record
        for record in cleaned.records
        if record.subject.member.status == "not_applicable"
        and record.subject.variant.name == "PAR_OV"
    )
    assert subset.edition_scope.kind == "intervals"
    assert tuple((i.start, i.end) for i in subset.edition_scope.intervals) == (
        ("2007", None),
    )
    # The undashed open-end reading sees the dashed subset as open too.
    partial = next(
        record for record in cleaned.records if record.subject.member.name == "PARTIELL"
    )
    assert partial.edition_scope.kind == "intervals"
    assert tuple((i.start, i.end) for i in partial.edition_scope.intervals) == (
        ("2010", None),
    )


@pytest.mark.parametrize(
    ("start", "end"),
    [
        # A supplied Data till together with a dashed start is a conflict.
        ("1999-", 2020),
        # Compound and free-text ranges stay unknown, dash or not.
        ("1999-2003 samt 2007-", None),
        ("Senaste tre år", None),
        ("2011 och 2013", None),
    ],
)
def test_trailing_dash_widens_nothing_else(
    tmp_path: Path, start: object, end: object
) -> None:
    import openpyxl

    path = tmp_path / "Metadata Patientregistret (PAR)_webb.xlsx"
    _write_source_workbook(path)
    workbook = openpyxl.load_workbook(path)
    workbook["Metadata - Variabelnivå"]["H2"] = start
    workbook["Metadata - Variabelnivå"]["I2"] = end
    workbook.save(path)
    register = parse_register_file(path)
    record = clean_sos_variable(register, register.variables[0], _revision(path))
    assert record.edition_scope.kind == "unknown"
