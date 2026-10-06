"""Literal source relationships, unbound code lists and source-finding ACKs at the build boundary.

Every case prepares a synthetic SCB register plus one SOS workbook through the real
input pipeline and asserts on the report ledger and the built SQLite artifact.
"""

from __future__ import annotations

import gzip
import json
import sqlite3
from dataclasses import dataclass
from typing import TYPE_CHECKING

import pytest
from _csv_fixtures import var_row, write_input_bundle, write_scb_input
from _prepared_fixtures import accept_prepared
from _sos_fixtures import DEFAULT_REGISTERS, write_sos_input
from openpyxl import load_workbook
from reg_meta.documentary import SourceCodeCrosswalkDeclaration
from reg_meta.source_evidence import canonical_sha256
from reg_meta_build.pipeline import build_catalog
from reg_meta_build.prepared_catalog import (
    ReferenceEvidence,
    open_prepared_catalog_sources,
    prepare_catalog_sources,
)
from reg_meta_build.source_naming import authored_naming_id

if TYPE_CHECKING:
    from pathlib import Path

SOS = "Socialstyrelsen/Metadata Syntetiskt register (SYN)_webb.xlsx"
SOS_REGISTER = "Syntetiskt register"
CROSSWALK = "Kodlista_CROSSWALK"
RETAINED_REASON = (
    "No supplied variable or encoding endpoint; retain literal codes unattached."
)
# The SCB family 1.101 is delivered under two columns; the partition below owns
# only VALUE, so the compiler reports ALTVALUE as an unassigned original column.
ALTVALUE_REF = json.dumps(
    {
        "source": "scb-registerinformation",
        "semantic_record_key": [
            "register:1",
            "variant:10",
            "edition:111",
            "variable:101",
            "member:1002",
        ],
    }
)


@dataclass(frozen=True)
class Sources:
    prepared: Path
    commit: str
    digest: str
    curation: Path

    def build(self, tmp_path: Path, label: str, registers: tuple[str, ...] = ()):
        report = tmp_path / f"{label}-report"
        output = tmp_path / f"{label}.db"
        result = build_catalog(
            self.prepared,
            self.commit,
            self.digest,
            output,
            report,
            curation_dir=self.curation,
            registers=registers,
            diagnostic=True,
        )
        with gzip.open(report / "events.jsonl.gz", "rt", encoding="utf-8") as stream:
            events = [json.loads(line) for line in stream]
        return result, events, output

    def append(self, register: str, text: str) -> None:
        path = self.curation / "registers" / register
        path.write_text(path.read_text(encoding="utf-8") + text, encoding="utf-8")


def _issues(events: list[dict], code: str) -> list[dict]:
    return [e for e in events if e["kind"] == "issue" and e["code"] == code]


def _crosswalk_dispositions(events: list[dict]) -> list[str]:
    return [
        e["disposition"]
        for e in events
        if e["kind"] == "prepared_evidence"
        and e["disposition"]
        in {"out_of_slice_relationship", "unbound_relationship", "retained_unattached"}
    ]


@pytest.fixture
def sources(tmp_path: Path) -> Sources:
    source = tmp_path / "source"
    write_scb_input(
        source,
        registerinformation_rows=[
            var_row(cvid=1001, var_id=101, colname="VALUE"),
            var_row(
                cvid=1002, var_id=101, colname="ALTVALUE", year="2021", regver_id=111
            ),
        ],
        unika_rows=[
            "TESTREG|Testregistret|Individer|Individer|GenericVar|VALUE|2020|2021|0|0|0"
        ],
        include=("registerinformation", "unika"),
    )
    workbook_path = next(
        write_sos_input(source, registers=DEFAULT_REGISTERS[:1]).glob("*.xlsx")
    )
    workbook = load_workbook(workbook_path)
    # One literal recode row (a code crosswalk declaration) and one code list
    # whose sheet names no variable it belongs to.
    crosswalk = workbook.create_sheet(CROSSWALK)
    crosswalk.append(
        ["Variabelnamn", "Tidsperiod", "Indata Kod", "Kodat till", "Beskrivning"]
    )
    crosswalk.append(["DIAGNOS", "2005-2015", "02", "03", "Recode declaration"])
    empty = workbook.create_sheet("Kodlista_EMPTY")
    empty.append(["Variabelnamn", "CODE"])
    empty.append(["Tidsperiod", "Kod", "Beskrivning"])
    workbook.save(workbook_path)
    workbook.close()
    bundle = write_input_bundle(tmp_path / "inputs", source)
    prepared = tmp_path / "prepared" / "catalog"
    manifest = prepare_catalog_sources(bundle, prepared)
    commit = accept_prepared(prepared)
    curation = tmp_path / "curation"
    (curation / "registers" / "scb").mkdir(parents=True)
    (curation / "registers" / "sos").mkdir()
    (curation / "classifications").mkdir()
    (curation / "registers/scb/sample.toml").write_text(
        '[register]\nprovider = "scb"\nslug = "sample"\nnative_id = "1"\n'
        '[[variant]]\nnative_id = "1.10"\nslug = "people"\n'
        '[[variable]]\nnative_id = "1.101.a"\nslug = "value"\n'
        '[[identity.partition]]\nvariable = "1.101"\n'
        'columns = { VALUE = "1.101.a" }\nunassigned_columns = ["ALTVALUE"]\n'
        'columns_ref = "exact fixture"\n',
        encoding="utf-8",
    )

    def sos_id(kind, member=None):
        return authored_naming_id(
            kind, provider="sos", register_key="syn", member_key=member
        )

    (curation / "registers/sos/syn.toml").write_text(
        '[register]\nprovider = "sos"\nslug = "syn"\n'
        f'name = "{SOS_REGISTER}"\nnative_id = "{sos_id("register")}"\n'
        + "".join(
            f'[[variant]]\nslug = "{slug}"\n'
            f'native_id = "{sos_id("register_variant", member)}"\n'
            for member, slug in (("SYN_A", "syn-a"), ("SYN_B", "syn-b"))
        )
        + "".join(
            f'[[variable]]\nslug = "{member.lower()}"\n'
            f'native_id = "{sos_id("variable", member)}"\n'
            for member in ("DIAGNOS", "KON")
        ),
        encoding="utf-8",
    )
    return Sources(prepared, commit, manifest.sha256, curation)


@pytest.mark.parametrize(
    ("registers", "deferred"),
    [
        pytest.param(("1",), True, id="sos-unselected"),
        pytest.param((SOS_REGISTER,), False, id="sos-selected"),
        pytest.param((), False, id="full"),
    ],
)
def test_unbound_crosswalk_defers_only_when_its_occurrence_source_is_unselected(
    sources: Sources, tmp_path: Path, registers: tuple[str, ...], deferred: bool
) -> None:
    _, events, _ = sources.build(tmp_path, "build", registers)
    crosswalk = [
        (e["code"], e["severity"], e["refs"][0]["source"])
        for e in events
        if e["kind"] == "issue" and e["subject"] == repr((f"sheet:{CROSSWALK}",))
    ]
    assert crosswalk == [
        ("deferred_out_of_slice_reference", "warning", SOS)
        if deferred
        else ("unbound_source_relationship", "error", SOS)
    ]
    assert _crosswalk_dispositions(events) == [
        "out_of_slice_relationship" if deferred else "unbound_relationship"
    ]


def test_retained_crosswalk_persists_literal_without_catalog_link(
    sources: Sources, tmp_path: Path
) -> None:
    prepared = open_prepared_catalog_sources(
        sources.prepared, expected_sha256=sources.digest, input_commit=sources.commit
    )
    declaration = next(
        e.declaration
        for e in prepared.iter_evidence()
        if isinstance(e, ReferenceEvidence)
        and isinstance(e.declaration, SourceCodeCrosswalkDeclaration)
    )
    table = next(
        t for t in prepared.records.iter_tables(source=SOS) if t.name == CROSSWALK
    )
    sources.append(
        "sos/syn.toml",
        f'\n[[documentary.retained]]\nsource = "{SOS}"\ntable = "{CROSSWALK}"\n'
        f'row = "{declaration.locator.physical_record}"\n'
        f'payload_sha256 = "{canonical_sha256(declaration.model_dump(mode="json"))}"\n'
        f'table_sha256 = "{canonical_sha256(table.model_dump(mode="json"))}"\n'
        f'reason = "{RETAINED_REASON}"\n'
        'evidence = "Complete exact table reviewed."\nnoted = "2026-10-02"\n',
    )
    _, events, output = sources.build(tmp_path, "retained")
    assert [
        (e["severity"], e["detail"])
        for e in _issues(events, "retained_unattached_source_relationship")
    ] == [("warning", RETAINED_REASON)]
    assert not _issues(events, "unbound_source_relationship")
    assert _crosswalk_dispositions(events) == ["retained_unattached"]
    with sqlite3.connect(output) as conn:
        assert conn.execute(
            "SELECT owner_variable_id, binding_status, kind, source_dataset "
            "FROM source_relationship"
        ).fetchall() == [(None, "retained_unattached", "code_crosswalk", SOS)]
        (raw,) = conn.execute(
            "SELECT declaration_json FROM source_relationship"
        ).fetchone()
        assert SourceCodeCrosswalkDeclaration.model_validate_json(raw) == declaration
        assert conn.execute(
            "SELECT COUNT(*) FROM source_relationship_variable"
        ).fetchone() == (0,)
        assert conn.execute(
            "SELECT r.slug, w.variable_id, json_extract(w.warning_json, '$.detail') "
            "FROM data_warning w JOIN register r USING (register_id) "
            "WHERE json_extract(w.warning_json, '$.code') = "
            "'retained_unattached_source_relationship'"
        ).fetchall() == [("syn", None, RETAINED_REASON)]
        assert not conn.execute("PRAGMA foreign_key_check").fetchall()


def test_unbound_code_list_issue_identity_is_its_exact_source_evidence(
    sources: Sources, tmp_path: Path
) -> None:
    _, events, _ = sources.build(tmp_path, "full")
    evidence = [e for e in events if e["kind"] == "value_source_issue"]
    assert {(e["descriptor_key"], e["physical_associations"]) for e in evidence} == {
        ("sheet:Kodlista_DIAGNOS", 2),
        ("sheet:Kodlista_EMPTY", 0),
    }
    expected = sorted(
        repr(
            (
                e["value_source"],
                e["revision_id"],
                e["descriptor_key"],
                canonical_sha256(e["descriptor"]),
                e["physical_associations"],
                tuple(e["raw_member_tokens"]),
            )
        )
        for e in evidence
    )
    issues = _issues(events, "unresolved_list_reference")
    assert {i["severity"] for i in issues} == {"error"}
    assert sorted(i["subject"] for i in issues) == expected
    assert all(e["revision_id"].startswith(f"{SOS}@sha256:") for e in evidence)


@pytest.mark.parametrize(
    ("ref", "acknowledged"),
    [
        pytest.param(ALTVALUE_REF, True, id="exact"),
        pytest.param(ALTVALUE_REF.replace("member:1002", "member:1001"), False),
    ],
)
def test_unassigned_original_column_finding_reaches_exact_ack(
    sources: Sources, tmp_path: Path, ref: str, acknowledged: bool
) -> None:
    sources.append(
        "scb/sample.toml",
        '\n[[acknowledge]]\ncode = "unassigned_original_columns"\n'
        f'subject = "1.101"\nrefs = [{json.dumps(ref)}]\n'
        'fields = ["identity", "column_name"]\n'
        'reason = "Preserve the unknown original column."\n'
        'evidence = "Exact reviewed source finding."\n',
    )
    _, events, _ = sources.build(tmp_path, "ack", ("1",))
    case_id = "curation/registers/scb/sample.toml#/acknowledge/1"
    assert [
        (i["severity"], i["acknowledged_by"], i["refs"])
        for i in _issues(events, "unassigned_original_columns")
    ] == [
        (
            "warning" if acknowledged else "error",
            case_id if acknowledged else None,
            [json.loads(ALTVALUE_REF)],
        )
    ]
    stale = [i["case_id"] for i in _issues(events, "stale_curation_entry")]
    assert stale == ([] if acknowledged else [case_id])
