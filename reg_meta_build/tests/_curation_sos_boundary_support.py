"""Synthetic SOS and thin-provider sources run through the real build for curation cases.

Every source is synthetic. One minimal SCB register (the snapshot every bundle needs)
is always present; each case adds SOS workbooks derived from the public
``_sos_fixtures.DEFAULT_REGISTERS`` shapes, or a thin-provider TOML, then prepares and
accepts the bundle and builds it with an authored curation tree. Cases assert on the
report event ledger and the built SQLite artifact only.
"""

from __future__ import annotations

import sqlite3
from dataclasses import dataclass, replace
from typing import TYPE_CHECKING

from _csv_fixtures import var_row, write_scb_input
from _pipeline_catalog_support import (
    prepare_accepted,
    report_events,
    write_curation_tree,
)
from _sos_fixtures import DEFAULT_REGISTERS, write_sos_input
from openpyxl import load_workbook
from reg_meta_build.pipeline import build_catalog
from reg_meta_build.source_naming import authored_naming_id

if TYPE_CHECKING:
    from collections.abc import Callable, Iterable
    from pathlib import Path

# The SYN default register carries a Deldatamängder sheet; SYT has none.
_WITH_SUBSETS, _WITHOUT_SUBSETS = DEFAULT_REGISTERS
_SUBSET = _WITH_SUBSETS.deldatamangder[0]
_VARIABLE = _WITH_SUBSETS.variables[0]

SCB_SAMPLE_TOML = (
    '[register]\nprovider = "scb"\nslug = "sample"\nnative_id = "1"\n'
    '[[variant]]\nnative_id = "1.10"\nslug = "people"\n'
    '[[variable]]\nnative_id = "1.101"\nslug = "value"\n'
)


def subset(name: str, *, label: str | None = None, aggregation: str | None = None):
    """A Deldatamängder row (2005-2015) derived from the SYN default subset."""
    return replace(
        _SUBSET,
        name=name,
        label=label if label is not None else name,
        description=None,
        data_from=2005,
        data_to=2015,
        aggregation_level=aggregation,
    )


def variable(
    name: str,
    token: str | None,
    *,
    data_type: str | None = "Sträng (text)",
    value_set: str | None = None,
    link: str | None = None,
    description: str | None = None,
):
    """A Variabelnivå row (2005-2015) derived from the SYN default variable."""
    return replace(
        _VARIABLE,
        name=name,
        deldatamangd=token,
        label=name.title(),
        description=description,
        data_type=data_type,
        data_from=2005,
        data_to=2015,
        value_set=value_set,
        external_classification=link,
    )


def write_workbook(
    source: Path,
    abbrev: str,
    title: str,
    variables: Iterable,
    subsets: Iterable = (),
) -> str:
    """Write one SOS workbook and return its source name.

    ``subsets`` empty writes no Deldatamängder sheet (the SYT shape). A blank
    ``Kopplingsvariabel`` column is added so every variable carries the delivered
    explicit "not a linkage variable" claim instead of an unknown flag.
    """
    shape = _WITH_SUBSETS if subsets else _WITHOUT_SUBSETS
    register = replace(
        shape,
        abbrev=abbrev,
        title_sv=title,
        description_sv=None,
        variables=tuple(variables),
        deldatamangder=tuple(subsets),
        kodlistor=(),
    )
    sos_dir = write_sos_input(source, registers=(register,))
    name = f"Metadata {title} ({abbrev})_webb.xlsx"
    workbook = load_workbook(sos_dir / name)
    sheet = workbook["Metadata - Variabelnivå"]
    sheet.cell(1, sheet.max_column + 1, "Kopplingsvariabel")
    workbook.save(sos_dir / name)
    workbook.close()
    return f"Socialstyrelsen/{name}"


def sos_register_toml(
    slug: str,
    name: str,
    *,
    variants: Iterable[str | tuple[str, str]] = (),
    variables: Iterable[str] = (),
    body: str = "",
) -> str:
    """A curated SOS register naming its delivered subsets and variable tokens."""

    def native(kind, member=None):
        return authored_naming_id(
            kind, provider="sos", register_key=slug, member_key=member
        )

    return (
        f'[register]\nprovider = "sos"\nslug = "{slug}"\nname = "{name}"\n'
        f'native_id = "{native("register")}"\n'
        + "".join(
            f'[[variant]]\nslug = "{slug}"\n'
            f'native_id = "{native("register_variant", member)}"\n'
            for member, slug in _named(variants)
        )
        + "".join(
            f'[[variable]]\nslug = "{slug}"\n'
            f'native_id = "{native("variable", member)}"\n'
            for member, slug in _named(variables)
        )
        + body
    )


def thin_register_toml(
    provider: str,
    key: str,
    *,
    variants: Iterable[str | tuple[str, str]] = (),
    variables: Iterable[str] = (),
    canonical_scb: bool = False,
) -> str:
    """A curated thin-provider register naming its authored variant and column keys."""

    def native(kind, member=None):
        return authored_naming_id(
            kind,
            provider=provider,
            register_key=key,
            member_key=member,
            canonical_scb=canonical_scb,
        )

    return (
        f'[register]\nprovider = "{provider}"\nslug = "{key}"\n'
        f'native_id = "{native("register")}"\n'
        + "".join(
            f'[[variant]]\nslug = "{slug}"\n'
            f'native_id = "{native("register_variant", member)}"\n'
            for member, slug in _named(variants)
        )
        + "".join(
            f'[[variable]]\nslug = "{slug}"\n'
            f'native_id = "{native("variable", member)}"\n'
            for member, slug in _named(variables)
        )
    )


def _named(members: Iterable[str | tuple[str, str]]) -> list[tuple[str, str]]:
    """(member, slug) pairs; a bare member slugs to its lower-cased ASCII words."""
    return [
        member
        if isinstance(member, tuple)
        else (
            member,
            member
            if member == "_default"
            else "-".join(
                "".join(c if c.isascii() and c.isalnum() else " " for c in member)
                .lower()
                .split()
            ),
        )
        for member in members
    ]


@dataclass(frozen=True)
class Build:
    result: dict
    events: list[dict]
    db: Path

    def issues(self, code: str | None = None) -> list[dict]:
        return [
            e
            for e in self.events
            if e["kind"] == "issue" and (code is None or e["code"] == code)
        ]

    def issue_cases(self, code: str) -> list[str]:
        return sorted(e["case_id"] for e in self.issues(code))

    def case_status(self) -> dict[str, str]:
        return {e["case_id"]: e["status"] for e in self.events if e["kind"] == "case"}

    def applied_cases(self) -> list[str]:
        """Every case id the ledger records on a source occurrence disposition."""
        return sorted(
            case
            for e in self.events
            if e["kind"] == "source_occurrence"
            for disposition in e["dispositions"]
            for case in disposition["cases"]
        )

    def rows(self, sql: str, *params) -> list[tuple]:
        with sqlite3.connect(self.db) as conn:
            return conn.execute(sql, params).fetchall()

    def states(self, register: str) -> list[tuple]:
        """(variable slug, variant slug, data_type, valid_from, valid_to) per state."""
        return sorted(
            self.rows(
                "SELECT v.slug, rv.slug, s.data_type, s.valid_from, s.valid_to "
                "FROM variable_state s JOIN variable v USING (variable_id) "
                "JOIN register_variant rv USING (register_variant_id) "
                "JOIN register r ON r.register_id = v.register_id WHERE r.slug = ?",
                register,
            )
        )


@dataclass(frozen=True)
class Prepared:
    path: Path
    commit: str
    digest: str

    def build(
        self,
        tmp_path: Path,
        label: str,
        curation: dict[str, str],
        *,
        registers: tuple[str, ...] = (),
    ) -> Build:
        """Build with an authored curation tree; ``curation`` maps register paths
        (``sos/x.toml``) to their text. The SCB sample register is always present."""
        root = write_curation_tree(
            tmp_path / f"{label}-curation",
            {
                f"registers/{relative}": text
                for relative, text in {
                    "scb/sample.toml": SCB_SAMPLE_TOML,
                    **curation,
                }.items()
            },
        )
        report = tmp_path / f"{label}-report"
        output = tmp_path / f"{label}.db"
        result = build_catalog(
            self.path,
            self.commit,
            self.digest,
            output,
            report,
            curation_dir=root,
            registers=registers,
            diagnostic=True,
        )
        return Build(result, report_events(report), output)


def prepare(tmp_path: Path, write_sources: Callable[[Path], object]) -> Prepared:
    """Prepare and accept the SCB sample plus whatever ``write_sources`` adds."""
    source = tmp_path / "source"
    write_scb_input(
        source,
        registerinformation_rows=[var_row(cvid=1001, var_id=101, colname="VALUE")],
        unika_rows=[
            "TESTREG|Testregistret|Individer|Individer|GenericVar|VALUE|2020|2020|0|0|0"
        ],
        include=("registerinformation", "unika"),
    )
    write_sources(source)
    return Prepared(*prepare_accepted(tmp_path, source))
