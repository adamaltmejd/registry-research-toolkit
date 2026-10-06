"""LISA workbook reader: delivered layouts, located layout drift and text cleaning."""

from __future__ import annotations

import hashlib
import re
import zipfile
from typing import TYPE_CHECKING

import pytest
from _lisa_fixtures import LISA_LAYOUT, QUALIFICATION_CONTEXT, write_lisa_workbook
from _source_inspection_fixtures import field_text
from openpyxl import load_workbook
from reg_meta.source_evidence import SourceField, SourceRevision
from reg_meta_build.input_snapshot import (
    LISA_DATASET_ID,
)
from reg_meta_build.source_records import (
    ScopeInterval,
    value_field,
)
from reg_meta_build.sources.lisa import LisaWorkbookError, read_lisa_source

if TYPE_CHECKING:
    from pathlib import Path


def _revision(path: Path, dataset: str = LISA_DATASET_ID) -> SourceRevision:
    payload = path.read_bytes()
    return SourceRevision.create(
        dataset=dataset,
        publisher="SCB",
        purpose="fixture source records",
        upstream_revision="fixture-2024-2025",
        artifact_path=path.name,
        artifact_size=len(payload),
        artifact_sha256=hashlib.sha256(payload).hexdigest(),
    )


def test_lisa_reader_preserves_four_layouts_sections_periods_and_occurrences(
    tmp_path: Path,
) -> None:
    path = write_lisa_workbook(tmp_path / "lisa.xlsx")
    source = read_lisa_source(path, _revision(path))
    records = source.records

    assert len(records) == 9
    assert {record.locators[0].physical_table for record in records} == {
        "Individ",
        "Individ årsoberoende",
        "Företag",
        "Arbetsställe",
    }
    person_numbers = [
        record for record in records if field_text(record, "column_name") == "PersonNr"
    ]
    assert len(person_numbers) == 4
    assert len({record.record_id for record in person_numbers}) == 4
    assert {record.context[1] for record in person_numbers} == {
        "Födelseland",
        "Avlidna",
        "In- och Utvandring",
    }
    assert all(
        record.edition_scope.kind == "year_independent" for record in person_numbers
    )

    rekr = next(
        record for record in records if field_text(record, "column_name") == "RekrBidr"
    )
    assert rekr.original_period_text == "2003-2008\n2013"
    assert [(item.start, item.end) for item in rekr.edition_scope.intervals] == [
        ("2003", "2008"),
        ("2013", "2013"),
    ]
    workplace = next(
        record
        for record in records
        if field_text(record, "column_name") == "Ast_LoneSum"
    )
    assert [(item.start, item.end) for item in workplace.edition_scope.intervals] == [
        ("1990", "2021"),
        ("2024", "2024"),
    ]
    company = next(
        record
        for record in records
        if field_text(record, "column_name") == "ForetagKod"
    )
    assert company.edition_scope.intervals == (ScopeInterval(start="2024", end="2024"),)
    assert company.subject.variant.name == "company"

    ampoltyp = next(
        record for record in records if field_text(record, "column_name") == "AmPolTyp"
    )
    assert ampoltyp.locators[0].physical_record == "row:600"
    assert ampoltyp.subject.variable.native_id == "AmPolTyp"
    assert ampoltyp.subject.variable.name == "AmPolTyp"
    assert field_text(ampoltyp, "description") == (
        "Typ av arbetsmarknadspolitisk åtgärd"
    )
    assert ampoltyp.fields.sensitivity == SourceField(
        status="value", value=False, raw_value="Nej"
    )
    assert ampoltyp.fields.availability == value_field(True)

    ku2 = next(
        record
        for record in records
        if field_text(record, "column_name") == "KU2YrkStalln"
    )
    assert ku2.fields.sensitivity == SourceField(
        status="value", value=True, raw_value="I vissa fall"
    )
    assert ku2.context[-2:] == (
        (
            "continuation-note: Före 2010 finns inte kod 5 "
            "(företagare i eget AB). För åren 1993- finns istället variabeln "
            "KU1Faman som anges som 1 om personen är en företagare i eget AB."
        ),
        "continuation-note: RAMS-Jobb",
    )
    assert ku2.locators[0].physical_cells[-2:] == ("Individ!B323", "Individ!F323")
    recognized_context = tuple(
        item
        for item in source.worksheet_context
        if item.startswith("worksheet-context ")
    )
    expected_context = {
        f"worksheet-context {sheet}!A{row}: {text}"
        for sheet, spec in LISA_LAYOUT.items()
        for row, text in spec["text_rows"].items()
    }
    assert set(recognized_context) == expected_context
    assert len(recognized_context) == len(expected_context)
    assert set(recognized_context) >= QUALIFICATION_CONTEXT

    footnotes = tuple(
        item
        for item in source.worksheet_context
        if item.startswith("worksheet-footnote ")
    )
    assert len(footnotes) == 5
    assert footnotes[0] == ("worksheet-footnote Individ!B815: _ftnref2")
    assert footnotes[-1].startswith(
        "worksheet-footnote Individ!B819: 7 Från och med årgång 2020"
    )
    assert all(
        "Individ!B815" not in record.locators[0].physical_cells for record in records
    )
    assert all(
        all(
            not item.startswith(("worksheet-context ", "worksheet-footnote "))
            for item in record.context
        )
        for record in records
    )


@pytest.mark.parametrize(
    ("mutation", "expected"),
    (
        ("shift-header", r"Individ!A3:F3"),
        ("change-section", r"Individ årsoberoende!A30"),
        ("section-qualifier", r"Individ!A5:F5"),
        ("bad-period", r"Individ!C600"),
        ("extra-column", r"expected 6 columns, got 7"),
        ("valid-control", None),
    ),
)
def test_lisa_reader_rejects_changed_actual_layout_with_coordinates(
    tmp_path: Path, mutation: str, expected: str | None
) -> None:
    path = write_lisa_workbook(tmp_path / "lisa.xlsx")
    workbook = load_workbook(path)
    if mutation == "shift-header":
        sheet = workbook["Individ"]
        for column in range(1, 7):
            sheet.cell(4, column, sheet.cell(3, column).value)
            sheet.cell(3, column).value = None
    elif mutation == "change-section":
        workbook["Individ årsoberoende"]["A30"] = "Ny tabell"
    elif mutation == "section-qualifier":
        workbook["Individ"]["B5"] = "Unexpected qualifier"
    elif mutation == "bad-period":
        workbook["Individ"]["C600"] = "1990 till 2024"
    elif mutation == "extra-column":
        workbook["Individ"]["G3"] = "Ny kolumn"
    else:
        assert mutation == "valid-control"
    workbook.save(path)
    workbook.close()

    if expected is None:
        source = read_lisa_source(path, _revision(path))
        assert "worksheet-context Individ!A5: Demografiska variabler" in (
            source.worksheet_context
        )
        return

    with pytest.raises(LisaWorkbookError, match=expected):
        read_lisa_source(path, _revision(path))


@pytest.mark.parametrize(
    ("cell", "value", "expected"),
    (
        ("A322", "Other", r"Individ!A323"),
        ("B323", "changed continuation", r"Individ!A323:F323"),
        ("B816", "changed footer", r"Individ!A816:F816"),
        ("B324", "arbitrary unkeyed text", r"Individ!A324:F324"),
    ),
)
def test_lisa_reader_accepts_only_observed_unkeyed_context_rows(
    tmp_path: Path, cell: str, value: str, expected: str
) -> None:
    path = write_lisa_workbook(tmp_path / "lisa.xlsx")
    workbook = load_workbook(path)
    workbook["Individ"][cell] = value
    workbook.save(path)
    workbook.close()

    with pytest.raises(LisaWorkbookError, match=expected):
        read_lisa_source(path, _revision(path))


def test_lisa_cleaning_uses_shared_text_rules_without_changing_code_spelling(
    tmp_path: Path,
) -> None:
    path = write_lisa_workbook(tmp_path / "lisa.xlsx")
    workbook = load_workbook(path)
    sheet = workbook["Individ"]
    raw_description = "Åtga\u0308rd\r\n\r\n  - Förklaring\u00a0 "
    sheet["A600"] = " A\u030a_01 "
    sheet["B600"] = raw_description
    sheet["E600"] = " I\u00a0  vissa\tfall "
    sheet["F600"] = " Arbets\u00a0 förmedlingen "
    workbook.save(path)
    workbook.close()

    source = read_lisa_source(path, _revision(path))
    record = next(
        record for record in source.records if record.fields.column_name.value == "Å_01"
    )

    assert record.fields.column_name.raw_value == " A\u030a_01 "
    assert record.fields.description.value == "Åtgärd\n\n  - Förklaring"
    # The workbook XML reader normalizes CRLF before returning the cell text.
    assert record.fields.description.raw_value == raw_description.replace("\r\n", "\n")
    assert record.fields.sensitivity.value is True
    assert record.fields.sensitivity.raw_value == " I\u00a0  vissa\tfall "
    assert record.fields.base_register.value == "Arbets förmedlingen"


def test_lisa_reader_does_not_trust_stale_sheet_dimensions(tmp_path: Path) -> None:
    # A workbook whose stored <dimension> understates the used range still reads
    # completely: the reader loads the cells, not the declared extent.
    path = write_lisa_workbook(tmp_path / "lisa.xlsx")
    expected = [
        record.record_id for record in read_lisa_source(path, _revision(path)).records
    ]
    stale = tmp_path / "stale.xlsx"
    with (
        zipfile.ZipFile(path) as source,
        zipfile.ZipFile(stale, "w", zipfile.ZIP_DEFLATED) as target,
    ):
        for info in source.infolist():
            data = source.read(info)
            if info.filename.startswith("xl/worksheets/sheet"):
                data = re.sub(
                    rb'<dimension ref="[^"]*"\s*/>', b'<dimension ref="A1"/>', data
                )
            target.writestr(info, data)

    records = read_lisa_source(stale, _revision(stale)).records

    assert [record.record_id for record in records] == expected
