"""Write an `.xlsx` delivery from a readable JSON workbook spec.

`cases/prepare/README.md` ("Workbook spec") is the format. A spec lists sheets as
rows of raw cells, the way a provider delivers them, plus the few cell properties a
reader distinguishes: number formats, hyperlinks, bold section headings, calendar
cells, hidden sheets, merged ranges, and two producer quirks openpyxl cannot write
(a formula's cached result and a numeric cell stored as `7.0`, rewritten in the
sheet XML after saving). A spec may start from an existing workbook, so a case states
only how its delivery departs from a shared base.
"""

from __future__ import annotations

import shutil
import xml.etree.ElementTree as ET
import zipfile
from datetime import datetime
from typing import TYPE_CHECKING, Any

import openpyxl
from openpyxl.styles import Font
from openpyxl.worksheet.hyperlink import Hyperlink

if TYPE_CHECKING:
    from pathlib import Path

    from openpyxl.cell import Cell

_MAIN = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
_CELL_KEYS = frozenset(
    {"value", "number_format", "hyperlink", "location", "bold", "datetime"}
    | {"cached", "stored"}
)
_SPEC_KEYS = frozenset(
    {"file", "from", "sheets", "cells", "remove", "rename", "order", "state"}
    | {"merge", "dimension"}
)


def write_workbook(spec: dict[str, Any], path: Path, *, base: Path | None) -> Path:
    """Write ``spec`` to ``path``, starting from the workbook ``base`` if given.

    The operations apply in a fixed order: remove, rename, sheets (rows appended,
    a sheet created where absent), cells, order, state, merge; then the sheet-XML
    rewrites (`cached`, `stored`, `dimension`).
    """
    if unknown := sorted(spec.keys() - _SPEC_KEYS):
        raise ValueError(f"unknown workbook spec keys {unknown}")
    path.parent.mkdir(parents=True, exist_ok=True)
    if base is not None and not spec.keys() - {"file", "from"}:
        # Nothing to change: deliver the base's bytes, not an openpyxl re-save
        # (which drops the inline-string type of an empty text cell).
        shutil.copyfile(base, path)
        return path
    if base is None:
        workbook = openpyxl.Workbook()
        workbook.remove(workbook.active)
    else:
        workbook = openpyxl.load_workbook(base)
    for name in spec.get("remove", ()):
        del workbook[name]
    for old, new in spec.get("rename", {}).items():
        workbook[old].title = new
    rewrites: dict[str, dict[str, dict[str, Any]]] = {}
    for name, rows in spec.get("sheets", {}).items():
        sheet = workbook[name] if name in workbook else workbook.create_sheet(name)
        start = 1 if _empty(sheet) else sheet.max_row + 1
        for offset, row in enumerate(rows):
            for column, cell in enumerate(row, start=1):
                target = sheet.cell(start + offset, column)
                _write_cell(target, cell, rewrites.setdefault(name, {}))
    for reference, cell in spec.get("cells", {}).items():
        name, coordinate = reference.rsplit("!", 1)
        _write_cell(workbook[name][coordinate], cell, rewrites.setdefault(name, {}))
    for offset, name in enumerate(spec.get("order", ())):
        workbook.move_sheet(name, offset=offset - workbook.sheetnames.index(name))
    for name, state in spec.get("state", {}).items():
        workbook[name].sheet_state = state
    for name, ranges in spec.get("merge", {}).items():
        for cell_range in ranges:
            workbook[name].merge_cells(cell_range)
    order = list(workbook.sheetnames)
    workbook.save(path)
    workbook.close()
    dimensions = spec.get("dimension", {})
    if any(rewrites.values()) or dimensions:
        _rewrite_sheets(path, order, rewrites, dimensions)
    return path


def _empty(sheet) -> bool:
    return sheet.max_row == 1 and sheet.max_column == 1 and sheet["A1"].value is None


def _write_cell(target: Cell, cell: Any, rewrites: dict[str, dict[str, Any]]) -> None:
    if not isinstance(cell, dict):
        target.value = cell
        return
    if unknown := sorted(cell.keys() - _CELL_KEYS):
        raise ValueError(f"unknown cell keys {unknown} at {target.coordinate}")
    target.value = (
        datetime.fromisoformat(cell["datetime"])
        if "datetime" in cell
        else cell.get("value")
    )
    if "number_format" in cell:
        target.number_format = cell["number_format"]
    if "hyperlink" in cell:
        target.hyperlink = cell["hyperlink"]
    if "location" in cell:
        # An internal link: a sheet reference, no external target.
        target.hyperlink = Hyperlink(
            ref=target.coordinate, location=cell["location"], display=target.value
        )
    if cell.get("bold"):
        target.font = Font(bold=True)
    if "cached" in cell or "stored" in cell:
        rewrites[target.coordinate] = cell


def _rewrite_sheets(
    path: Path,
    order: list[str],
    rewrites: dict[str, dict[str, dict[str, Any]]],
    dimensions: dict[str, str],
) -> None:
    """Rewrite cell values and dimensions in the saved sheet XML.

    openpyxl names the sheet parts `sheet1.xml`, `sheet2.xml`, ... in workbook order.
    """
    ET.register_namespace("", _MAIN)
    parts = {
        f"xl/worksheets/sheet{index}.xml": name
        for index, name in enumerate(order, start=1)
    }
    with zipfile.ZipFile(path) as source:
        members = [(info, source.read(info)) for info in source.infolist()]
    staged = path.with_suffix(".rewrite")
    with zipfile.ZipFile(staged, "w", zipfile.ZIP_DEFLATED) as target:
        for info, payload in members:
            name = parts.get(info.filename)
            if name is not None and (rewrites.get(name) or name in dimensions):
                payload = _rewrite_sheet(
                    payload, rewrites.get(name, {}), dimensions.get(name)
                )
            target.writestr(info, payload)
    shutil.move(staged, path)


def _rewrite_sheet(
    payload: bytes, cells: dict[str, dict[str, Any]], dimension: str | None
) -> bytes:
    document = ET.fromstring(payload)
    namespace = {"s": _MAIN}
    if dimension is not None:
        found = document.find("s:dimension", namespace)
        assert found is not None, "sheet has no stored dimension"
        found.set("ref", dimension)
    for element in document.iterfind(".//s:c", namespace):
        cell = cells.get(element.get("r", ""))
        if cell is None:
            continue
        value = element.find("s:v", namespace)
        if value is None:
            value = ET.SubElement(element, f"{{{_MAIN}}}v")
        if "cached" in cell:
            # A formula's delivered result: openpyxl writes the formula only.
            element.set("t", "str")
            value.text = cell["cached"]
        else:
            # The number exactly as another producer stores it, such as `7.0`.
            value.text = cell["stored"]
    return ET.tostring(document, xml_declaration=True, encoding="UTF-8")
