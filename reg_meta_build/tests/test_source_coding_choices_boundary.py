"""Checked coding curation at the build-catalog boundary.

Every case prepares synthetic sources through the real input pipeline, writes
register-file curation, builds a diagnostic catalog and asserts on the built SQLite
artifact and the report ledger.

The SCB fixture is one register (native id 1, variant 1.10) whose Registerinformation
rows and Vardemangder code lists are synthetic, modelled on the deleted IR tests in
`test_source_coding_choices.py` (a column delivered under two competing code lists
"keep" and "other"; a column whose earlier edition was exported with a blank column
name; a preliminary edition superseded by its final edition). The SOS fixture is the
synthetic SYN workbook from `_sos_fixtures` with its KON `Värdemängd` cell replaced by the two decimal-comma enumerations of the
deleted test (themselves modelled on retained Socialstyrelsen cells).

Guard digests and source-authority blocks are authored the way a curator authors
them: computed from the accepted prepared records with the public capture helpers
(`acknowledgement_evidence_sha256`, `capture_expectations`, the value-binding
session). The oracles are independent of that authoring: which codes and labels a
state carries, which period it covers, whether a data warning is written and which
curation entry the ledger reports stale.
"""

from __future__ import annotations

import json
import re
import sqlite3
from dataclasses import dataclass, replace
from typing import TYPE_CHECKING, Literal

import pytest
from _csv_fixtures import var_row, write_input_bundle, write_scb_input
from _pipeline_catalog_support import report_issues
from _prepared_fixtures import accept_prepared
from _sos_fixtures import DEFAULT_REGISTERS, write_sos_input
from openpyxl import load_workbook
from reg_meta.errors import RegMetaError
from reg_meta_build.curation_tree import load_curation_tree
from reg_meta_build.pipeline import build_catalog
from reg_meta_build.prepared_catalog import (
    open_prepared_catalog_sources,
    prepare_catalog_sources,
)
from reg_meta_build.source_coding import coding_source_sha256
from reg_meta_build.source_curation import (
    acknowledgement_evidence_sha256,
    capture_expectations,
)
from reg_meta_build.source_naming import authored_naming_id
from reg_meta_build.source_records import SourceFields
from reg_meta_build.source_value_bindings import bind_code_lists, open_value_bindings

if TYPE_CHECKING:
    from pathlib import Path

    from reg_meta_build.source_records import SourceRecord

SCB = "scb-registerinformation"
SCB_REGISTER = (
    '[register]\nprovider = "scb"\nslug = "sample"\nnative_id = "1"\n'
    '[[variant]]\nnative_id = "1.10"\nslug = "people"\n'
    '[[variable]]\nnative_id = "1.101"\nslug = "value"\n'
)
UNIKA = ["TESTREG|Testregistret|Individer|Individer|GenericVar|VALUE|2010|2020|0|0|0"]
# Two competing complete lists for CVID 1001 (VALUE 2020) and two for CVID 1002
# (OTHER 2020). Item validity covers every fixture year.
VALUE_LISTS = ["keep|1|01|Label|1001|7001", "other|1|02|Other|1001|7002"]
OTHER_LISTS = ["keep2|1|11|Kept|1002|7011", "other2|1|12|Dropped|1002|7012"]
VALIDITY = [
    f"{item}|2000-01-01|2030-12-31" for item in ("7001", "7002", "7003", "7011", "7012")
]


@dataclass(frozen=True)
class Prepared:
    prepared: Path
    commit: str
    digest: str

    def records(self, source: str) -> tuple[SourceRecord, ...]:
        opened = open_prepared_catalog_sources(
            self.prepared, expected_sha256=self.digest, input_commit=self.commit
        )
        return tuple(r for r in opened.records.records if r.source == source)

    def evidence_sha256(self, records: tuple[SourceRecord, ...]) -> str:
        """The complete-evidence guard a curator pins for these column originals."""
        opened = open_prepared_catalog_sources(
            self.prepared, expected_sha256=self.digest, input_commit=self.commit
        )
        with open_value_bindings(opened.value_sources) as sessions:
            claims = tuple(
                claim
                for record in records
                for claim in bind_code_lists(record, sessions).claims
            )
        return acknowledgement_evidence_sha256(
            records, (coding_source_sha256(claim) for claim in claims)
        )

    def build(self, tmp_path: Path, label: str, curation: Path):
        output = tmp_path / f"{label}.db"
        report = tmp_path / f"{label}-report"
        result = build_catalog(
            self.prepared,
            self.commit,
            self.digest,
            output,
            report,
            curation_dir=curation,
            diagnostic=True,
        )
        assert result["status"] == "diagnostic_complete"
        return output, report_issues(report)


def _prepare(tmp_path: Path, label: str, source: Path) -> Prepared:
    bundle = write_input_bundle(tmp_path / label / "inputs", source)
    prepared = tmp_path / label / "prepared" / "catalog"
    manifest = prepare_catalog_sources(bundle, prepared)
    return Prepared(prepared, accept_prepared(prepared), manifest.sha256)


def prepare_scb(
    tmp_path: Path,
    label: str,
    rows: list[str],
    lists: list[str] | None = None,
    summaries: tuple[str, ...] = (),
) -> Prepared:
    source = tmp_path / label / "source"
    write_scb_input(
        source,
        registerinformation_rows=rows,
        vardemangder_rows=lists or [],
        unika_rows=[*UNIKA, *summaries],
        valid_dates_rows=VALIDITY,
        include=("registerinformation", "unika")
        + (("vardemangder", "valid_dates") if lists else ()),
    )
    return _prepare(tmp_path, label, source)


def scb_curation(tmp_path: Path, label: str, body: str = "") -> Path:
    curation = tmp_path / label / "curation"
    (curation / "registers" / "scb").mkdir(parents=True)
    (curation / "classifications").mkdir()
    (curation / "registers" / "scb" / "sample.toml").write_text(
        SCB_REGISTER + body, encoding="utf-8"
    )
    return curation


def state_codes(db: Path, column: str) -> list[tuple[str, str, str | None, str | None]]:
    with sqlite3.connect(db) as conn:
        return conn.execute(
            "SELECT s.valid_from, s.valid_to, c.code, c.label FROM variable_state s "
            "LEFT JOIN value_set_member m ON m.value_set_id = s.value_set_id "
            "LEFT JOIN value_code c ON c.code_id = m.code_id "
            "WHERE s.delivery_column_name = ? ORDER BY 1, 3",
            (column,),
        ).fetchall()


def member(records: tuple[SourceRecord, ...], cvid: str) -> tuple[SourceRecord, ...]:
    return tuple(
        r for r in records if r.locators[0].semantic_record_key[-1] == f"member:{cvid}"
    )


def stale_entries(issues: list[dict]) -> list[str]:
    return sorted(
        issue["case_id"] for issue in issues if issue["code"] == "stale_curation_entry"
    )


def issue_codes(issues: list[dict], subject: str) -> list[str]:
    return sorted(issue["code"] for issue in issues if subject in str(issue["subject"]))


def choice(variable: str, column: str, keep: str, over: str, digest: str) -> str:
    return (
        f'\n[[coding.choice]]\nvariable = "{variable}"\nvariant = "people"\n'
        f'column = "{column}"\n'
        'periods = [["2020-01-01", "2020-06-30"], ["2020-07-01", "2020-12-31"]]\n'
        f'keep = "{keep}"\nover = ["{over}"]\n'
        f'expected_evidence_sha256 = "{digest}"\n'
        'reason = "Reviewed competing lists"\nsource = "fixture"\n'
    )


VALUE_2020 = var_row(cvid=1001, var_id=101, colname="VALUE")
OTHER_2020 = var_row(cvid=1002, var_id=101, colname="OTHER")
VALUE_CHOICE = "curation/registers/scb/sample.toml#/coding.choice/1/period/"


# -- Compact coding choices (raw source fingerprints) ---------------------------


def test_competing_lists_without_a_choice_publish_no_value_set(tmp_path: Path) -> None:
    sources = prepare_scb(tmp_path, "src", [VALUE_2020], VALUE_LISTS)
    db, issues = sources.build(tmp_path, "plain", scb_curation(tmp_path, "plain"))
    assert state_codes(db, "VALUE") == [("2020-01-01", "2020-12-31", None, None)]
    assert issue_codes(issues, "scb/sample/value") == ["conflicting_code_memberships"]


def test_compact_choice_applies_the_kept_list_to_every_reviewed_period(
    tmp_path: Path,
) -> None:
    sources = prepare_scb(tmp_path, "src", [VALUE_2020], VALUE_LISTS)
    digest = sources.evidence_sha256(sources.records(SCB))
    curation = scb_curation(
        tmp_path, "cur", choice("1.101", "VALUE", "keep", "other", digest)
    )
    db, issues = sources.build(tmp_path, "built", curation)
    assert state_codes(db, "VALUE") == [
        ("2020-01-01", "2020-06-30", "01", "Label"),
        ("2020-07-01", "2020-12-31", "01", "Label"),
    ]
    assert issues == []


def test_compact_choices_on_two_columns_check_each_columns_own_lists(
    tmp_path: Path,
) -> None:
    sources = prepare_scb(
        tmp_path,
        "src",
        [VALUE_2020, OTHER_2020],
        VALUE_LISTS + OTHER_LISTS,
        ("TESTREG|Testregistret|Individer|Individer|GenericVar|OTHER|2020|2020|0|0|0",),
    )
    records = sources.records(SCB)
    value = sources.evidence_sha256(member(records, "1001"))
    other = sources.evidence_sha256(member(records, "1002"))
    curation = scb_curation(
        tmp_path,
        "cur",
        # Both spellings stay one variable, so one application checks both columns.
        '[[variable]]\nnative_id = "1.101.a"\nslug = "value-a"\n'
        '[[identity.partition]]\nvariable = "1.101"\n'
        'columns = { VALUE = "1.101.a", OTHER = "1.101.a" }\n'
        'columns_ref = "exact fixture"\n'
        + choice("1.101.a", "VALUE", "keep", "other", value)
        + choice("1.101.a", "OTHER", "keep2", "other2", other),
    )
    db, issues = sources.build(tmp_path, "built", curation)
    assert [code for *_, code, _ in state_codes(db, "VALUE")] == ["01", "01"]
    assert [code for *_, code, _ in state_codes(db, "OTHER")] == ["11", "11"]
    assert stale_entries(issues) == []


def test_changed_competing_list_makes_the_reviewed_choice_stale(
    tmp_path: Path,
) -> None:
    reviewed = prepare_scb(tmp_path, "reviewed", [VALUE_2020], VALUE_LISTS)
    digest = reviewed.evidence_sha256(reviewed.records(SCB))
    changed = prepare_scb(
        tmp_path,
        "changed",
        [VALUE_2020],
        [VALUE_LISTS[0], "other|1|03|Other|1001|7003"],
    )
    curation = scb_curation(
        tmp_path, "cur", choice("1.101", "VALUE", "keep", "other", digest)
    )
    db, issues = changed.build(tmp_path, "built", curation)
    assert state_codes(db, "VALUE") == [("2020-01-01", "2020-12-31", None, None)]
    assert stale_entries(issues) == [VALUE_CHOICE + "1", VALUE_CHOICE + "2"]


def test_each_build_checks_a_choice_against_its_own_source_lists(
    tmp_path: Path,
) -> None:
    """One process builds two accepted inputs whose competing list changed between
    them; each build carries a choice reviewed for its own input, so both apply."""
    first = prepare_scb(tmp_path, "first", [VALUE_2020], VALUE_LISTS)
    second = prepare_scb(
        tmp_path,
        "second",
        [VALUE_2020],
        [VALUE_LISTS[0], "other|1|03|Other|1001|7003"],
    )
    for label, sources in (("first", first), ("second", second)):
        digest = sources.evidence_sha256(sources.records(SCB))
        curation = scb_curation(
            tmp_path, f"{label}-cur", choice("1.101", "VALUE", "keep", "other", digest)
        )
        db, issues = sources.build(tmp_path, label, curation)
        assert [code for *_, code, _ in state_codes(db, "VALUE")] == ["01", "01"]
        assert stale_entries(issues) == []


# -- Peer guards over an effective (errata-pooled) column -----------------------

BLANK_2010 = var_row(cvid=1010, var_id=101, colname="", year="2010", regver_id=110)
VALUE_2013 = var_row(cvid=1013, var_id=101, colname="VALUE", year="2013", regver_id=113)
DELIVERED_2010 = (
    '\n[[errata.delivered]]\nvariant = "people"\ncolumn = "VALUE"\n'
    'versions = ["2010"]\nevidence = "accepted delivery"\nnoted = "2026-09-28"\n'
)


def warning(digest: str, year: int = 2010) -> str:
    return (
        '\n[[coding.warning]]\nvariable = "1.101"\nvariant = "people"\n'
        f'column = "VALUE"\nperiods = [["{year}-01-01", "{year}-12-31"]]\n'
        'fields = ["data_type", "coding"]\n'
        f'expected_evidence_sha256 = "{digest}"\n'
        'data_warning = "Retained source metadata conflict"\n'
        'reason = "Preserve source type and own coding"\nsource = "Exact source rows"\n'
    )


def warning_periods(db: Path) -> list[tuple[str, str]]:
    with sqlite3.connect(db) as conn:
        return [
            (payload["valid_from"], payload["valid_to"])
            for (raw,) in conn.execute("SELECT warning_json FROM data_warning")
            for payload in (json.loads(raw),)
        ]


def test_coding_warning_peers_include_rows_pooled_into_the_column(
    tmp_path: Path,
) -> None:
    """The blank-named 2010 row joins VALUE through `[[errata.delivered]]`; a warning
    reviewed over both original rows applies to the pooled 2010 state."""
    sources = prepare_scb(tmp_path, "src", [BLANK_2010, VALUE_2013])
    digest = sources.evidence_sha256(sources.records(SCB))
    curation = scb_curation(tmp_path, "cur", DELIVERED_2010 + warning(digest))
    db, issues = sources.build(tmp_path, "built", curation)
    assert warning_periods(db) == [("2010-01-01", "2010-12-31")]
    assert stale_entries(issues) == []


def test_coding_warning_reviewed_without_the_pooled_row_is_stale(
    tmp_path: Path,
) -> None:
    sources = prepare_scb(tmp_path, "src", [BLANK_2010, VALUE_2013])
    donor_only = member(sources.records(SCB), "1013")
    curation = scb_curation(
        tmp_path,
        "cur",
        DELIVERED_2010 + warning(sources.evidence_sha256(donor_only)),
    )
    db, issues = sources.build(tmp_path, "built", curation)
    assert warning_periods(db) == []
    assert stale_entries(issues) == [
        "curation/registers/scb/sample.toml#/coding.warning/1/period/1"
    ]


# SCB delivers VALUE 2020 twice in one variant: a preliminary and a final edition.
# The SCB edition-name rule (DESIGN.md: a final edition supersedes `YYYY, preliminär
# version`, retaining its records as support) takes the preliminary row out of the
# column's catalog peers. Synthetic rows, modelled on that documented rule.
PRELIMINARY_2020 = var_row(
    cvid=1020,
    var_id=101,
    colname="VALUE",
    versionname="2020, preliminär version",
    regver_id=120,
)
FINAL_2020 = var_row(
    cvid=1021,
    var_id=101,
    colname="VALUE",
    versionname="2020, slutlig version",
    regver_id=121,
)


@pytest.mark.parametrize(
    "reviewed,stale",
    [(("1021",), False), (("1020", "1021"), True)],
    ids=["final-only", "with-preliminary"],
)
def test_coding_warning_peers_exclude_a_superseded_preliminary_row(
    tmp_path: Path, reviewed: tuple[str, ...], stale: bool
) -> None:
    sources = prepare_scb(tmp_path, "src", [PRELIMINARY_2020, FINAL_2020])
    records = tuple(
        record for cvid in reviewed for record in member(sources.records(SCB), cvid)
    )
    curation = scb_curation(
        tmp_path, "cur", warning(sources.evidence_sha256(records), year=2020)
    )
    db, issues = sources.build(tmp_path, "built", curation)
    assert warning_periods(db) == ([] if stale else [("2020-01-01", "2020-12-31")])
    assert stale_entries(issues) == (
        ["curation/registers/scb/sample.toml#/coding.warning/1/period/1"]
        if stale
        else []
    )


# -- Reviewed decimal-comma enumeration certificates (SOS) ----------------------

SINGLE = (
    "0=giltigt pnr,  8=Ogiltigt pnr",
    (("0", "giltigt pnr"), ("8", "Ogiltigt pnr")),
    (",",),
)
MIXED = (
    "0=(ogitligt), 1=enbart trygghetslarm; 9=uppgift saknas",
    (("0", "(ogitligt)"), ("1", "enbart trygghetslarm"), ("9", "uppgift saknas")),
    (",", ";"),
)
SOS_KON = (
    "('Socialstyrelsen/Metadata Syntetiskt register (SYN)_webb.xlsx', 'sos', "
    "'register', 'name', 'Syntetiskt register', 'variable', 'native-str', 'KON'"
)
DOCUMENTED = "curation/registers/sos/syn.toml#/coding.documented/1/period/1"


def sos_id(
    kind: Literal["register", "register_variant", "variable"],
    member: str | None = None,
) -> str:
    return authored_naming_id(
        kind, provider="sos", register_key="syn", member_key=member
    )


def prepare_sos(
    tmp_path: Path,
    label: str,
    cell: str,
    *,
    data_from: int = 2005,
    data_to: int | None = 2015,
    code_list: bool = False,
) -> Prepared:
    source = tmp_path / label / "source"
    write_scb_input(
        source,
        registerinformation_rows=[VALUE_2020],
        unika_rows=UNIKA,
        include=("registerinformation", "unika"),
    )
    syn = DEFAULT_REGISTERS[0]
    kon = replace(
        syn.variables[2], value_set=cell, data_from=data_from, data_to=data_to
    )
    path = next(
        write_sos_input(
            source, registers=(replace(syn, variables=(*syn.variables[:2], kon)),)
        ).glob("*.xlsx")
    )
    workbook = load_workbook(path)
    sheet = workbook["Metadata - Variabelnivå"]
    # A delivered blank Kopplingsvariabel is the provider's explicit "not an
    # identifier" claim; without the column KON's flags stay unknown and withheld.
    column = sheet.max_column + 1
    sheet.cell(row=1, column=column, value="Kopplingsvariabel")
    for row in range(2, sheet.max_row + 1):
        sheet.cell(row=row, column=column, value="")
    if code_list:
        codes = workbook.create_sheet("Kodlista_KON")
        codes.append(["Tidsperiod", "Kod", "Beskrivning"])
        codes.append(["2005-2015", "0", "giltigt pnr"])
    workbook.save(path)
    workbook.close()
    return _prepare(tmp_path, label, source)


def _toml(value: object) -> str:
    """Render a JSON value as one TOML inline value (the committed-file form)."""
    if isinstance(value, dict):
        return (
            "{"
            + ", ".join(
                f"{json.dumps(k)} = {_toml(v)}"
                for k, v in value.items()
                if v is not None
            )
            + "}"
        )
    if isinstance(value, list | tuple):
        return "[" + ", ".join(_toml(item) for item in value) + "]"
    if isinstance(value, bool):
        return "true" if value else "false"
    return json.dumps(value, ensure_ascii=False)


def certificate(
    sources: Prepared,
    pairs: tuple[tuple[str, str], ...],
    separators: tuple[str, ...],
    *,
    own_scope: bool = False,
    scope_end: str | None = None,
) -> str:
    """A `[[coding.documented]]` entry whose authority is the reviewed KON row."""
    opened = open_prepared_catalog_sources(
        sources.prepared, expected_sha256=sources.digest, input_commit=sources.commit
    )
    record = next(
        r
        for r in opened.records.records
        if r.source.startswith("Socialstyrelsen/") and r.subject.member.name == "KON"
    )
    revision = next(
        v
        for v in opened.records.manifest.revisions
        if v.revision_id == record.source_revision_id
    )
    authority = {
        "revision": revision.model_dump(mode="json"),
        "locators": [locator.model_dump(mode="json") for locator in record.locators],
        "records": [
            expectation.model_dump(mode="json")
            for expectation in capture_expectations(
                (record,),
                fields=tuple(SourceFields.model_fields),
                parents=True,
                coding=True,
            )
        ],
        "codings": [],
        "raw_codings": [],
    }
    authority["enumeration"] = {
        "field": "representation",
        "syntax": "ascii-decimal-comma-equals",
        "assignment_separators": list(separators),
        "lines": [f"{code}={label}" for code, label in pairs],
    }
    if own_scope:
        scope = record.edition_scope.model_dump(mode="json")
        if scope_end is not None:
            scope["intervals"][0]["end"] = scope_end
        authority["source_scope"] = scope
    return (
        "\n[[coding.documented]]\n"
        f'variable = "{sos_id("variable", "KON")}"\n'
        f'variant = "{sos_id("register_variant", "SYN_A")}"\n'
        'column = "KON"\n'
        + (
            "periods = []\n"
            if own_scope
            else 'periods = [["2005-01-01", "2015-12-31"]]\n'
        )
        + f"members = {_toml([list(pair) for pair in pairs])}\n"
        'version_label = "Exact source decimal meanings"\n'
        'reason = "Reviewed own enumeration"\nsource = "fixture"\n'
        f"source_authority = {_toml(authority)}\n"
    )


def sos_curation(tmp_path: Path, label: str, body: str = "") -> Path:
    curation = scb_curation(tmp_path, label)
    (curation / "registers" / "sos").mkdir()
    (curation / "registers" / "sos" / "syn.toml").write_text(
        '[register]\nprovider = "sos"\nslug = "syn"\n'
        f'name = "Syntetiskt register"\nnative_id = "{sos_id("register")}"\n'
        + "".join(
            f'[[variant]]\nslug = "{slug}"\n'
            f'native_id = "{sos_id("register_variant", member)}"\n'
            for member, slug in (("SYN_A", "syn-a"), ("SYN_B", "syn-b"))
        )
        + "".join(
            f'[[variable]]\nslug = "{member.lower()}"\n'
            f'native_id = "{sos_id("variable", member)}"\n'
            for member in ("DIAGNOS", "KON")
        )
        + body,
        encoding="utf-8",
    )
    return curation


@pytest.mark.parametrize("case", [SINGLE, MIXED], ids=["comma", "comma-semicolon"])
def test_unresolved_decimal_comma_cell_publishes_no_members_without_review(
    tmp_path: Path, case
) -> None:
    sources = prepare_sos(tmp_path, "src", case[0])
    db, issues = sources.build(tmp_path, "plain", sos_curation(tmp_path, "plain"))
    assert state_codes(db, "KON") == [("2005-01-01", "2015-12-31", None, None)]
    assert "unresolved_member_list" in issue_codes(issues, SOS_KON)


def test_reviewed_decimal_comma_certificate_states_exact_source_members(
    tmp_path: Path,
) -> None:
    cell, pairs, separators = SINGLE
    sources = prepare_sos(tmp_path, "src", cell)
    curation = sos_curation(tmp_path, "cur", certificate(sources, pairs, separators))
    db, issues = sources.build(tmp_path, "built", curation)
    assert state_codes(db, "KON") == [
        ("2005-01-01", "2015-12-31", code, label) for code, label in pairs
    ]
    assert stale_entries(issues) == []


def test_reviewed_certificate_with_its_own_open_source_scope_stays_open(
    tmp_path: Path,
) -> None:
    cell, pairs, separators = MIXED
    sources = prepare_sos(tmp_path, "src", cell, data_to=None)
    curation = sos_curation(
        tmp_path, "cur", certificate(sources, pairs, separators, own_scope=True)
    )
    db, issues = sources.build(tmp_path, "built", curation)
    assert state_codes(db, "KON") == [
        ("2005-01-01", "9999-12-31", code, label) for code, label in pairs
    ]
    assert stale_entries(issues) == []


@pytest.mark.parametrize("data_from,data_to", [(2006, None), (2005, 2024)])
def test_changed_source_scope_makes_an_own_scope_certificate_stale(
    tmp_path: Path, data_from: int, data_to: int | None
) -> None:
    cell, pairs, separators = MIXED
    reviewed = prepare_sos(tmp_path, "reviewed", cell, data_to=None)
    changed = prepare_sos(
        tmp_path, "changed", cell, data_from=data_from, data_to=data_to
    )
    curation = sos_curation(
        tmp_path, "cur", certificate(reviewed, pairs, separators, own_scope=True)
    )
    db, issues = changed.build(tmp_path, "built", curation)
    assert {code for *_, code, _ in state_codes(db, "KON")} == {None}
    assert DOCUMENTED in stale_entries(issues)


@pytest.mark.parametrize(
    "changed",
    [
        "0=giltigt pnr,  8=Ogiltigt pnr, 4=samordningsnummer",
        "0=giltigt pnr,  9=Ogiltigt pnr",
        "0=giltigt pnr,  8=Annat pnr",
        "Prefix 0=giltigt pnr,  8=Ogiltigt pnr",
    ],
    ids=["added-member", "changed-code", "changed-label", "prefixed"],
)
def test_changed_source_cell_makes_the_certificate_stale(
    tmp_path: Path, changed: str
) -> None:
    cell, pairs, separators = SINGLE
    reviewed = prepare_sos(tmp_path, "reviewed", cell)
    curation = sos_curation(tmp_path, "cur", certificate(reviewed, pairs, separators))
    db, issues = prepare_sos(tmp_path, "changed", changed).build(
        tmp_path, "built", curation
    )
    assert {code for *_, code, _ in state_codes(db, "KON")} == {None}
    assert DOCUMENTED in stale_entries(issues)


def test_certificate_refuses_a_column_that_now_has_a_source_code_list(
    tmp_path: Path,
) -> None:
    cell, pairs, separators = SINGLE
    reviewed = prepare_sos(tmp_path, "reviewed", cell)
    curation = sos_curation(tmp_path, "cur", certificate(reviewed, pairs, separators))
    listed = prepare_sos(tmp_path, "listed", cell, code_list=True)
    db, issues = listed.build(tmp_path, "built", curation)
    assert ("2005-01-01", "2015-12-31", "8", "Ogiltigt pnr") not in state_codes(
        db, "KON"
    )
    assert DOCUMENTED in stale_entries(issues)


@pytest.mark.parametrize(
    "members,scope_end,separators,reason",
    [
        (None, "9999-12-31", None, "requires one exact supplied interval"),
        (
            (("0", "giltigt pnr"), ("0", "Ogiltigt pnr")),
            None,
            None,
            "enumeration codes must be unique",
        ),
        (
            (("0", "giltigt pnr"), ("x", "Ogiltigt pnr")),
            None,
            None,
            "literal ASCII decimal code and label lines",
        ),
        (
            (("0", "giltigt pnr, annat"), ("8", "Ogiltigt pnr")),
            None,
            None,
            "literal ASCII decimal code and label lines",
        ),
        (None, None, (",", ","), "separators require literal decimal assignments"),
        (None, None, ("|",), "assignment_separators.*Input should be ',' or ';'"),
        (None, None, (), "assignment_separators.*at least 1 item"),
        (
            (("0", "giltigt pnr; annat"), ("8", "Ogiltigt pnr")),
            None,
            (",", ";"),
            "literal ASCII decimal code and label lines",
        ),
    ],
    ids=[
        "literal-open-end",
        "duplicate-code",
        "non-decimal-code",
        "label-separator",
        "repeated-separator",
        "unknown-separator",
        "no-separator",
        "label-second-separator",
    ],
)
def test_malformed_certificate_fails_curation_load_at_its_entry(
    tmp_path: Path,
    members: tuple[tuple[str, str], ...] | None,
    scope_end: str | None,
    separators: tuple[str, ...] | None,
    reason: str,
) -> None:
    """The entry's `members` and its enumeration lines agree, so only the
    malformed part named by the case id can reject it; `reason` pins that part's
    own located message (read from a run of each case, then checked against the
    validator that owns the rule)."""
    cell, pairs, default_separators = SINGLE
    sources = prepare_sos(tmp_path, "src", cell, data_to=None)
    body = certificate(
        sources,
        pairs if members is None else members,
        default_separators if separators is None else separators,
        own_scope=scope_end is not None,
        scope_end=scope_end,
    )
    with pytest.raises(RegMetaError) as caught:
        load_curation_tree(sos_curation(tmp_path, "cur", body))
    assert "curation/registers/sos/syn.toml [[coding.documented" in caught.value.message
    assert re.search(reason, caught.value.message)
