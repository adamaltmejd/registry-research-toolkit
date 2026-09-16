"""SCB reference declarations remain source facts, not catalog relationships."""

from __future__ import annotations

import hashlib
from typing import TYPE_CHECKING

import pytest
from _csv_fixtures import (
    PIPE,
    TIMESERIES_HEADER,
    timeseries_row,
    write_scb_input,
    write_scb_snapshot,
)
from openpyxl import Workbook
from pydantic import TypeAdapter
from reg_meta_build.input_snapshot import open_scb_snapshot
from reg_meta_build.source_records import SourceRevision
from reg_meta_build.source_reference_records import SourceReferenceDeclaration
from reg_meta_build.sources.scb_reference_records import (
    ScbReferenceSourceError,
    clean_timeseries_row,
    read_scb_column_types,
    read_scb_events,
    read_scb_join_keys,
)

if TYPE_CHECKING:
    from pathlib import Path


def _revision(path: Path) -> SourceRevision:
    payload = path.read_bytes()
    return SourceRevision.create(
        dataset="scb-reference-fixture",
        publisher="SCB",
        purpose="source fixture",
        upstream_revision="fixture-1",
        artifact_path=f"source/{path.name}",
        artifact_size=len(payload),
        artifact_sha256=hashlib.sha256(payload).hexdigest(),
    )


def _cells(*, action="Ersatt av", entity="Variabel"):
    raw = (
        " A  name ",
        action,
        " A paragraph\r\n\r\n  an indented line  ",
        entity,
        " 0007 ",
        "",
        None,
    )
    return {
        name: (value is not None, value, value or "")
        for name, value in zip(TIMESERIES_HEADER.split(PIPE), raw, strict=True)
    }


@pytest.mark.parametrize(
    "native,neutral",
    [
        ("Avslutad", "retired"),
        ("Tidsseriebrott", "series_break"),
        ("Ersatt av", "replaced_by"),
        ("Ersätter", "replaces"),
        ("Unknown action", "unknown"),
    ],
)
def test_event_action_decodes_without_reversing_or_parsing_native_tokens(
    tmp_path: Path, native: str, neutral: str
) -> None:
    path = tmp_path / "Timeseries.csv"
    path.write_bytes(b"fixture")
    event = clean_timeseries_row(
        TIMESERIES_HEADER.split(PIPE), 8, _cells(action=native), _revision(path)
    )
    assert event.action == neutral
    assert event.action_label.raw_value == native
    assert event.first_token.raw_value == " 0007 "
    assert event.first_token.interpreted_value == " 0007 "
    assert event.second_token.present and event.second_token.raw_value == ""
    assert not event.document_token.present and event.document_token.raw_value is None
    assert event.name.value == "A name"
    assert event.description.value == " A paragraph\n\n  an indented line"
    assert event.locator.physical_record == "row:8"
    assert (
        TypeAdapter(SourceReferenceDeclaration).validate_json(event.model_dump_json())
        == event
    )


@pytest.mark.parametrize(
    "native,neutral",
    [
        ("Register", "register"),
        ("RegisterVariant", "variant"),
        ("RegisterVersion", "edition"),
        ("Variabel", "variable"),
        ("AktuellVariabel", "member"),
        ("Population", "population"),
        ("VS_Class", "opaque"),
        ("PopContext", "opaque"),
        ("New entity", "opaque"),
    ],
)
def test_event_entity_kind_keeps_unsupported_native_kinds_opaque(
    tmp_path: Path, native: str, neutral: str
) -> None:
    path = tmp_path / "Timeseries.csv"
    path.write_bytes(b"fixture")
    event = clean_timeseries_row(
        TIMESERIES_HEADER.split(PIPE), 2, _cells(entity=native), _revision(path)
    )
    assert event.entity_kind == neutral
    assert event.entity_label.raw_value == native


def test_event_reader_uses_accepted_raw_pin_and_retains_duplicates(
    tmp_path: Path,
) -> None:
    raw = tmp_path / "raw"
    row = timeseries_row(
        handelse="Ersätter", entitet="AktuellVariabel", id1="0007", id2="8"
    )
    write_scb_input(raw, timeseries_rows=[row, row])
    path = raw / "SCB" / "Timeseries.csv"
    snapshot = open_scb_snapshot(write_scb_snapshot(tmp_path / "accepted", path.parent))
    revision = _revision(path)
    result = read_scb_events(snapshot, revision)
    assert len(result.declarations) == 2
    assert (
        result.declarations[0].locator.semantic_record_key
        == result.declarations[1].locator.semantic_record_key
    )
    assert (
        result.declarations[0].locator.physical_record
        != result.declarations[1].locator.physical_record
    )
    assert result.tables[0].rows[0].role == "header"
    assert len(result.tables[0].rows) == 3
    wrong = revision.model_copy(update={"artifact_size": revision.artifact_size + 1})
    with pytest.raises(ScbReferenceSourceError, match="declared revision"):
        read_scb_events(snapshot, wrong)


def _sql(tmp_path: Path, body: str, *, prefix: str = "", suffix: str = ""):
    path = tmp_path / "Tabelldefinitioner.sql"
    text = (
        prefix
        + "CREATE TABLE [dbo].[Delivery](\r\n"
        + body
        + "\r\n) ON [PRIMARY]\r\nGO\r\n"
        + suffix
    )
    path.write_bytes(text.encode("cp1252"))
    return path, _revision(path)


def test_sql_retains_comments_duplicate_columns_exact_types_and_unspecified_nullability(
    tmp_path: Path,
) -> None:
    comment = "/* Comment: CREATE TABLE [dbo].[Ignore]([X] [int]) ON [PRIMARY] */\r\n"
    path, revision = _sql(
        tmp_path,
        "[Kod] [varchar](0040) NOT NULL,\r\n[Kod] [varchar](0040) NULL,\r\n[År] [smallint]",
        prefix=comment,
    )
    result = read_scb_column_types(path, revision)
    assert len(result.declarations) == 3
    first, second, unknown = result.declarations
    assert first.table_name.value == "Delivery"
    assert first.column_name.value == second.column_name.value == "Kod"
    assert first.data_type.value == "varchar"
    assert first.declared_width.value == 40 and first.declared_width.raw_value == "0040"
    assert first.nullable.value is False and second.nullable.value is True
    assert unknown.column_name.value == "År" and unknown.data_type.value == "smallint"
    assert unknown.nullable.status == unknown.declared_width.status == "unknown"
    assert not unknown.delivered_cells[-1].present
    raw = result.tables[0].rows[0].cells[0].raw_value
    assert raw.encode("cp1252") == path.read_bytes()
    _, start, end = first.locator.physical_record.split(":")
    assert raw[int(start) : int(end)] == "[Kod] [varchar](0040) NOT NULL"


@pytest.mark.parametrize(
    "body,prefix,suffix",
    [
        ("[A] [decimal](9,2) NULL", "", ""),
        ("[A] [int] DEFAULT 0", "", ""),
        ("[A] [int], PRIMARY KEY ([A])", "", ""),
        ("[A] [int] [B] [int]", "", ""),
        ("[A] [int],", "", ""),
        ("", "", ""),
        ("[A] [int]", "DROP TABLE [Other];\n", ""),
        ("[A] [int]", "", "ALTER TABLE [Delivery] ADD [B] [int]"),
        ("[A] [int]", "", "/* unterminated"),
    ],
)
def test_sql_cannot_silently_skip_unsupported_syntax(
    tmp_path: Path, body, prefix, suffix
) -> None:
    path, revision = _sql(tmp_path, body, prefix=prefix, suffix=suffix)
    with pytest.raises(ScbReferenceSourceError, match="SQL"):
        read_scb_column_types(path, revision)


def _workbook(
    tmp_path: Path, *, extra_sheet=False, header=("Tabell", "ID-kolumn", "Beskrivning")
):
    path = tmp_path / "ID-kolumner.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Blad1"
    sheet.append(header)
    sheet.append((" Timeseries ", "ID1", "Possibly related to Table.Key; no assertion"))
    sheet.append((" Timeseries ", "ID1", "  Indented paragraph\n\n  second line  "))
    sheet.append((None, "=A2", ""))
    sheet["C2"].hyperlink = "https://example.invalid/evidence"
    if extra_sheet:
        workbook.create_sheet("Unexpected")
    workbook.save(path)
    workbook.close()
    return path, _revision(path)


def test_join_key_reader_preserves_prose_duplicates_cells_and_unknowns(
    tmp_path: Path,
) -> None:
    path, revision = _workbook(tmp_path)
    result = read_scb_join_keys(path, revision)
    first, second, incomplete = result.declarations
    assert first.table_name.value == second.table_name.value == "Timeseries"
    assert first.column_name.value == "ID1"
    assert first.description.value == "Possibly related to Table.Key; no assertion"
    assert second.description.value == "  Indented paragraph\n\n  second line"
    assert (
        first.delivered_cells[2].hyperlink_target == "https://example.invalid/evidence"
    )
    assert incomplete.table_name.status == incomplete.column_name.status == "unknown"
    assert incomplete.delivered_cells[1].raw_value == "=A2"
    assert incomplete.delivered_cells[1].storage_type == "f"
    assert first.locator.physical_cells == ("Blad1!A2", "Blad1!B2", "Blad1!C2")
    assert len(result.tables[0].rows) == 4


@pytest.mark.parametrize(
    "options",
    [
        {"extra_sheet": True},
        {"header": ("table", "column", "description")},
        {"header": ("Tabell", "ID-kolumn", "Beskrivning", "Extra")},
    ],
)
def test_join_key_unknown_workbook_contract_fails(tmp_path: Path, options) -> None:
    path, revision = _workbook(tmp_path, **options)
    with pytest.raises(ScbReferenceSourceError, match="ID-kolumner"):
        read_scb_join_keys(path, revision)


@pytest.mark.parametrize(
    "reader,fixture", [(read_scb_column_types, _sql), (read_scb_join_keys, _workbook)]
)
def test_side_readers_require_exact_revision(tmp_path: Path, reader, fixture) -> None:
    path, revision = (
        fixture(tmp_path, "[A] [int]") if fixture is _sql else fixture(tmp_path)
    )
    path.write_bytes(path.read_bytes() + b" ")
    with pytest.raises(ScbReferenceSourceError, match="declared revision"):
        reader(path, revision)
